import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from batch.lib.overview_batch_export import build_overview_batch_export
from batch.lib.overview_batch_export import batch_line_for_row
from batch.lib.data_summary_exporter import save_workbook_data


class OverviewBatchExportTests(unittest.TestCase):
    def test_builds_channel_and_token_batch_lines(self):
        workbook = Workbook()
        workbook.active.title = "总览"
        channels = workbook.create_sheet("test-渠道管理")
        channels.append(["名称", "分组", "类型", "状态", "模型列表", "模型重定向", "base_url", "备注", "密钥", "部署时间"])
        channels.append([
            "gateway-alpha", "gateway", "OpenAI", "已启用",
            "model-a, model-b", "model-a->upstream-a", "https://upstream.example.com/", "备注", "sk-test", "",
        ])
        tokens = workbook.create_sheet("test-令牌管理")
        tokens.append(["名称", "项目环境", "状态", "剩余额度", "分组", "可用模型", "包含渠道", "IP限制", "创建时间", "过期时间"])
        tokens.append(["token-alpha", "test", "已启用", "120/500", "gateway", "model-a、model-b", "gateway-alpha", "无限制", "", ""])

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "总览.xlsx"
            workbook.save(path)
            channels, channel_stats = build_overview_batch_export(path, "channels")
            tokens, token_stats = build_overview_batch_export(path, "tokens")

        self.assertEqual(channels, [
            "gateway-alpha\tgateway\tOpenAI\tmodel-a\tmodel-a->upstream-a\thttps://upstream.example.com\tsk-test",
            "gateway-alpha\tgateway\tOpenAI\tmodel-b\tmodel-b->model-b\thttps://upstream.example.com\tsk-test",
        ])
        self.assertEqual(channel_stats["rows"], 1)
        self.assertEqual(channel_stats["lines"], 2)
        self.assertEqual(tokens, ["token-alpha\tgateway\tmodel-a、model-b\t120"])
        self.assertEqual(token_stats["rows"], 1)
        self.assertEqual(token_stats["lines"], 1)

    def test_save_workbook_data_writes_batch_config_column(self):
        channel = {
            "名称": "gateway-alpha",
            "分组": "gateway",
            "类型": "OpenAI",
            "状态": "已启用",
            "模型列表": "model-a",
            "模型重定向": "",
            "base_url": "https://upstream.example.com/",
            "密钥": "sk-test",
            "部署时间": "",
        }
        token = {
            "名称": "token-alpha",
            "项目环境": "test",
            "状态": "已启用",
            "剩余额度": "120/500",
            "分组": "gateway",
            "可用模型": "model-a、model-b",
            "包含渠道": "gateway-alpha",
            "IP限制": "无限制",
            "创建时间": "",
            "过期时间": "",
        }

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "summary.xlsx"
            save_workbook_data(
                path,
                channel_env_rows={"test": [channel], "uat": [], "prod": []},
                token_env_rows={"test": [token], "uat": [], "prod": []},
            )
            from openpyxl import load_workbook

            workbook = load_workbook(path)
            channel_sheet = workbook["test-渠道管理"]
            token_sheet = workbook["test-令牌管理"]
            self.assertEqual(channel_sheet.cell(row=1, column=12).value, "批量配置")
            self.assertEqual(channel_sheet.cell(row=2, column=12).value, "gateway-alpha\tgateway\tOpenAI\tmodel-a\tmodel-a->model-a\thttps://upstream.example.com\tsk-test")
            self.assertEqual(token_sheet.cell(row=1, column=11).value, "批量配置")
            self.assertEqual(token_sheet.cell(row=2, column=11).value, "token-alpha\tgateway\tmodel-a、model-b\t120")
            workbook.close()

    def test_azure_channel_requires_api_version(self):
        row = {
            "名称": "azure-alpha",
            "分组": "azure",
            "类型": "Azure OpenAI",
            "状态": "已启用",
            "模型列表": "gpt-x",
            "模型重定向": "",
            "base_url": "https://azure.example.com",
            "密钥": "sk-test",
            "默认API版本": "2025-04-01-preview",
        }
        self.assertEqual(
            batch_line_for_row(row, "channels"),
            "azure-alpha\tazure\tAzure OpenAI\tgpt-x\tgpt-x->gpt-x\thttps://azure.example.com\t2025-04-01-preview\tsk-test",
        )

        row.pop("默认API版本")
        self.assertEqual(batch_line_for_row(row, "channels"), "")

    def test_models_self_map_when_redirect_is_blank(self):
        row = {
            "名称": "heditest-alpha",
            "分组": "heditest",
            "类型": "OpenAI",
            "状态": "已启用",
            "模型列表": "heditest-gpt-4, heditest-gpt-5",
            "模型重定向": "",
            "base_url": "https://heditest.example.com",
            "密钥": "sk-test",
            "部署时间": "",
        }

        self.assertEqual(
            batch_line_for_row(row, "channels"),
            "\n".join(
                [
                    "heditest-alpha\theditest\tOpenAI\theditest-gpt-4\theditest-gpt-4->heditest-gpt-4\thttps://heditest.example.com\tsk-test",
                    "heditest-alpha\theditest\tOpenAI\theditest-gpt-5\theditest-gpt-5->heditest-gpt-5\thttps://heditest.example.com\tsk-test",
                ]
            ),
        )

    def test_skips_disabled_rows_by_default(self):
        workbook = Workbook()
        workbook.active.title = "总览"
        channels = workbook.create_sheet("uat-渠道管理")
        channels.append(["名称", "分组", "类型", "状态", "模型列表", "模型重定向", "base_url", "备注", "密钥", "部署时间"])
        channels.append(["disabled", "uat", "OpenAI", "已禁用", "model", "model->model", "https://upstream.example.com", "", "sk-test", ""])

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "总览.xlsx"
            workbook.save(path)
            lines, stats = build_overview_batch_export(path, "channels")

        self.assertEqual(lines, [])
        self.assertEqual(stats["skipped_disabled"], 1)


if __name__ == "__main__":
    unittest.main()
