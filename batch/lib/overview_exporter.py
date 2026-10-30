# -*- coding: utf-8 -*-
"""⑥ 一表总览：把③ 网关汇总与⑤ 两侧上游写入同一工作簿。"""

from __future__ import annotations

import shutil
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from .data_summary import (
    ENVS,
    SHEET_CHANNELS,
    SHEET_DUP_CHANNELS,
    SHEET_DUP_TOKENS,
    SHEET_TOKENS,
)
from .data_summary_exporter import excel_dir, latest_excel_path, pick_latest_path, sort_excel_files_latest_first
from .data_summary_exporter import apply_channel_remark_mismatch_style
from .deploy_time import sort_rows_by_deploy_time_desc
from .upstream_azure import AZURE_HEADERS
from .upstream_exporter import (
    SHEET_AZURE,
    SHEET_AZURE_LEGACY,
    SHEET_VOLC,
    _apply_style,
    _write_sheet,
    latest_azure_upstream_excel_path,
    latest_combined_upstream_excel_path,
    latest_upstream_excel_path,
)
from .upstream_volc import UPSTREAM_HEADERS

OVERVIEW_STEM = "ai网关令牌及渠道列表_一表总览"
OVERVIEW_STEM_LEGACY = "一表总览"
SHEET_INDEX = "总览"
TIMESTAMP_FORMAT = "%Y%m%d%H%M%S"
INDEX_HEADERS = ["分类", "工作表", "行数", "说明"]
_UPSTREAM_SHEET_ORDER = (SHEET_VOLC, SHEET_AZURE, SHEET_AZURE_LEGACY)


def build_overview_filename(when: datetime | None = None) -> str:
    stamp = (when or datetime.now()).strftime(TIMESTAMP_FORMAT)
    return f"{OVERVIEW_STEM}_{stamp}.xlsx"


def build_overview_path(root_dir: Path, when: datetime | None = None) -> Path:
    return excel_dir(root_dir) / build_overview_filename(when)


def list_overview_excel_files(root_dir: Path) -> list[Path]:
    directory = excel_dir(root_dir)
    if not directory.exists():
        return []
    found: dict[str, Path] = {}
    for pattern in (f"{OVERVIEW_STEM}_*.xlsx", f"{OVERVIEW_STEM_LEGACY}_*.xlsx"):
        for path in directory.glob(pattern):
            found[str(path.resolve())] = path
    return sort_excel_files_latest_first(list(found.values()))


def latest_overview_excel_path(root_dir: Path) -> Path | None:
    files = list_overview_excel_files(root_dir)
    return files[0] if files else None


def _sheet_row_count(ws) -> int:
    return max(0, int(ws.max_row or 1) - 1)


def load_sheet_rows(path: Path | None, preferred_names: list[str]) -> list[dict[str, Any]]:
    """从已有 Excel 读取指定工作表为 dict 行（无文件则空）。"""
    if not path or not Path(path).exists():
        return []
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        title = next((name for name in preferred_names if name in wb.sheetnames), "")
        if not title:
            return []
        ws = wb[title]
        rows_iter = ws.iter_rows(values_only=True)
        headers = [str(h or "").strip() for h in next(rows_iter, [])]
        out: list[dict[str, Any]] = []
        for row in rows_iter:
            if not row or not any(cell not in (None, "") for cell in row):
                continue
            item: dict[str, Any] = {}
            for idx, header in enumerate(headers):
                if not header:
                    continue
                value = row[idx] if idx < len(row) else ""
                item[header] = "" if value is None else value
            out.append(item)
        return out
    finally:
        wb.close()


def resolve_gateway_source(root_dir: Path, explicit: Path | None = None) -> Path | None:
    if explicit and Path(explicit).exists():
        return Path(explicit)
    return pick_latest_path(latest_excel_path(root_dir), latest_overview_excel_path(root_dir))


def _volc_candidate_paths(root_dir: Path) -> list[Path]:
    return [
        path
        for path in (
            latest_combined_upstream_excel_path(root_dir),
            latest_upstream_excel_path(root_dir),
            latest_overview_excel_path(root_dir),
        )
        if path
    ]


def _azure_candidate_paths(root_dir: Path) -> list[Path]:
    return [
        path
        for path in (
            latest_combined_upstream_excel_path(root_dir),
            latest_azure_upstream_excel_path(root_dir),
            latest_overview_excel_path(root_dir),
        )
        if path
    ]


def resolve_volc_source_path(root_dir: Path) -> Path | None:
    for path in sort_excel_files_latest_first(_volc_candidate_paths(root_dir)):
        if load_sheet_rows(path, [SHEET_VOLC]):
            return path
    return None


def resolve_azure_source_path(root_dir: Path) -> Path | None:
    for path in sort_excel_files_latest_first(_azure_candidate_paths(root_dir)):
        if load_sheet_rows(path, [SHEET_AZURE, SHEET_AZURE_LEGACY]):
            return path
    return None


def resolve_volc_rows(
    root_dir: Path,
    live_rows: list[dict[str, Any]] | None = None,
    source: Path | None = None,
) -> list[dict[str, Any]]:
    if live_rows:
        return list(live_rows)
    if source and Path(source).exists():
        return load_sheet_rows(Path(source), [SHEET_VOLC])
    path = resolve_volc_source_path(root_dir)
    return load_sheet_rows(path, [SHEET_VOLC]) if path else []


def resolve_azure_rows(
    root_dir: Path,
    live_rows: list[dict[str, Any]] | None = None,
    source: Path | None = None,
) -> list[dict[str, Any]]:
    if live_rows:
        return list(live_rows)
    if source and Path(source).exists():
        return load_sheet_rows(Path(source), [SHEET_AZURE, SHEET_AZURE_LEGACY])
    path = resolve_azure_source_path(root_dir)
    return load_sheet_rows(path, [SHEET_AZURE, SHEET_AZURE_LEGACY]) if path else []


def describe_overview_sources(root_dir: Path) -> dict[str, str]:
    gateway = resolve_gateway_source(root_dir)
    volc = resolve_volc_source_path(root_dir)
    azure = resolve_azure_source_path(root_dir)
    return {
        "gateway": gateway.name if gateway else "",
        "volc": volc.name if volc else "",
        "azure": azure.name if azure else "",
    }


def _sheet_category(title: str) -> tuple[str, str]:
    if title in SHEET_CHANNELS.values():
        env = next((e for e, name in SHEET_CHANNELS.items() if name == title), "")
        return "网关渠道", f"③ {env or ''} 环境渠道".strip()
    if title in SHEET_TOKENS.values():
        env = next((e for e, name in SHEET_TOKENS.items() if name == title), "")
        return "网关令牌", f"③ {env or ''} 环境令牌".strip()
    if title == SHEET_DUP_CHANNELS:
        return "网关对照", "③ 跨环境重复渠道"
    if title == SHEET_DUP_TOKENS:
        return "网关对照", "③ 跨环境重复令牌"
    if title == SHEET_VOLC:
        return "上游", "⑤ 火山云接入点"
    if title in {SHEET_AZURE, SHEET_AZURE_LEGACY}:
        return "上游", "⑤ Azure OpenAI 部署"
    if title == SHEET_INDEX:
        return "目录", "本工作簿工作表清单"
    return "其它", ""


def _write_index_sheet(wb) -> None:
    if SHEET_INDEX in wb.sheetnames:
        del wb[SHEET_INDEX]
    ws = wb.create_sheet(SHEET_INDEX, 0)
    rows: list[dict[str, Any]] = []
    for title in wb.sheetnames:
        if title == SHEET_INDEX:
            continue
        category, note = _sheet_category(title)
        rows.append(
            {
                "分类": category,
                "工作表": title,
                "行数": _sheet_row_count(wb[title]),
                "说明": note,
            }
        )
    _write_sheet(ws, INDEX_HEADERS, rows)


def _order_overview_sheets(wb) -> None:
    """总览页固定第一，上游表固定排在所有环境表之后。"""
    existing = {ws.title: ws for ws in wb._sheets}
    ordered = []
    if SHEET_INDEX in existing:
        ordered.append(existing[SHEET_INDEX])
    ordered.extend(
        ws
        for title, ws in existing.items()
        if title not in {SHEET_INDEX, *_UPSTREAM_SHEET_ORDER}
    )
    ordered.extend(existing[title] for title in _UPSTREAM_SHEET_ORDER if title in existing)
    wb._sheets = ordered


def _ensure_channel_remark_columns(wb) -> None:
    """兼容旧版网关汇总：总览渠道表始终在 base_url 后保留备注列。"""
    for title in [*SHEET_CHANNELS.values(), SHEET_DUP_CHANNELS]:
        if title not in wb.sheetnames:
            continue
        ws = wb[title]
        headers = [str(cell.value or "").strip() for cell in ws[1]]
        if "base_url" not in headers or "备注" in headers:
            continue
        base_col = headers.index("base_url") + 1
        remark_col = base_col + 1
        ws.insert_cols(remark_col)
        for row in range(1, (ws.max_row or 1) + 1):
            source = ws.cell(row=row, column=base_col)
            target = ws.cell(row=row, column=remark_col)
            target._style = copy(source._style)
            target.number_format = source.number_format
            target.protection = copy(source.protection)
            target.alignment = copy(source.alignment)
            if row == 1:
                target.value = "备注"

        if ws.auto_filter.ref:
            ws.auto_filter.ref = f"A1:{ws.cell(row=ws.max_row, column=ws.max_column).coordinate}"
        apply_channel_remark_mismatch_style(ws)

    for title in [*SHEET_CHANNELS.values(), SHEET_DUP_CHANNELS]:
        if title in wb.sheetnames:
            apply_channel_remark_mismatch_style(wb[title])


def save_overview_workbook(
    path: Path,
    *,
    gateway_path: Path | None = None,
    volc_rows: list[dict[str, Any]] | None = None,
    azure_rows: list[dict[str, Any]] | None = None,
) -> Path:
    """同一工作簿：③ 全部工作表 + 火山云 + Azure OpenAI + 总览目录。"""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    gateway = Path(gateway_path) if gateway_path else None
    if gateway and gateway.exists() and gateway.resolve() != path.resolve():
        shutil.copy2(gateway, path)
        wb = load_workbook(path)
    elif path.exists() and gateway and gateway.resolve() == path.resolve():
        wb = load_workbook(path)
    else:
        wb = Workbook()
        placeholder = wb.active
        placeholder.title = "_empty"

    _ensure_channel_remark_columns(wb)

    for name in list(wb.sheetnames):
        if name in {SHEET_VOLC, SHEET_AZURE, SHEET_AZURE_LEGACY, SHEET_INDEX, "_empty"}:
            del wb[name]

    ws_volc = wb.create_sheet(SHEET_VOLC)
    _write_sheet(ws_volc, list(UPSTREAM_HEADERS), sort_rows_by_deploy_time_desc(volc_rows))
    ws_azure = wb.create_sheet(SHEET_AZURE)
    _write_sheet(ws_azure, list(AZURE_HEADERS), sort_rows_by_deploy_time_desc(azure_rows))
    _write_index_sheet(wb)
    _order_overview_sheets(wb)

    if not wb.sheetnames:
        ws = wb.create_sheet(SHEET_INDEX)
        _apply_style(ws, INDEX_HEADERS)

    wb.save(path)
    wb.close()
    return path


def refresh_overview_index(path: Path) -> None:
    """补密钥后刷新首页「总览」行数，不改其它工作表。"""
    target = Path(path)
    if not target.exists():
        return
    wb = load_workbook(target)
    try:
        _write_index_sheet(wb)
        _order_overview_sheets(wb)
        wb.save(target)
    finally:
        wb.close()


def overview_workbook_stats(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {"rows": 0, "sheet": SHEET_INDEX, "sheets": {}, "total_rows": 0}
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets: dict[str, int] = {}
        for name in wb.sheetnames:
            sheets[name] = _sheet_row_count(wb[name])
        gateway_sheets = [
            name
            for name in wb.sheetnames
            if name in {*SHEET_CHANNELS.values(), *SHEET_TOKENS.values(), SHEET_DUP_CHANNELS, SHEET_DUP_TOKENS}
        ]
        return {
            "rows": sum(v for k, v in sheets.items() if k != SHEET_INDEX),
            "sheet": SHEET_INDEX,
            "sheets": sheets,
            "total_rows": sum(v for k, v in sheets.items() if k != SHEET_INDEX),
            "gateway_sheets": gateway_sheets,
            "volc_rows": int(sheets.get(SHEET_VOLC) or 0),
            "azure_rows": int(sheets.get(SHEET_AZURE) or sheets.get(SHEET_AZURE_LEGACY) or 0),
            "envs": [env for env in ENVS if SHEET_CHANNELS[env] in sheets],
        }
    finally:
        wb.close()
