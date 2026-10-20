"""一表总览到批量配置 TXT 的只读转换。"""

from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .data_summary import ENVS, SHEET_CHANNELS, SHEET_TOKENS


_TAB_OR_NEWLINE = re.compile(r"[\t\r\n]+")
_STATUS_DISABLED = {"2", "3", "4", "已禁用", "自动禁用", "已过期", "已耗尽", "false", "0"}
_UNLIMITED_MODELS = {"无限制", "不限", "unlimited", "all"}
_API_VERSION_RE = re.compile(r"^\d{4}-\d{2}-\d{2}(-preview)?$")


def _text(value: Any) -> str:
    if value is None:
        return ""
    return _TAB_OR_NEWLINE.sub(" ", str(value)).strip()


def _enabled(value: Any) -> bool:
    text = _text(value).strip().lower()
    if not text:
        return True
    if text == "已启用":
        return True
    try:
        return int(float(text)) == 1
    except (TypeError, ValueError):
        return text not in _STATUS_DISABLED


def _cell(row: dict[str, Any], *names: str) -> str:
    for name in names:
        if name in row:
            value = _text(row.get(name))
            if value:
                return value
    return ""


def _read_rows(path: Path, title: str) -> list[dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if title not in wb.sheetnames:
            return []
        ws = wb[title]
        rows = ws.iter_rows(values_only=True)
        headers = [_text(value) for value in next(rows, ())]
        result: list[dict[str, Any]] = []
        for values in rows:
            if not values or not any(value not in (None, "") for value in values):
                continue
            row = {
                header: values[index] if index < len(values) else ""
                for index, header in enumerate(headers)
                if header
            }
            if _text(row.get("名称")):
                result.append(row)
        return result
    finally:
        wb.close()


def _mapping(raw: str) -> dict[str, str]:
    if not raw:
        return {}
    try:
        value = json.loads(raw)
        if isinstance(value, dict):
            return {str(k).strip(): str(v).strip() for k, v in value.items() if str(k).strip() and str(v).strip()}
    except (json.JSONDecodeError, TypeError):
        pass
    result: dict[str, str] = {}
    for part in re.split(r"[,，;；|\n\r]+", raw):
        if "->" not in part:
            continue
        source, target = part.split("->", 1)
        source_text, target_text = _text(source), _text(target)
        if source_text and target_text:
            result[source_text] = target_text
    return result


def _token_quota(raw: str) -> str:
    text = _text(raw).replace(",", "")
    if not text or text.lower() in {"无限额度", "无限", "unlimited"}:
        return ""
    if "/" in text:
        text = text.split("/", 1)[0].strip()
    try:
        value = int(float(text))
    except (TypeError, ValueError):
        return ""
    return str(max(value, 0))


def _channel_lines(path: Path, *, include_disabled: bool) -> tuple[list[str], dict[str, int]]:
    lines: list[str] = []
    stats = {"rows": 0, "lines": 0, "skipped_disabled": 0, "skipped_missing_key": 0, "skipped_incomplete": 0}
    for env in ENVS:
        for row in _read_rows(path, SHEET_CHANNELS[env]):
            stats["rows"] += 1
            if not include_disabled and not _enabled(row.get("状态")):
                stats["skipped_disabled"] += 1
                continue
            name = _cell(row, "名称")
            group = _cell(row, "分组")
            channel_type = _cell(row, "类型")
            models_raw = _cell(row, "模型列表", "模型", "models")
            models = [item.strip() for item in re.split(r"[,，;；|]+", models_raw) if item.strip()]
            base_url = _cell(row, "base_url", "Base URL").rstrip("/")
            redirect = _mapping(_cell(row, "模型重定向", "model_mapping"))
            key = _cell(row, "密钥", "秘钥", "key")
            api_version = _cell(row, "默认API版本", "api_version", "default_api_version", "other")
            if not name or not group or not channel_type or not models or not base_url:
                stats["skipped_incomplete"] += 1
                continue
            if not key:
                stats["skipped_missing_key"] += 1
                continue
            if _is_azure(channel_type) and not api_version:
                stats["skipped_incomplete"] += 1
                continue
            for model in models:
                target = redirect.get(model) or model
                parts = [name, group, channel_type, model, f"{model}->{target}", base_url]
                if api_version:
                    parts.append(api_version)
                parts.append(key)
                lines.append("\t".join(part.strip() for part in parts))
                stats["lines"] += 1
    return lines, stats


def _is_azure(channel_type: str) -> bool:
    return channel_type.strip().lower() in {"azure openai", "azure", "azure_openai", "3"}


def _token_lines(path: Path, *, include_disabled: bool) -> tuple[list[str], dict[str, int]]:
    lines: list[str] = []
    stats = {"rows": 0, "lines": 0, "skipped_disabled": 0, "skipped_unlimited_models": 0, "skipped_incomplete": 0}
    for env in ENVS:
        for row in _read_rows(path, SHEET_TOKENS[env]):
            stats["rows"] += 1
            if not include_disabled and not _enabled(row.get("状态")):
                stats["skipped_disabled"] += 1
                continue
            name = _cell(row, "名称")
            group = _cell(row, "分组")
            models_raw = _cell(row, "可用模型", "模型限制", "token_model_limit")
            if not name or not group:
                stats["skipped_incomplete"] += 1
                continue
            if not models_raw or models_raw.lower() in _UNLIMITED_MODELS:
                stats["skipped_unlimited_models"] += 1
                continue
            models = [item.strip() for item in re.split(r"[,，;；|]+", models_raw) if item.strip()]
            if not models:
                stats["skipped_incomplete"] += 1
                continue
            quota = _token_quota(_cell(row, "剩余额度", "额度", "额度(tokens)", "remain_quota"))
            parts = [name, group, "、".join(models)]
            if quota:
                parts.append(quota)
            lines.append("\t".join(parts))
            stats["lines"] += 1
    return lines, stats


def build_overview_batch_export(
    path: Path,
    kind: str,
    *,
    include_disabled: bool = False,
) -> tuple[list[str], dict[str, Any]]:
    target = Path(path)
    if not target.exists():
        raise FileNotFoundError("Excel 文件不存在")
    normalized = str(kind or "").strip().lower()
    if normalized == "channels":
        lines, stats = _channel_lines(target, include_disabled=include_disabled)
    elif normalized == "tokens":
        lines, stats = _token_lines(target, include_disabled=include_disabled)
    else:
        raise ValueError("kind 仅支持 channels 或 tokens")
    metadata = {
        "kind": normalized,
        "file": str(target.name),
        "include_disabled": bool(include_disabled),
        **stats,
    }
    return lines, metadata


def build_batch_export_filename(kind: str) -> str:
    label = "渠道批量配置" if str(kind).lower() == "channels" else "令牌批量配置"
    return f"一表总览_{label}_{datetime.now().strftime('%Y%m%d%H%M%S')}.txt"
