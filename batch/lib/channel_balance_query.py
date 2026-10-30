"""渠道余额查询：已用额度 / 剩余余额筛选。"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any

from .api_client import NewAPIClient, authenticate_client
from .channel_types import get_channel_types, resolve_channel_type
from .log_exporter import DEFAULT_QUOTA_PER_UNIT

BALANCE_EXCEL_STEM = "渠道余额"
BALANCE_TIMESTAMP_FORMAT = "%Y%m%d%H%M%S"
_LEGACY_BALANCE_PREFIX = "channel_balance_"


def excel_dir(root_dir: Path) -> Path:
    return Path(root_dir) / "excel"


def build_balance_filename(when: datetime | None = None) -> str:
    stamp = (when or datetime.now()).strftime(BALANCE_TIMESTAMP_FORMAT)
    return f"{BALANCE_EXCEL_STEM}_{stamp}.xlsx"


def list_balance_excel_files(root_dir: Path) -> list[Path]:
    directory = excel_dir(root_dir)
    if not directory.exists():
        return []
    files = list(directory.glob(f"{BALANCE_EXCEL_STEM}_*.xlsx"))
    files.extend(directory.glob(f"{_LEGACY_BALANCE_PREFIX}*.xlsx"))
    uniq = {path.resolve(): path for path in files if path.is_file()}
    return sorted(uniq.values(), key=lambda path: path.stat().st_mtime, reverse=True)


def latest_balance_excel_path(root_dir: Path) -> Path | None:
    files = list_balance_excel_files(root_dir)
    return files[0] if files else None

# 常用区间预设（单位：USD）→ (min_inclusive, max_inclusive)；None 表示不限
USED_RANGE_PRESETS: dict[str, tuple[float | None, float | None]] = {
    "all": (None, None),
    "eq0": (0.0, 0.0),
    "gt0": (1e-9, None),
    "le1": (None, 1.0),
    "le5": (None, 5.0),
    "ge1": (1.0, None),
    "ge10": (10.0, None),
    "ge50": (50.0, None),
    "ge100": (100.0, None),
    "r1_10": (1.0, 10.0),
    "r10_100": (10.0, 100.0),
    "ge1000": (1000.0, None),
}

REMAIN_RANGE_PRESETS: dict[str, tuple[float | None, float | None]] = {
    "all": (None, None),
    "eq0": (0.0, 0.0),
    "gt0": (1e-9, None),
    "le1": (None, 1.0),
    "le5": (None, 5.0),
    "lt1": (None, 0.999999),
    "lt10": (None, 9.999999),
    "ge1": (1.0, None),
    "ge10": (10.0, None),
    "ge100": (100.0, None),
}

USED_PRESET_LABELS: dict[str, str] = {
    "all": "全部",
    "eq0": "= $0",
    "gt0": "> $0",
    "le1": "≤ $1",
    "le5": "≤ $5",
    "ge1": "≥ $1",
    "ge10": "≥ $10",
    "ge50": "≥ $50",
    "ge100": "≥ $100",
    "r1_10": "$1 ~ $10",
    "r10_100": "$10 ~ $100",
    "ge1000": "≥ $1000",
}

REMAIN_PRESET_LABELS: dict[str, str] = {
    "all": "全部",
    "eq0": "= $0（耗尽）",
    "gt0": "> $0",
    "le1": "≤ $1",
    "le5": "≤ $5",
    "lt1": "< $1",
    "lt10": "< $10",
    "ge1": "≥ $1",
    "ge10": "≥ $10",
    "ge100": "≥ $100",
}

# 配置时间快捷项：含今天在内的近 N 个自然日（1 天 = 今天 00:00 ~ 23:59:59）
TIME_PRESET_DAYS: dict[str, int | None] = {
    "all": None,
    "1d": 1,
    "3d": 3,
    "5d": 5,
    "15d": 15,
    "30d": 30,
    "60d": 60,
    "180d": 180,
    "365d": 365,
}

TIME_PRESET_LABELS: dict[str, str] = {
    "all": "全部",
    "1d": "1天",
    "3d": "3天",
    "5d": "5天",
    "15d": "15天",
    "30d": "30天",
    "60d": "60天",
    "180d": "180天",
    "365d": "365天",
}

STATUS_LABELS = {
    1: "已启用",
    2: "已禁用",
    3: "自动禁用",
}


@dataclass
class ChannelBalanceQueryResult:
    total: int = 0
    matched_count: int = 0
    query_used: bool = True
    query_remain: bool = True
    quota_per_unit: float = DEFAULT_QUOTA_PER_UNIT
    channel_type: str = ""
    type_id: int | None = None
    used_preset: str = "all"
    remain_preset: str = "all"
    time_preset: str = "all"
    items: list[dict[str, Any]] = field(default_factory=list)
    presets: dict[str, Any] = field(default_factory=dict)
    extra: dict[str, Any] = field(default_factory=dict)


def _to_float(value: Any, default: float = 0.0) -> float:
    try:
        if value in (None, ""):
            return default
        return float(value)
    except (TypeError, ValueError):
        return default


def used_quota_to_usd(used_quota: Any, quota_per_unit: float) -> float:
    unit = quota_per_unit or DEFAULT_QUOTA_PER_UNIT
    if unit <= 0:
        unit = DEFAULT_QUOTA_PER_UNIT
    return _to_float(used_quota) / unit


def format_usd(value: float) -> str:
    if abs(value) < 1e-9:
        return "$0"
    if abs(value) >= 100:
        return f"${value:,.2f}"
    if abs(value) >= 1:
        return f"${value:.2f}"
    return f"${value:.4f}".rstrip("0").rstrip(".") if "." in f"{value:.4f}" else f"${value:.4f}"


def format_unix_time(value: Any) -> str:
    """将 Unix 时间戳格式化为可读时间；无效/0 返回空串。"""
    try:
        ts = int(value)
    except (TypeError, ValueError):
        return ""
    if ts <= 0:
        return ""
    try:
        return datetime.fromtimestamp(ts).strftime("%Y-%m-%d %H:%M:%S")
    except (OSError, OverflowError, ValueError):
        return ""


def resolve_range(
    preset: str,
    presets: dict[str, tuple[float | None, float | None]],
    min_usd: float | None,
    max_usd: float | None,
) -> tuple[float | None, float | None]:
    """优先使用自定义 min/max；否则用预设。"""
    if min_usd is not None or max_usd is not None:
        return min_usd, max_usd
    key = (preset or "all").strip() or "all"
    return presets.get(key, presets["all"])


def in_range(value: float, min_v: float | None, max_v: float | None) -> bool:
    if min_v is not None and max_v is not None and abs(min_v - max_v) < 1e-12:
        # 等值区间（如 = $0）
        return abs(value - min_v) < 1e-6
    if min_v is not None and value < min_v - 1e-12:
        return False
    if max_v is not None and value > max_v + 1e-12:
        return False
    return True


def parse_query_timestamp(raw: Any) -> int | None:
    """解析 Unix 秒或常见日期字符串；空值返回 None。"""
    if raw in (None, ""):
        return None
    if isinstance(raw, bool):
        raise ValueError(f"时间无效: {raw}")
    if isinstance(raw, (int, float)):
        ts = int(raw)
        return ts if ts > 0 else None
    text = str(raw).strip()
    if not text:
        return None
    try:
        ts = int(float(text))
        return ts if ts > 0 else None
    except (TypeError, ValueError):
        pass
    normalized = text.replace("Z", "").replace("/", "-")
    if "T" in normalized:
        normalized = normalized.replace("T", " ", 1)
    if len(normalized) >= 19:
        normalized = normalized[:19]
    elif len(normalized) == 16:
        normalized += ":00"
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            return int(datetime.strptime(normalized, fmt).timestamp())
        except ValueError:
            continue
    raise ValueError(f"时间无效: {raw}")


def resolve_created_time_range(
    preset: str = "all",
    start_ts: int | None = None,
    end_ts: int | None = None,
    *,
    now: datetime | None = None,
) -> tuple[int | None, int | None]:
    """自定义起止优先；否则按快捷天数（含今天的近 N 个自然日）。"""
    if start_ts is not None or end_ts is not None:
        start_n = int(start_ts) if start_ts is not None else None
        end_n = int(end_ts) if end_ts is not None else None
        if start_n is not None and start_n <= 0:
            start_n = None
        if end_n is not None and end_n <= 0:
            end_n = None
        if start_n is not None and end_n is not None and start_n > end_n:
            raise ValueError("开始时间不能晚于结束时间")
        return start_n, end_n
    key = (preset or "all").strip() or "all"
    days = TIME_PRESET_DAYS.get(key)
    if days is None:
        return None, None
    current = now or datetime.now()
    start_dt = (current - timedelta(days=int(days) - 1)).replace(
        hour=0, minute=0, second=0, microsecond=0
    )
    end_dt = current.replace(hour=23, minute=59, second=59, microsecond=0)
    return int(start_dt.timestamp()), int(end_dt.timestamp())


def in_created_time_range(
    created_time: Any,
    start_ts: int | None,
    end_ts: int | None,
) -> bool:
    if start_ts is None and end_ts is None:
        return True
    try:
        ts = int(created_time or 0)
    except (TypeError, ValueError):
        ts = 0
    if ts <= 0:
        return False
    if start_ts is not None and ts < start_ts:
        return False
    if end_ts is not None and ts > end_ts:
        return False
    return True


def format_time_range_label(start_ts: int | None, end_ts: int | None, preset: str = "all") -> str:
    if start_ts is None and end_ts is None:
        return TIME_PRESET_LABELS.get((preset or "all").strip() or "all", "全部")
    start_s = format_unix_time(start_ts) or "不限"
    end_s = format_unix_time(end_ts) or "不限"
    return f"{start_s} ~ {end_s}"


def channel_balance_row(
    channel: dict[str, Any],
    *,
    quota_per_unit: float,
    type_label: str = "",
) -> dict[str, Any]:
    used_raw = channel.get("used_quota")
    remain_raw = channel.get("balance")
    used_usd = used_quota_to_usd(used_raw, quota_per_unit)
    remain_usd = _to_float(remain_raw)
    status = int(channel.get("status") or 0)
    created_raw = channel.get("created_time") or 0
    balance_updated_raw = channel.get("balance_updated_time") or 0
    return {
        "id": int(channel.get("id") or 0),
        "name": str(channel.get("name") or ""),
        "group": str(channel.get("group") or ""),
        "type": int(channel.get("type") or 0),
        "type_label": type_label or str(channel.get("type") or ""),
        "status": status,
        "status_label": STATUS_LABELS.get(status, f"状态{status}"),
        "used_quota": int(_to_float(used_raw)),
        "used_usd": round(used_usd, 6),
        "used_usd_label": format_usd(used_usd),
        "balance": remain_usd,
        "remain_usd": round(remain_usd, 6),
        "remain_usd_label": format_usd(remain_usd),
        "created_time": int(_to_float(created_raw)),
        "created_time_label": format_unix_time(created_raw) or "—",
        "balance_updated_time": int(_to_float(balance_updated_raw)),
        "balance_updated_time_label": format_unix_time(balance_updated_raw) or "—",
        "models": str(channel.get("models") or ""),
        "test_model": str(channel.get("test_model") or ""),
        "activity_label": "N/A",
        "activity_verdict": "na",
    }


def filter_balance_rows(
    rows: list[dict[str, Any]],
    *,
    query_used: bool,
    query_remain: bool,
    used_min: float | None,
    used_max: float | None,
    remain_min: float | None,
    remain_max: float | None,
    start_ts: int | None = None,
    end_ts: int | None = None,
) -> list[dict[str, Any]]:
    matched: list[dict[str, Any]] = []
    for row in rows:
        if not in_created_time_range(row.get("created_time"), start_ts, end_ts):
            continue
        if query_used and not in_range(float(row["used_usd"]), used_min, used_max):
            continue
        if query_remain and not in_range(float(row["remain_usd"]), remain_min, remain_max):
            continue
        matched.append(row)
    return matched


def attach_activity_to_rows(client: Any, rows: list[dict[str, Any]]) -> None:
    """按模型补近一月活跃。接口失败或无模型时写 N/A，不影响余额主流程。"""
    from .channel_activity import (
        ACTIVITY_BATCH_LIMIT,
        apply_na_fill,
        build_activity_fill,
        parse_channel_model_names,
    )

    if not rows:
        return
    pairs: list[dict[str, Any]] = []
    seen: set[tuple[int, str]] = set()
    for row in rows:
        from .channel_activity import channel_row_id

        channel_id = channel_row_id(row)
        for model in parse_channel_model_names(row):
            key = (channel_id, model)
            if channel_id <= 0 or key in seen:
                continue
            seen.add(key)
            pairs.append({"channel_id": channel_id, "model_name": model})
    stats_map: dict[tuple[int, str], dict[str, Any]] = {}
    fetch_ok = False
    if pairs:
        try:
            stats = client.fetch_channel_activity_stats(pairs, batch_size=ACTIVITY_BATCH_LIMIT)
            fetch_ok = True
            for item in stats or []:
                cid = int(item.get("channel_id") or 0)
                model = str(item.get("model_name") or "").strip()
                if cid and model:
                    stats_map[(cid, model)] = item
        except Exception:  # noqa: BLE001
            fetch_ok = False
            stats_map = {}
    for row in rows:
        fill = build_activity_fill(row, stats_map, fetch_ok=fetch_ok) if pairs else apply_na_fill(row)
        if not parse_channel_model_names(row):
            fill = apply_na_fill(row)
        row["activity_label"] = fill["activity_label"]
        row["activity_verdict"] = fill["activity_verdict"]
        row["activity_details"] = fill["details"]
        row["近一月活跃"] = fill["activity_label"]


def run_channel_balance_query(
    base_url: str,
    auth: dict[str, Any],
    query_used: bool = True,
    query_remain: bool = True,
    channel_type: str = "",
    group: str = "",
    keyword: str = "",
    status: int | None = None,
    used_preset: str = "all",
    remain_preset: str = "all",
    used_min: float | None = None,
    used_max: float | None = None,
    remain_min: float | None = None,
    remain_max: float | None = None,
    time_preset: str = "all",
    start_timestamp: int | None = None,
    end_timestamp: int | None = None,
    *,
    include_activity: bool = False,
    sort_by: str = "used_usd",
    sort_desc: bool = True,
) -> ChannelBalanceQueryResult:
    if not query_used and not query_remain:
        raise ValueError("请至少勾选「已用」或「剩余」之一")

    client = NewAPIClient(base_url)
    authenticate_client(client, auth)

    quota_per_unit = DEFAULT_QUOTA_PER_UNIT
    try:
        status_payload = client.get_status()
        data = status_payload.get("data") if isinstance(status_payload, dict) else None
        if isinstance(data, dict) and data.get("quota_per_unit"):
            quota_per_unit = float(data["quota_per_unit"])
        elif isinstance(status_payload, dict) and status_payload.get("quota_per_unit"):
            quota_per_unit = float(status_payload["quota_per_unit"])
    except Exception:  # noqa: BLE001
        pass

    types, _source = get_channel_types(base_url)
    type_id: int | None = None
    type_label = ""
    type_name_by_id = {int(item["value"]): str(item.get("label") or item["value"]) for item in types}

    if channel_type.strip():
        type_id = resolve_channel_type(channel_type.strip(), types)
        type_label = next(
            (str(item.get("label") or "") for item in types if int(item["value"]) == type_id),
            channel_type.strip(),
        )

    channels = client.list_all_channels()
    group_key = (group or "").strip()
    keyword_l = (keyword or "").strip().lower()

    rows: list[dict[str, Any]] = []
    for ch in channels:
        if type_id is not None and int(ch.get("type") or -1) != type_id:
            continue
        if group_key and str(ch.get("group") or "").strip() != group_key:
            continue
        if status is not None and int(ch.get("status") or 0) != int(status):
            continue
        if keyword_l:
            hay = f"{ch.get('id', '')} {ch.get('name', '')} {ch.get('group', '')}".lower()
            if keyword_l not in hay:
                continue
        cid = int(ch.get("type") or 0)
        rows.append(
            channel_balance_row(
                ch,
                quota_per_unit=quota_per_unit,
                type_label=type_name_by_id.get(cid, str(cid)),
            )
        )

    used_min_r, used_max_r = resolve_range(used_preset, USED_RANGE_PRESETS, used_min, used_max)
    remain_min_r, remain_max_r = resolve_range(
        remain_preset, REMAIN_RANGE_PRESETS, remain_min, remain_max
    )
    start_ts, end_ts = resolve_created_time_range(
        time_preset,
        start_timestamp,
        end_timestamp,
    )

    # 仅对勾选的指标应用区间；未勾选的指标不参与筛选
    matched = filter_balance_rows(
        rows,
        query_used=query_used,
        query_remain=query_remain,
        used_min=used_min_r if query_used else None,
        used_max=used_max_r if query_used else None,
        remain_min=remain_min_r if query_remain else None,
        remain_max=remain_max_r if query_remain else None,
        start_ts=start_ts,
        end_ts=end_ts,
    )

    sort_key = (sort_by or "used_usd").strip()
    if sort_key not in {"id", "name", "used_usd", "remain_usd", "balance", "status"}:
        sort_key = "used_usd" if query_used else "remain_usd"
    if sort_key == "balance":
        sort_key = "remain_usd"

    matched.sort(
        key=lambda r: (
            float(r.get(sort_key) or 0) if sort_key in {"used_usd", "remain_usd"} else str(r.get(sort_key) or ""),
            int(r.get("id") or 0),
        ),
        reverse=bool(sort_desc) if sort_key in {"used_usd", "remain_usd", "id", "status"} else False,
    )
    if sort_key in {"name"} and sort_desc:
        matched.reverse()

    if include_activity:
        try:
            attach_activity_to_rows(client, matched)
        except Exception:  # noqa: BLE001
            from .channel_activity import apply_na_fill

            for row in matched:
                fill = apply_na_fill(row)
                row["activity_label"] = fill["activity_label"]
                row["activity_verdict"] = fill["activity_verdict"]

    return ChannelBalanceQueryResult(
        total=len(channels),
        matched_count=len(matched),
        query_used=query_used,
        query_remain=query_remain,
        quota_per_unit=quota_per_unit,
        channel_type=type_label or channel_type,
        type_id=type_id,
        used_preset=used_preset if query_used else "",
        remain_preset=remain_preset if query_remain else "",
        time_preset=time_preset or "all",
        items=matched,
        presets={
            "used": [{"key": k, "label": USED_PRESET_LABELS[k]} for k in USED_RANGE_PRESETS],
            "remain": [{"key": k, "label": REMAIN_PRESET_LABELS[k]} for k in REMAIN_RANGE_PRESETS],
            "time": [
                {"key": k, "label": TIME_PRESET_LABELS[k], "days": TIME_PRESET_DAYS[k]}
                for k in TIME_PRESET_DAYS
            ],
        },
        extra={
            "filtered_before_range": len(rows),
            "used_range": {"min": used_min_r, "max": used_max_r} if query_used else None,
            "remain_range": {"min": remain_min_r, "max": remain_max_r} if query_remain else None,
            "time_range": {
                "preset": time_preset or "all",
                "start_timestamp": start_ts,
                "end_timestamp": end_ts,
                "label": format_time_range_label(start_ts, end_ts, time_preset),
            },
            "group": group_key,
            "keyword": keyword,
            "status_filter": status,
            "sort_by": sort_key,
            "sort_desc": sort_desc,
        },
    )


EXPORT_COLUMNS = [
    ("环境", "env"),
    ("站点", "site_label"),
    ("ID", "id"),
    ("名称", "name"),
    ("分组", "group"),
    ("类型", "type_label"),
    ("状态", "status_label"),
    ("已用(USD)", "used_usd_label"),
    ("剩余(USD)", "remain_usd_label"),
    ("配置时间", "created_time_label"),
    ("余额更新时间", "balance_updated_time_label"),
    ("近一月活跃", "activity_label"),
]


def build_export_table(
    items: list[dict[str, Any]],
    *,
    include_used: bool = True,
    include_remain: bool = True,
    include_activity: bool = False,
) -> tuple[list[str], list[list[Any]]]:
    cols = []
    for header, key in EXPORT_COLUMNS:
        if key == "used_usd_label" and not include_used:
            continue
        if key == "remain_usd_label" and not include_remain:
            continue
        if key == "activity_label" and not include_activity:
            continue
        cols.append((header, key))
    headers = [h for h, _ in cols]
    rows: list[list[Any]] = []
    for item in items:
        rows.append([item.get(key, "") for _, key in cols])
    return headers, rows


def balance_sheet_title(env: str) -> str:
    """Excel 工作表名：test-渠道余额 / uat-渠道余额 / prod-渠道余额。"""
    key = str(env or "other").strip().lower() or "other"
    # 非法字符替换，Excel 表名 ≤ 31
    safe = re.sub(r'[:\\/?*\[\]]', "-", key)[:12]
    return f"{safe}-渠道余额"[:31]


def group_items_by_env(items: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    grouped: dict[str, list[dict[str, Any]]] = {}
    for item in items:
        env = str(item.get("env") or "other").strip().lower() or "other"
        grouped.setdefault(env, []).append(item)
    return grouped


def _ordered_envs(envs: list[str]) -> list[str]:
    from .data_summary import ENVS

    preferred = list(ENVS)
    rest = sorted(e for e in envs if e not in preferred)
    return [e for e in preferred if e in envs] + rest


def export_channel_balance_xlsx(
    items: list[dict[str, Any]],
    *,
    include_used: bool = True,
    include_remain: bool = True,
    include_activity: bool = False,
    sheet_title: str = "渠道余额",
    by_env: bool = True,
) -> tuple[bytes, str]:
    """导出渠道余额为 Excel。

    by_env=True 时按环境分表，表名形如 ``uat-渠道余额``；
    无环境字段时退回单表 ``sheet_title``。
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    except ImportError as exc:
        raise RuntimeError("导出 Excel 需要安装 openpyxl，请执行 pip install openpyxl") from exc

    if not items:
        raise ValueError("没有可导出的渠道数据")

    header_fill = PatternFill("solid", fgColor="2F5496")
    header_font = Font(name="宋体", color="FFFFFF", bold=True, size=11)
    data_font = Font(name="宋体", size=11)
    grid = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )
    widths = {
        "环境": 10,
        "站点": 22,
        "ID": 10,
        "名称": 28,
        "分组": 18,
        "类型": 14,
        "状态": 10,
        "已用(USD)": 12,
        "剩余(USD)": 12,
        "配置时间": 20,
        "余额更新时间": 20,
        "近一月活跃": 36,
    }

    def _write_sheet(ws, sheet_items: list[dict[str, Any]]) -> None:
        headers, rows = build_export_table(
            sheet_items,
            include_used=include_used,
            include_remain=include_remain,
            include_activity=include_activity,
        )
        # 单环境表可去掉「环境」列以更干净；多环境合并表保留
        envs_in_sheet = {str(x.get("env") or "") for x in sheet_items}
        if len(envs_in_sheet) == 1 and "环境" in headers:
            env_idx = headers.index("环境")
            headers = [h for i, h in enumerate(headers) if i != env_idx]
            rows = [[c for i, c in enumerate(row) if i != env_idx] for row in rows]
        ws.append(headers)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.border = grid
            cell.alignment = Alignment(horizontal="center", vertical="center")
        status_idx = headers.index("状态") if "状态" in headers else -1
        activity_idx = headers.index("近一月活跃") if "近一月活跃" in headers else -1
        for row in rows:
            ws.append(row)
        for row_cells in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
            status_val = ""
            if status_idx >= 0:
                status_val = str(row_cells[status_idx].value or "").strip()
            row_font = (
                Font(name="宋体", size=11, color="FFFF0000")
                if status_val == "已禁用"
                else data_font
            )
            for cell in row_cells:
                cell.font = row_font
                cell.border = grid
                wrap = activity_idx >= 0 and cell.column == activity_idx + 1
                cell.alignment = Alignment(vertical="top" if wrap else "center", wrap_text=wrap)
            if activity_idx >= 0:
                text = str(row_cells[activity_idx].value or "")
                lines = max(1, text.count("\n") + 1)
                ws.row_dimensions[row_cells[0].row].height = max(18, 15 * min(lines, 8))
        for idx, header in enumerate(headers, start=1):
            letter = ws.cell(1, idx).column_letter
            ws.column_dimensions[letter].width = widths.get(header, 14)
        ws.freeze_panes = "A2"

    wb = Workbook()
    grouped = group_items_by_env(items) if by_env else {}
    has_env = by_env and any(str(it.get("env") or "").strip() for it in items)

    if has_env and grouped:
        first = True
        for env in _ordered_envs(list(grouped.keys())):
            title = balance_sheet_title(env)
            if first:
                ws = wb.active
                ws.title = title
                first = False
            else:
                ws = wb.create_sheet(title)
            _write_sheet(ws, grouped[env])
    else:
        ws = wb.active
        ws.title = (sheet_title or "渠道余额")[:31]
        _write_sheet(ws, items)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue(), build_balance_filename()


@dataclass
class MultiSiteBalanceQueryResult:
    sites_ok: int = 0
    sites_failed: int = 0
    matched_count: int = 0
    query_used: bool = True
    query_remain: bool = True
    items: list[dict[str, Any]] = field(default_factory=list)
    by_env: dict[str, list[dict[str, Any]]] = field(default_factory=dict)
    site_results: list[dict[str, Any]] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    presets: dict[str, Any] = field(default_factory=dict)
    sheet_names: list[str] = field(default_factory=list)
    time_range: dict[str, Any] = field(default_factory=dict)


def run_multi_site_channel_balance_query(
    *,
    sites: list[Any],
    query_used: bool = True,
    query_remain: bool = True,
    channel_type: str = "",
    group: str = "",
    keyword: str = "",
    status: int | None = None,
    used_preset: str = "all",
    remain_preset: str = "all",
    used_min: float | None = None,
    used_max: float | None = None,
    remain_min: float | None = None,
    remain_max: float | None = None,
    time_preset: str = "all",
    start_timestamp: int | None = None,
    end_timestamp: int | None = None,
    sort_by: str = "used_usd",
    sort_desc: bool = True,
    include_activity: bool = False,
) -> MultiSiteBalanceQueryResult:
    """对多个 NA 站点执行渠道余额查询，结果按环境归类。"""
    from .data_summary import SiteConfig, detect_env, enrich_site_from_env
    from .url_utils import normalize_na_base_url

    if not query_used and not query_remain:
        raise ValueError("请至少勾选「已用」或「剩余」之一")
    if not sites:
        raise ValueError("请至少配置一个站点")

    out = MultiSiteBalanceQueryResult(query_used=query_used, query_remain=query_remain)
    all_items: list[dict[str, Any]] = []

    for site in sites:
        if isinstance(site, SiteConfig):
            cfg = site
        elif isinstance(site, dict):
            base_url = str(site.get("base_url") or "").strip()
            if not base_url:
                continue
            user_id_raw = site.get("user_id")
            try:
                user_id = int(user_id_raw) if user_id_raw not in (None, "") else None
            except (TypeError, ValueError):
                user_id = None
            cfg = SiteConfig(
                base_url=base_url,
                username=str(site.get("username") or "").strip(),
                password=str(site.get("password") or ""),
                env=str(site.get("env") or "").strip().lower(),
                label=str(site.get("label") or "").strip(),
                otp_code=str(site.get("otp_code") or site.get("verification_code") or "").strip(),
                totp_secret=str(site.get("totp_secret") or site.get("otp_secret") or "").strip(),
                access_token=str(site.get("access_token") or "").strip(),
                user_id=user_id,
                verify_ssl=bool(site.get("verify_ssl", True)),
            )
        else:
            continue

        enrich_site_from_env(cfg)
        base_url = normalize_na_base_url(cfg.base_url)
        env = detect_env(base_url, cfg.env)
        label = cfg.label or base_url
        auth: dict[str, Any] = {}
        if str(cfg.access_token or "").strip():
            auth = {
                "access_token": str(cfg.access_token).strip(),
                "user_id": cfg.user_id,
                "otp_code": str(cfg.otp_code or "").strip(),
                "totp_secret": str(cfg.totp_secret or "").strip(),
            }
        elif cfg.username and cfg.password:
            auth = {
                "username": cfg.username,
                "password": cfg.password,
                "otp_code": str(cfg.otp_code or "").strip(),
                "totp_secret": str(cfg.totp_secret or "").strip(),
            }
        else:
            out.sites_failed += 1
            msg = f"[{env}] {label}: 缺少用户名/密码"
            out.errors.append(msg)
            out.site_results.append(
                {"env": env, "label": label, "base_url": base_url, "success": False, "message": msg, "matched": 0}
            )
            continue

        try:
            result = run_channel_balance_query(
                base_url=base_url,
                auth=auth,
                query_used=query_used,
                query_remain=query_remain,
                channel_type=channel_type,
                group=group,
                keyword=keyword,
                status=status,
                used_preset=used_preset,
                remain_preset=remain_preset,
                used_min=used_min,
                used_max=used_max,
                remain_min=remain_min,
                remain_max=remain_max,
                time_preset=time_preset,
                start_timestamp=start_timestamp,
                end_timestamp=end_timestamp,
                sort_by=sort_by,
                sort_desc=sort_desc,
                include_activity=include_activity,
            )
            for item in result.items:
                enriched = dict(item)
                enriched["env"] = env
                enriched["site_label"] = label
                enriched["base_url"] = base_url
                all_items.append(enriched)
            out.sites_ok += 1
            out.site_results.append(
                {
                    "env": env,
                    "label": label,
                    "base_url": base_url,
                    "success": True,
                    "matched": result.matched_count,
                    "total": result.total,
                    "quota_per_unit": result.quota_per_unit,
                }
            )
            if not out.presets:
                out.presets = result.presets
        except Exception as exc:  # noqa: BLE001
            out.sites_failed += 1
            msg = f"[{env}] {label}: {exc}"
            out.errors.append(msg)
            out.site_results.append(
                {"env": env, "label": label, "base_url": base_url, "success": False, "message": str(exc), "matched": 0}
            )

    out.items = all_items
    out.matched_count = len(all_items)
    out.by_env = group_items_by_env(all_items)
    out.sheet_names = [balance_sheet_title(e) for e in _ordered_envs(list(out.by_env.keys()))]
    start_ts, end_ts = resolve_created_time_range(time_preset, start_timestamp, end_timestamp)
    out.time_range = {
        "preset": time_preset or "all",
        "start_timestamp": start_ts,
        "end_timestamp": end_ts,
        "label": format_time_range_label(start_ts, end_ts, time_preset),
    }
    if not out.presets:
        out.presets = {
            "used": [{"key": k, "label": USED_PRESET_LABELS[k]} for k in USED_RANGE_PRESETS],
            "remain": [{"key": k, "label": REMAIN_PRESET_LABELS[k]} for k in REMAIN_RANGE_PRESETS],
            "time": [
                {"key": k, "label": TIME_PRESET_LABELS[k], "days": TIME_PRESET_DAYS[k]}
                for k in TIME_PRESET_DAYS
            ],
        }
    return out
