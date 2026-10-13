# -*- coding: utf-8 -*-
"""excel/ 产物列表：按类型列出，禁止跨目录读取。"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from .channel_balance_query import BALANCE_EXCEL_STEM
from .data_summary import EXCEL_STEM
from .data_summary_exporter import (
    excel_dir,
    excel_recency_key,
    filename_timestamp,
    is_overview_excel,
    is_staging_excel,
    list_summary_excel_files,
)
from .overview_analyzer import ANALYSIS_STEM, list_analysis_excel_files
from .overview_exporter import OVERVIEW_STEM, list_overview_excel_files
from .upstream_exporter import (
    UPSTREAM_AZURE_STEM,
    UPSTREAM_COMBINED_STEM,
    UPSTREAM_EXCEL_STEM,
)
from .volc_billing import list_bill_excel_files
from .azure_billing import list_azure_bill_excel_files
from .bill_combined import list_combined_bill_excel_files

KINDS = ("summary", "upstream", "overview", "analysis", "balance", "logs", "bill", "bill_azure", "bill_combined")
_EMPTY_SECRET_VALUES = frozenset({"", "n/a", "na", "-", "—", "无", "空", "none", "null"})
_SECRET_HEADER_KEYS = frozenset(
    {
        "密钥",
        "密钥一",
        "密钥二",
        "apikey",
        "apikey汇总",
        "rawapikey",
    }
)


def excel_root(root_dir: Path) -> Path:
    return excel_dir(root_dir).resolve()


def resolve_excel_path(root_dir: Path, raw: str | Path | None) -> Path:
    """只允许 excel/ 下的已有 .xlsx。id 可为相对路径或文件名。"""
    text = str(raw or "").strip().replace("\\", "/")
    if not text:
        raise ValueError("请选择 Excel 文件")
    if ".." in text or text.startswith("/") or text.startswith("\\"):
        raise ValueError("Excel 文件名无效")
    name = Path(text).name
    if not name.lower().endswith(".xlsx"):
        raise ValueError("仅支持 .xlsx")
    if name.startswith(".") or "/" in name or "\\" in name:
        raise ValueError("Excel 文件名无效")
    dest = (excel_root(root_dir) / name).resolve()
    try:
        dest.relative_to(excel_root(root_dir))
    except ValueError as exc:
        raise ValueError("文件必须位于 excel/ 目录") from exc
    if not dest.is_file():
        raise FileNotFoundError(f"未找到文件：{name}")
    return dest


def _rel(root_dir: Path, path: Path) -> str:
    return str(path.resolve().relative_to(Path(root_dir).resolve())).replace("\\", "/")


def _mtime_iso(path: Path) -> str:
    try:
        return datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds")
    except OSError:
        return ""


def _norm_header(label: Any) -> str:
    return str(label or "").strip().lower().replace(" ", "").replace("　", "")


def _is_secret_header(label: Any) -> bool:
    text = str(label or "").strip()
    if not text or "摘要" in text:
        return False
    return _norm_header(text) in _SECRET_HEADER_KEYS


def _is_filled_secret(value: Any) -> bool:
    text = str(value or "").strip()
    if not text:
        return False
    return text.lower() not in _EMPTY_SECRET_VALUES


def workbook_has_filled_secrets(path: Path | None) -> bool:
    """有密钥列但单元格全空不算含密钥；密钥摘要列也不算。"""
    src = Path(path) if path else None
    if not src or not src.is_file() or src.suffix.lower() != ".xlsx":
        return False
    from openpyxl import load_workbook

    try:
        wb = load_workbook(src, read_only=True, data_only=True)
    except Exception:
        return False
    try:
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            header = next(rows, None)
            if not header:
                continue
            secret_idxs = [i for i, cell in enumerate(header) if _is_secret_header(cell)]
            if not secret_idxs:
                continue
            for row in rows:
                if not row:
                    continue
                for idx in secret_idxs:
                    if idx < len(row) and _is_filled_secret(row[idx]):
                        return True
    finally:
        wb.close()
    return False


def _artifact(root_dir: Path, path: Path, kind: str, extra: dict[str, Any] | None = None) -> dict[str, Any]:
    # 智能分析表只有密钥摘要，不当作明文密钥。其余看单元格是否真有值，而不是「有密钥列」。
    has_secrets = False if kind == "analysis" else workbook_has_filled_secrets(path)
    item = {
        "id": path.name,
        "kind": kind,
        "filename": path.name,
        "path": _rel(root_dir, path),
        "created_at": filename_timestamp(path) or _mtime_iso(path),
        "mtime": _mtime_iso(path),
        "has_secrets": has_secrets,
        "recency": excel_recency_key(path),
    }
    if extra:
        item.update(extra)
    return item


def list_upstream_files(root_dir: Path) -> list[Path]:
    directory = excel_dir(root_dir)
    if not directory.exists():
        return []
    found: dict[str, Path] = {}
    for stem in (UPSTREAM_COMBINED_STEM, UPSTREAM_EXCEL_STEM, UPSTREAM_AZURE_STEM):
        for path in directory.glob(f"{stem}_*.xlsx"):
            found[str(path.resolve())] = path
    from .data_summary_exporter import sort_excel_files_latest_first

    return sort_excel_files_latest_first(list(found.values()))


def list_balance_files(root_dir: Path) -> list[Path]:
    directory = excel_dir(root_dir)
    if not directory.exists():
        return []
    from .data_summary_exporter import sort_excel_files_latest_first

    return sort_excel_files_latest_first(list(directory.glob(f"{BALANCE_EXCEL_STEM}_*.xlsx")))


def list_log_files(root_dir: Path) -> list[Path]:
    directory = excel_dir(root_dir)
    if not directory.exists():
        return []
    from .data_summary_exporter import sort_excel_files_latest_first

    files = [p for p in directory.glob("usage_logs_*.xlsx") if not is_staging_excel(p)]
    files += list(directory.glob("usage_logs_*.csv"))
    return sort_excel_files_latest_first(files)


def _paths_for_kind(root_dir: Path, kind_key: str) -> list[Path]:
    """返回某一种产物的全部候选路径，不读取工作簿内容。"""
    if kind_key == "summary":
        return list_summary_excel_files(root_dir)
    if kind_key == "upstream":
        return list_upstream_files(root_dir)
    if kind_key == "overview":
        return list_overview_excel_files(root_dir)
    if kind_key == "analysis":
        return list_analysis_excel_files(root_dir)
    if kind_key == "balance":
        return list_balance_files(root_dir)
    if kind_key == "bill":
        return list_bill_excel_files(root_dir)
    if kind_key == "bill_azure":
        return list_azure_bill_excel_files(root_dir)
    if kind_key == "bill_combined":
        return list_combined_bill_excel_files(root_dir)
    return list_log_files(root_dir)


def list_artifacts(root_dir: Path, kind: str, *, limit: int = 20) -> list[dict[str, Any]]:
    kind_key = str(kind or "").strip().lower()
    if kind_key not in KINDS:
        raise ValueError(f"kind 须为 {', '.join(KINDS)}")
    cap = max(1, min(int(limit or 20), 50))
    paths = _paths_for_kind(root_dir, kind_key)
    return [_artifact(root_dir, path, kind_key) for path in paths[:cap]]


def artifact_kind_for_filename(root_dir: Path, raw: str | Path | None) -> str | None:
    """按已有产物清单确定文件真实类型，供通用下载接口做二次授权校验。"""
    name = Path(str(raw or "")).name
    if not name:
        return None
    for kind in KINDS:
        if any(path.name == name for path in _paths_for_kind(root_dir, kind)):
            return kind
    return None


def latest_artifacts_bundle(root_dir: Path) -> list[dict[str, Any]]:
    bundle: list[dict[str, Any]] = []
    for kind in ("summary", "upstream", "overview", "analysis", "balance", "bill", "bill_azure", "bill_combined"):
        items = list_artifacts(root_dir, kind, limit=1)
        if items:
            bundle.append(items[0])
    return bundle


def is_gateway_summary_name(name: str) -> bool:
    path = Path(name)
    return path.name.startswith(EXCEL_STEM) and not is_overview_excel(path) and path.suffix.lower() == ".xlsx"


def is_overview_name(name: str) -> bool:
    return is_overview_excel(name) or Path(name).name.startswith(OVERVIEW_STEM)


def is_analysis_name(name: str) -> bool:
    return Path(name).name.startswith(f"{ANALYSIS_STEM}_")
