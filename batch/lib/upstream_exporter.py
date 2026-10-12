# -*- coding: utf-8 -*-
"""上游信息 Excel 导出（独立于 AI 网关汇总表）。"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .data_summary_exporter import sort_excel_files_latest_first
from .deploy_time import sort_rows_by_deploy_time_desc
from .upstream_azure import AZURE_HEADERS
from .upstream_volc import UPSTREAM_HEADERS

UPSTREAM_EXCEL_STEM = "上游信息_火山云"
SHEET_VOLC = "火山云"
UPSTREAM_AZURE_STEM = "上游信息_Azure"
SHEET_AZURE = "Azure OpenAI"
SHEET_AZURE_LEGACY = "openai"
UPSTREAM_COMBINED_STEM = "上游信息_汇总"
TIMESTAMP_FORMAT = "%Y%m%d%H%M%S"

_GRID_SIDE = Side(style="thin", color="FF000000")
_HEADER_FILL = PatternFill("solid", fgColor="FF2F5496")
_HEADER_FONT = Font(name="宋体", bold=True, color="FFFFFFFF", size=11)
_DATA_FONT = Font(name="宋体", size=11)


def excel_dir(root_dir: Path) -> Path:
    return root_dir / "excel"


def build_upstream_filename(when: datetime | None = None) -> str:
    stamp = (when or datetime.now()).strftime(TIMESTAMP_FORMAT)
    return f"{UPSTREAM_EXCEL_STEM}_{stamp}.xlsx"


def build_upstream_path(root_dir: Path, when: datetime | None = None) -> Path:
    return excel_dir(root_dir) / build_upstream_filename(when)


def list_upstream_excel_files(root_dir: Path) -> list[Path]:
    directory = excel_dir(root_dir)
    if not directory.exists():
        return []
    files = list(directory.glob(f"{UPSTREAM_EXCEL_STEM}_*.xlsx"))
    return sort_excel_files_latest_first(files)


def latest_upstream_excel_path(root_dir: Path) -> Path | None:
    files = list_upstream_excel_files(root_dir)
    return files[0] if files else None


def build_azure_upstream_filename(when: datetime | None = None) -> str:
    stamp = (when or datetime.now()).strftime(TIMESTAMP_FORMAT)
    return f"{UPSTREAM_AZURE_STEM}_{stamp}.xlsx"


def build_azure_upstream_path(root_dir: Path, when: datetime | None = None) -> Path:
    return excel_dir(root_dir) / build_azure_upstream_filename(when)


def list_azure_upstream_excel_files(root_dir: Path) -> list[Path]:
    directory = excel_dir(root_dir)
    if not directory.exists():
        return []
    files = list(directory.glob(f"{UPSTREAM_AZURE_STEM}_*.xlsx"))
    return sort_excel_files_latest_first(files)


def latest_azure_upstream_excel_path(root_dir: Path) -> Path | None:
    files = list_azure_upstream_excel_files(root_dir)
    return files[0] if files else None


def build_combined_upstream_filename(when: datetime | None = None) -> str:
    stamp = (when or datetime.now()).strftime(TIMESTAMP_FORMAT)
    return f"{UPSTREAM_COMBINED_STEM}_{stamp}.xlsx"


def build_combined_upstream_path(root_dir: Path, when: datetime | None = None) -> Path:
    return excel_dir(root_dir) / build_combined_upstream_filename(when)


def list_combined_upstream_excel_files(root_dir: Path) -> list[Path]:
    directory = excel_dir(root_dir)
    if not directory.exists():
        return []
    files = list(directory.glob(f"{UPSTREAM_COMBINED_STEM}_*.xlsx"))
    return sort_excel_files_latest_first(files)


def latest_combined_upstream_excel_path(root_dir: Path) -> Path | None:
    files = list_combined_upstream_excel_files(root_dir)
    return files[0] if files else None


def _apply_style(ws, headers: list[str]) -> None:
    max_col = len(headers)
    max_row = ws.max_row or 0
    if max_row < 1 or max_col < 1:
        return
    border = Border(left=_GRID_SIDE, right=_GRID_SIDE, top=_GRID_SIDE, bottom=_GRID_SIDE)
    for row_idx in range(1, max_row + 1):
        for col_idx in range(1, max_col + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=col_idx > 1)
            if row_idx == 1:
                cell.fill = _HEADER_FILL
                cell.font = _HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cell.font = _DATA_FONT
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(max_col)}{max_row}"
    for col_idx, header in enumerate(headers, start=1):
        max_len = len(str(header))
        for row_idx in range(2, min(max_row + 1, 200)):
            value = ws.cell(row=row_idx, column=col_idx).value
            if value is not None:
                first_line = str(value).splitlines()[0] if str(value) else ""
                cap = 80 if header == "API Key汇总" else 48
                max_len = max(max_len, min(len(first_line), cap))
        width = max(max_len + 2, 10)
        if header == "API Key汇总":
            width = max(width, 36)
        if header == "近一月活跃":
            width = max(width, 28)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(width, 88)


def _write_sheet(ws, headers: list[str], rows: list[dict[str, Any]]) -> None:
    cols = list(headers)
    fill_azure_tags = "项目" in cols and "负责人" in cols
    for col_idx, header in enumerate(cols, start=1):
        ws.cell(row=1, column=col_idx, value=header)
    for row_idx, row in enumerate(rows or [], start=2):
        for col_idx, header in enumerate(cols, start=1):
            value = row.get(header, "")
            if fill_azure_tags and header in ("项目", "负责人") and str(value or "").strip() == "":
                value = "N/A"
            ws.cell(row=row_idx, column=col_idx, value="" if value is None else value)
    _apply_style(ws, cols)


def save_upstream_workbook(
    path: Path,
    rows: list[dict[str, Any]],
    *,
    headers: list[str] | None = None,
    sheet_name: str = SHEET_VOLC,
) -> Path:
    """写入独立上游工作簿；与 ai网关令牌及渠道列表_* 互不覆盖。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    cols = list(headers or UPSTREAM_HEADERS)
    ordered = sort_rows_by_deploy_time_desc(rows) if "部署时间" in cols else list(rows or [])
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    _write_sheet(ws, cols, ordered)
    wb.save(path)
    wb.close()
    return path


def save_combined_upstream_workbook(
    path: Path,
    volc_rows: list[dict[str, Any]] | None = None,
    azure_rows: list[dict[str, Any]] | None = None,
) -> Path:
    """同一工作簿写入「火山云」「Azure OpenAI」两张工作表。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws_volc = wb.active
    ws_volc.title = SHEET_VOLC
    _write_sheet(ws_volc, list(UPSTREAM_HEADERS), sort_rows_by_deploy_time_desc(volc_rows))
    ws_azure = wb.create_sheet(SHEET_AZURE)
    _write_sheet(ws_azure, list(AZURE_HEADERS), sort_rows_by_deploy_time_desc(azure_rows))
    wb.save(path)
    wb.close()
    return path


def upstream_workbook_stats(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {"rows": 0, "sheet": SHEET_VOLC, "sheets": {}, "total_rows": 0}
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets: dict[str, int] = {}
        for name in wb.sheetnames:
            ws = wb[name]
            sheets[name] = max(0, int(ws.max_row or 1) - 1)
        has_volc = SHEET_VOLC in sheets
        has_azure = SHEET_AZURE in sheets or SHEET_AZURE_LEGACY in sheets
        if has_volc and has_azure:
            title = f"{SHEET_VOLC} + {SHEET_AZURE}"
            rows = sum(sheets.values())
        elif has_volc:
            title = SHEET_VOLC
            rows = sheets.get(SHEET_VOLC, 0)
        elif SHEET_AZURE in sheets:
            title = SHEET_AZURE
            rows = sheets.get(SHEET_AZURE, 0)
        elif SHEET_AZURE_LEGACY in sheets:
            title = SHEET_AZURE_LEGACY
            rows = sheets.get(SHEET_AZURE_LEGACY, 0)
        elif wb.sheetnames:
            title = wb.sheetnames[0]
            rows = sheets.get(title, 0)
        else:
            return {"rows": 0, "sheet": "", "sheets": sheets, "total_rows": 0}
        return {
            "rows": int(rows or 0),
            "sheet": title,
            "sheets": sheets,
            "total_rows": sum(sheets.values()),
        }
    finally:
        wb.close()
