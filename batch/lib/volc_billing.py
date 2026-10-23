# -*- coding: utf-8 -*-
"""上游账单 · 火山云费用中心（按天 / 按月汇总导出）。"""

from __future__ import annotations

import calendar
import re
from dataclasses import dataclass, field
from datetime import datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from .volc_openapi import BILLING_HOST, VolcOpenAPIClient, VolcOpenAPIError
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_CN_TZ = timezone(timedelta(hours=8))
_GRID_SIDE = Side(style="thin", color="FF000000")
_HEADER_FILL = PatternFill("solid", fgColor="FF2F5496")
_HEADER_FONT = Font(name="宋体", bold=True, color="FFFFFFFF", size=11)
_DATA_FONT = Font(name="宋体", size=11)
_SUM_FONT = Font(name="宋体", bold=True, size=11)
_SUM_FILL = PatternFill("solid", fgColor="FFE7E6E6")
COST_FIELD = "费用"
COST_COL_HEADER = "费用（¥）"
SAVING_PLAN_FIELD = "节省计划抵扣总额"
SAVING_PLAN_COL_HEADER = "节省计划抵扣总额（¥）"
COST_NUMBER_FORMAT = '"¥"#,##0.00'
SUM_LABEL = "合计"
BILLING_SERVICE = "billing"
BILLING_VERSION = "2022-01-01"
EXCEL_STEM = "上游账单_火山云"
TIMESTAMP_FORMAT = "%Y%m%d%H%M%S"
SHEET_DAY = "按天"
SHEET_MONTH = "按月"
SHEET_NOTE = "说明"
DAY_HEADERS = ["日期", "项目名", "产品", "配置名称", "分拆项ID", "费用", "节省计划抵扣总额"]
MONTH_HEADERS = ["日期", "项目名", "产品", "配置名称", "分拆项ID", "费用", "节省计划抵扣总额"]
_DEFAULT_PROJECT_LABELS = frozenset(
    {
        "default",
        "默认项目",
        "默认",
        "default(默认项目)",
        "default（默认项目）",
    }
)
_PERIOD_RE = re.compile(r"^(\d{4})[-/]?(\d{1,2})$")
_DATE_RE = re.compile(r"^(\d{4})[-/]?(\d{1,2})[-/]?(\d{1,2})$")
GROUP_PERIOD_MONTH = 0
GROUP_PERIOD_DAY = 1
GROUP_TERM_ITEM = 0
PAGE_LIMIT = 300
MAX_BILL_MONTHS = 24

PERMISSION_HINT = (
    "当前 IAM Access Key 调不通费用中心。"
    "密钥仍用 env/sites.json 的 upstream.volc（与「上游信息」相同的 AKLT…）。"
    "请在 https://console.volcengine.com/iam 给该用户附加账单查询权限"
    "（至少 billing:ListBillDetail；要「分拆项ID」还需开通分账账单并授权 billing:ListSplitBillDetail）。"
    "控制台对照：https://console.volcengine.com/finance/bill/detail"
)


@dataclass
class BillCollectStats:
    bill_period: str = ""
    source_action: str = ""
    day_rows: int = 0
    month_rows: int = 0
    raw_count: int = 0
    skipped_empty_or_default: int = 0
    split_used: bool = False
    split_error: str = ""
    errors: list[str] = field(default_factory=list)
    rows_day: list[dict[str, Any]] = field(default_factory=list)
    rows_month: list[dict[str, Any]] = field(default_factory=list)


def default_bill_period(*, now: datetime | None = None) -> str:
    clock = now or datetime.now(_CN_TZ)
    if clock.tzinfo is None:
        clock = clock.replace(tzinfo=_CN_TZ)
    clock = clock.astimezone(_CN_TZ)
    year, month = clock.year, clock.month - 1
    if month < 1:
        year -= 1
        month = 12
    return f"{year:04d}-{month:02d}"


def parse_bill_period(raw: Any) -> str:
    text = str(raw or "").strip()
    if not text:
        return default_bill_period()
    found = _PERIOD_RE.match(text.replace("年", "-").replace("月", ""))
    if not found:
        raise ValueError("账期须为 YYYY-MM，例如 2026-08")
    year, month = int(found.group(1)), int(found.group(2))
    if month < 1 or month > 12:
        raise ValueError("账期月份须在 1–12")
    return f"{year:04d}-{month:02d}"


def bill_month_label(bill_period: Any = None) -> str:
    """账期 YYYY-MM →「9月」（不补零）。解析失败时用默认账期。"""
    try:
        period = parse_bill_period(bill_period)
    except ValueError:
        period = default_bill_period()
    month = int(period.split("-")[1])
    return f"{month}月"


def expand_bill_period_range(start: Any, end: Any) -> list[str]:
    """闭区间展开为 YYYY-MM 列表，起止颠倒会自动对调。一次最多 24 个月。"""
    first = parse_bill_period(start)
    last = parse_bill_period(end)
    if first > last:
        first, last = last, first
    year, month = [int(part) for part in first.split("-")]
    end_year, end_month = [int(part) for part in last.split("-")]
    out: list[str] = []
    while (year, month) <= (end_year, end_month):
        out.append(f"{year:04d}-{month:02d}")
        if len(out) > MAX_BILL_MONTHS:
            raise ValueError(f"一次最多导出 {MAX_BILL_MONTHS} 个月，请缩小账期范围")
        month += 1
        if month > 12:
            month = 1
            year += 1
    return out


def bill_periods_from_payload(body: dict[str, Any] | None) -> list[str]:
    """从请求体解析账期列表：bill_periods、起止范围，或单个 bill_period。"""
    pack = body or {}
    raw_list = pack.get("bill_periods") if pack.get("bill_periods") not in (None, "") else pack.get("periods")
    if isinstance(raw_list, str):
        raw_list = [part.strip() for part in raw_list.replace("，", ",").split(",") if part.strip()]
    if isinstance(raw_list, list) and raw_list:
        seen: list[str] = []
        for item in raw_list:
            period = parse_bill_period(item)
            if period not in seen:
                seen.append(period)
        if len(seen) > MAX_BILL_MONTHS:
            raise ValueError(f"一次最多导出 {MAX_BILL_MONTHS} 个月，请缩小账期范围")
        return seen
    start = pack.get("bill_period_from") or pack.get("period_from") or pack.get("start")
    end = pack.get("bill_period_to") or pack.get("period_to") or pack.get("end")
    if str(start or "").strip() or str(end or "").strip():
        return expand_bill_period_range(start or end, end or start)
    return [parse_bill_period(pack.get("bill_period") or pack.get("period"))]


def excel_dir(root_dir: Path) -> Path:
    return Path(root_dir) / "excel"


def build_bill_filename(when: datetime | None = None, *, bill_period: Any = None) -> str:
    stamp = (when or datetime.now()).strftime(TIMESTAMP_FORMAT)
    return f"{EXCEL_STEM}_{bill_month_label(bill_period)}_{stamp}.xlsx"


def build_bill_path(root_dir: Path, when: datetime | None = None, *, bill_period: Any = None) -> Path:
    return excel_dir(root_dir) / build_bill_filename(when, bill_period=bill_period)


def list_bill_excel_files(root_dir: Path) -> list[Path]:
    from .data_summary_exporter import sort_excel_files_latest_first

    directory = excel_dir(root_dir)
    if not directory.exists():
        return []
    return sort_excel_files_latest_first(list(directory.glob(f"{EXCEL_STEM}_*.xlsx")))


def latest_bill_excel_path(root_dir: Path) -> Path | None:
    files = list_bill_excel_files(root_dir)
    return files[0] if files else None


def _parse_money(value: Any) -> Decimal:
    text = (
        str(value or "")
        .strip()
        .replace(",", "")
        .replace("¥", "")
        .replace("￥", "")
        .replace(" ", "")
    )
    if not text:
        return Decimal("0")
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        return Decimal("0")


def _format_money(value: Decimal) -> float:
    quantized = value.quantize(Decimal("0.000001"))
    return float(quantized)


def is_blank_or_default_project(*names: Any) -> bool:
    """空项目名、火山「默认项目」均排除。"""
    labels = [str(name or "").strip() for name in names]
    if not any(labels):
        return True
    for text in labels:
        if not text:
            continue
        folded = text.lower().replace(" ", "")
        if folded in {item.lower() for item in _DEFAULT_PROJECT_LABELS}:
            return True
        if folded.startswith("default(") or folded.startswith("default（"):
            return True
    return False


def filter_named_projects(rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], int]:
    kept: list[dict[str, Any]] = []
    skipped = 0
    for row in rows:
        if not row:
            continue
        if is_blank_or_default_project(row.get("项目名")):
            skipped += 1
            continue
        kept.append(row)
    return kept, skipped


def _pick_text(item: dict[str, Any], *keys: str) -> str:
    for key in keys:
        text = str(item.get(key) or "").strip()
        if text:
            return text
    return ""


def _expense_date(item: dict[str, Any], bill_period: str) -> str:
    raw = _pick_text(item, "ExpenseDate", "ExpenseTime", "BillPeriod")
    if raw:
        head = raw.replace(".", "-").replace("/", "-").split(" ")[0]
        found = _DATE_RE.match(head.replace("-", "")) if "-" not in head and len(head) >= 8 else _DATE_RE.match(head)
        if not found and len(head) >= 8 and head[:8].isdigit():
            found = _DATE_RE.match(head[:8])
        if found:
            year, month, day = int(found.group(1)), int(found.group(2)), int(found.group(3))
            last = calendar.monthrange(year, month)[1]
            day = min(max(day, 1), last)
            return f"{year:04d}-{month:02d}-{day:02d}"
        month_hit = _PERIOD_RE.match(head)
        if month_hit:
            return f"{int(month_hit.group(1)):04d}-{int(month_hit.group(2)):02d}"
    return bill_period


def _payable(item: dict[str, Any]) -> Decimal:
    for key in (
        "PayableAmount",
        "SettlePayableAmount",
        "PreTaxPayableAmount",
        "OriginalBillAmount",
        "PretaxAmount",
    ):
        if item.get(key) not in (None, ""):
            return _parse_money(item.get(key))
    return Decimal("0")


def _saving_plan_deduction(item: dict[str, Any]) -> Decimal:
    raw = item.get("SavingPlanDeductionDiscountTotalAmount")
    if raw in (None, ""):
        return Decimal("0")
    return _parse_money(raw)


def map_bill_item(item: dict[str, Any], *, bill_period: str) -> dict[str, Any]:
    if not isinstance(item, dict):
        return {}
    date = _expense_date(item, bill_period)
    return {
        "日期": date,
        "项目名": _pick_text(item, "ProjectDisplayName", "Project"),
        "产品": _pick_text(item, "ProductZh", "Product"),
        "配置名称": _pick_text(item, "ConfigName"),
        "分拆项ID": _pick_text(item, "SplitItemID", "SplitItemId", "SplitItemName"),
        "费用": _format_money(_payable(item)),
        SAVING_PLAN_FIELD: _format_money(_saving_plan_deduction(item)),
    }


def _row_key(row: dict[str, Any], *, month: bool = False) -> tuple[str, str, str, str, str]:
    date = str(row.get("日期") or "")
    if month:
        date = date[:7]
    return (
        date,
        str(row.get("项目名") or ""),
        str(row.get("产品") or ""),
        str(row.get("配置名称") or ""),
        str(row.get("分拆项ID") or ""),
    )


def aggregate_rows(rows: list[dict[str, Any]], *, month: bool = False) -> list[dict[str, Any]]:
    buckets: dict[tuple[str, str, str, str, str], tuple[Decimal, Decimal]] = {}
    for row in rows:
        if not row:
            continue
        key = _row_key(row, month=month)
        cost, saving_plan = buckets.get(key, (Decimal("0"), Decimal("0")))
        buckets[key] = (
            cost + _parse_money(row.get(COST_FIELD)),
            saving_plan + _parse_money(row.get(SAVING_PLAN_FIELD)),
        )
    out = [
        {
            "日期": key[0],
            "项目名": key[1],
            "产品": key[2],
            "配置名称": key[3],
            "分拆项ID": key[4],
            "费用": _format_money(amount[0]),
            SAVING_PLAN_FIELD: _format_money(amount[1]),
        }
        for key, amount in buckets.items()
    ]
    out.sort(key=lambda r: (str(r["日期"]), str(r["项目名"]), str(r["产品"]), str(r["配置名称"]), str(r["分拆项ID"])))
    return out


def _list_result_batch(result: dict[str, Any]) -> list[dict[str, Any]]:
    batch = result.get("List") or result.get("list") or []
    if not isinstance(batch, list):
        return []
    return [item for item in batch if isinstance(item, dict)]


def list_billing_pages(
    client: VolcOpenAPIClient,
    *,
    action: str,
    bill_period: str,
    group_period: int,
) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    offset = 0
    while True:
        result = client.request(
            method="POST",
            host=BILLING_HOST,
            service=BILLING_SERVICE,
            action=action,
            version=BILLING_VERSION,
            body={
                "BillPeriod": bill_period,
                "Limit": PAGE_LIMIT,
                "Offset": offset,
                "NeedRecordNum": 1,
                "IgnoreZero": 0,
                "GroupPeriod": int(group_period),
                "GroupTerm": GROUP_TERM_ITEM,
            },
        )
        batch = _list_result_batch(result)
        items.extend(batch)
        total = int(result.get("Total") or 0)
        offset += len(batch)
        if not batch or (total > 0 and offset >= total) or len(batch) < PAGE_LIMIT:
            break
        if offset > 200_000:
            break
    return items


def _is_permission_error(exc: VolcOpenAPIError) -> bool:
    code = str(exc.code or "").lower()
    text = str(exc).lower()
    tokens = ("accessdenied", "forbidden", "unauthorized", "nopermission", "permissiondenied", "notauthorized")
    return any(token in code or token in text for token in tokens)


def fetch_bill_items(client: VolcOpenAPIClient, bill_period: str) -> tuple[list[dict[str, Any]], str, str]:
    """优先分账明细（含分拆项ID），失败则退回账单明细。返回 (items, action, split_error)。"""
    split_error = ""
    try:
        items = list_billing_pages(
            client,
            action="ListSplitBillDetail",
            bill_period=bill_period,
            group_period=GROUP_PERIOD_DAY,
        )
        return items, "ListSplitBillDetail", ""
    except VolcOpenAPIError as exc:
        split_error = str(exc)
        try:
            items = list_billing_pages(
                client,
                action="ListBillDetail",
                bill_period=bill_period,
                group_period=GROUP_PERIOD_DAY,
            )
            return items, "ListBillDetail", split_error
        except VolcOpenAPIError as inner:
            if _is_permission_error(exc) or _is_permission_error(inner):
                raise VolcOpenAPIError(
                    f"{PERMISSION_HINT} 原始错误：{inner}",
                    code=inner.code or exc.code,
                    request_id=inner.request_id or exc.request_id,
                    raw=inner.raw,
                ) from inner
            raise


def collect_volc_bills(
    access_key_id: str,
    secret_access_key: str,
    *,
    region: str = "cn-beijing",
    bill_period: str | None = None,
    client: VolcOpenAPIClient | None = None,
) -> BillCollectStats:
    period = parse_bill_period(bill_period)
    api = client or VolcOpenAPIClient(access_key_id, secret_access_key, region=region)
    raw, action, split_error = fetch_bill_items(api, period)
    named: list[dict[str, Any]] = []
    skipped = 0
    for item in raw:
        row = map_bill_item(item, bill_period=period)
        if not row:
            skipped += 1
            continue
        if is_blank_or_default_project(
            row.get("项目名"),
            item.get("Project") if isinstance(item, dict) else "",
            item.get("ProjectDisplayName") if isinstance(item, dict) else "",
        ):
            skipped += 1
            continue
        named.append(row)
    day_rows = aggregate_rows(named, month=False)
    month_rows = aggregate_rows(named, month=True)
    stats = BillCollectStats(
        bill_period=period,
        source_action=action,
        day_rows=len(day_rows),
        month_rows=len(month_rows),
        raw_count=len(raw),
        skipped_empty_or_default=skipped,
        split_used=action == "ListSplitBillDetail",
        split_error=split_error,
        rows_day=day_rows,
        rows_month=month_rows,
    )
    if split_error and not stats.split_used:
        stats.errors.append(
            "分账接口不可用，已改用账单明细；「分拆项ID」可能为空。"
            f"（{split_error}）"
        )
    return stats


def excel_cost_headers(headers: list[str]) -> list[str]:
    return [
        COST_COL_HEADER if item == COST_FIELD
        else SAVING_PLAN_COL_HEADER if item == SAVING_PLAN_FIELD
        else item
        for item in headers
    ]


def sheet_data_rows(ws) -> int:
    max_row = int(ws.max_row or 1)
    if max_row <= 1:
        return 0
    last = str(ws.cell(row=max_row, column=1).value or "").strip()
    if last == SUM_LABEL:
        return max(0, max_row - 2)
    return max(0, max_row - 1)


def write_bill_sheet(ws, headers: list[str], rows: list[dict[str, Any]], *, add_sum: bool = True) -> None:
    cols = list(headers)
    display = excel_cost_headers(cols)
    cost_idx = next((i for i, name in enumerate(cols, start=1) if name == COST_FIELD), None)
    saving_plan_idx = next((i for i, name in enumerate(cols, start=1) if name == SAVING_PLAN_FIELD), None)
    for col_idx, header in enumerate(display, start=1):
        ws.cell(row=1, column=col_idx, value=header)
    data_rows = list(rows or [])
    for row_idx, row in enumerate(data_rows, start=2):
        for col_idx, header in enumerate(cols, start=1):
            value = row.get(header, "")
            cell = ws.cell(row=row_idx, column=col_idx, value="" if value is None else value)
            if col_idx in (cost_idx, saving_plan_idx) and value not in ("", None):
                cell.number_format = COST_NUMBER_FORMAT
    last_data = 1 + len(data_rows)
    if add_sum and cost_idx:
        sum_row = last_data + 1
        letter = get_column_letter(cost_idx)
        ws.cell(row=sum_row, column=1, value=SUM_LABEL)
        for money_idx, money_letter in (
            (cost_idx, letter),
            (saving_plan_idx, get_column_letter(saving_plan_idx) if saving_plan_idx else ""),
        ):
            if not money_idx:
                continue
            total = ws.cell(row=sum_row, column=money_idx)
            if data_rows:
                total.value = f"=SUM({money_letter}2:{money_letter}{last_data})"
            else:
                total.value = 0
            total.number_format = COST_NUMBER_FORMAT
    border = Border(left=_GRID_SIDE, right=_GRID_SIDE, top=_GRID_SIDE, bottom=_GRID_SIDE)
    max_row = ws.max_row or 0
    max_col = len(cols)
    for r_idx in range(1, max_row + 1):
        for c_idx in range(1, max_col + 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.border = border
            if r_idx == 1:
                cell.fill = _HEADER_FILL
                cell.font = _HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            elif add_sum and cost_idx and r_idx == max_row:
                cell.font = _SUM_FONT
                cell.fill = _SUM_FILL
                cell.alignment = Alignment(vertical="center")
            else:
                cell.font = _DATA_FONT
                cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    last_col = get_column_letter(max_col)
    filter_end = last_data if add_sum and cost_idx else max_row
    ws.auto_filter.ref = f"A1:{last_col}{max(1, filter_end)}"
    for col_idx, header in enumerate(display, start=1):
        width = max(len(str(header)) + 2, 12)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(width, 48)


def _write_sheet(ws, headers: list[str], rows: list[dict[str, Any]], *, add_sum: bool | None = None) -> None:
    has_cost = COST_FIELD in headers
    write_bill_sheet(ws, headers, rows, add_sum=has_cost if add_sum is None else add_sum)


def save_bill_workbook(
    path: Path,
    stats: BillCollectStats,
) -> Path:
    from openpyxl import Workbook

    dest = Path(path)
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws_day = wb.active
    ws_day.title = SHEET_DAY
    _write_sheet(ws_day, DAY_HEADERS, stats.rows_day)
    ws_month = wb.create_sheet(SHEET_MONTH)
    _write_sheet(ws_month, MONTH_HEADERS, stats.rows_month)
    note = wb.create_sheet(SHEET_NOTE)
    _write_sheet(
        note,
        ["项", "说明"],
        [
            {"项": "账期", "说明": stats.bill_period},
            {"项": "接口", "说明": stats.source_action or ""},
            {
                "项": "按天",
                "说明": "工作表「按天」：日期为 YYYY-MM-DD，按 项目名+产品+配置名称+分拆项ID 汇总应付金额",
            },
            {
                "项": "按月",
                "说明": "工作表「按月」：日期为 YYYY-MM，维度与按天相同，费用为当月合计",
            },
            {
                "项": COST_COL_HEADER,
                "说明": "优先 PayableAmount（应付），其次结算/未税金额；列头与单元格均标注人民币 ¥，末行为 SUM 合计",
            },
            {
                "项": SAVING_PLAN_COL_HEADER,
                "说明": "来自 SavingPlanDeductionDiscountTotalAmount（节省计划抵扣总额）；与应付金额分开汇总",
            },
            {
                "项": "分拆项ID",
                "说明": (
                    "来自 ListSplitBillDetail"
                    if stats.split_used
                    else "未开通或无权分账账单时为空；请在费用中心开通分账并给 IAM 授权 billing:ListSplitBillDetail"
                ),
            },
            {"项": "原始条数", "说明": str(stats.raw_count)},
            {
                "项": "排除项目",
                "说明": (
                    f"已排除项目名为空或默认项目 {int(stats.skipped_empty_or_default or 0)} 条"
                    "（default / 默认项目）"
                ),
            },
            {"项": "对照控制台", "说明": "https://console.volcengine.com/finance/bill/detail"},
        ],
    )
    wb.save(dest)
    wb.close()
    return dest


def run_volc_bill_export(
    root_dir: Path,
    access_key_id: str,
    secret_access_key: str,
    *,
    region: str = "cn-beijing",
    bill_period: str | None = None,
    client: VolcOpenAPIClient | None = None,
) -> tuple[Path, BillCollectStats]:
    stats = collect_volc_bills(
        access_key_id,
        secret_access_key,
        region=region,
        bill_period=bill_period,
        client=client,
    )
    dest = build_bill_path(root_dir, bill_period=stats.bill_period)
    save_bill_workbook(dest, stats)
    return dest, stats


def bill_workbook_stats(path: Path) -> dict[str, Any]:
    if not path or not Path(path).exists():
        return {"rows": 0, "sheet": SHEET_DAY, "sheets": {}, "total_rows": 0}
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = {name: sheet_data_rows(wb[name]) for name in wb.sheetnames}
        return {
            "rows": sheets.get(SHEET_DAY, 0),
            "sheet": SHEET_DAY,
            "sheets": sheets,
            "total_rows": sheets.get(SHEET_DAY, 0) + sheets.get(SHEET_MONTH, 0),
        }
    finally:
        wb.close()
