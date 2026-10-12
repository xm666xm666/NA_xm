"""Excel 汇总导出。"""

from __future__ import annotations

import re
import zlib
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .data_summary import (
    CHANNEL_HEADERS,
    DUP_CHANNEL_HEADERS,
    DUP_TOKEN_HEADERS,
    ENVS,
    EXCEL_STEM,
    SHEET_CHANNELS,
    SHEET_DUP_CHANNELS,
    SHEET_DUP_TOKENS,
    SHEET_TOKENS,
    TOKEN_HEADERS,
    build_duplicate_channels,
    build_duplicate_tokens,
    normalize_token_row,
)

TIMESTAMP_FORMAT = "%Y%m%d%H%M%S"
_FILE_STAMP = re.compile(r"(\d{8,14})")


def filename_timestamp(path: Path | str) -> str:
    """从文件名取出时间戳（同名前缀多份时用来比新旧）。"""
    found = _FILE_STAMP.findall(Path(path).stem)
    if not found:
        return ""
    return found[-1].ljust(14, "0")[:14]


def excel_recency_key(path: Path | str) -> str:
    """越大越新：优先文件名时间戳，没有则退回修改时间。"""
    stamp = filename_timestamp(path)
    if stamp:
        return stamp
    item = Path(path)
    try:
        return datetime.fromtimestamp(item.stat().st_mtime).strftime(TIMESTAMP_FORMAT)
    except OSError:
        return "0" * 14


def pick_latest_path(*paths: Path | str | None) -> Path | None:
    existing = [Path(item) for item in paths if item and Path(item).exists()]
    if not existing:
        return None
    unique: dict[str, Path] = {}
    for item in existing:
        unique[str(item.resolve())] = item
    return sorted(unique.values(), key=excel_recency_key, reverse=True)[0]


def sort_excel_files_latest_first(files: list[Path]) -> list[Path]:
    return sorted(files, key=excel_recency_key, reverse=True)


def excel_dir(root_dir: Path) -> Path:
    return root_dir / "excel"


def build_timestamped_filename(when: datetime | None = None) -> str:
    stamp = (when or datetime.now()).strftime(TIMESTAMP_FORMAT)
    return f"{EXCEL_STEM}_{stamp}.xlsx"


def is_staging_excel(path: Path | str) -> bool:
    name = Path(path).name.lower()
    return name.endswith(".part.xlsx")


def is_overview_excel(path: Path | str) -> bool:
    """⑥ 总览文件名带同一前缀，不能当作③「最新网关表」。"""
    name = Path(path).name
    return name.startswith(f"{EXCEL_STEM}_一表总览_") or name.startswith("一表总览_")


def build_staging_filename(when: datetime | None = None) -> str:
    stamp = (when or datetime.now()).strftime(TIMESTAMP_FORMAT)
    return f"{EXCEL_STEM}_{stamp}.part.xlsx"


def build_staging_path(root_dir: Path, when: datetime | None = None) -> Path:
    return excel_dir(root_dir) / build_staging_filename(when)


def publish_staging_excel(path: Path) -> Path:
    """把分站采集用的 .part.xlsx 改名为正式文件，之后才会出现在「最新一份」里。"""
    src = Path(path)
    if not src.is_file():
        raise FileNotFoundError("未找到待发布的 Excel")
    if not is_staging_excel(src):
        return src
    dest = src.with_name(src.name[: -len(".part.xlsx")] + ".xlsx")
    if dest.exists():
        dest.unlink()
    src.replace(dest)
    return dest


def build_timestamped_path(root_dir: Path, when: datetime | None = None) -> Path:
    return excel_dir(root_dir) / build_timestamped_filename(when)


def list_summary_excel_files(root_dir: Path) -> list[Path]:
    directory = excel_dir(root_dir)
    if not directory.exists():
        return []
    files = [
        path
        for path in directory.glob(f"{EXCEL_STEM}*.xlsx")
        if not is_staging_excel(path) and not is_overview_excel(path)
    ]
    return sort_excel_files_latest_first(files)


def latest_excel_path(root_dir: Path) -> Path | None:
    files = list_summary_excel_files(root_dir)
    return files[0] if files else None


def default_excel_path(root_dir: Path) -> Path:
    latest = latest_excel_path(root_dir)
    if latest:
        return latest
    return build_timestamped_path(root_dir)


_GRID_SIDE = Side(style="thin", color="FF000000")
_HEADER_FILL = PatternFill("solid", fgColor="FF2F5496")
_HEADER_FONT = Font(name="宋体", bold=True, color="FFFFFFFF", size=11)
_DATA_FONT = Font(name="宋体", size=11)
_DISABLED_FONT = Font(name="宋体", size=11, color="FFFF0000")
# 紫色用于备注与地址不一致的提示，需与“已禁用”的红色字体明确区分。
_REMARK_MISMATCH_FONT_COLOR = "FF7030A0"
# 分组底色不含红色系，避免与「已禁用」红字冲突
_GROUP_FILLS = (
    PatternFill("solid", fgColor="FFD9E2F3"),
    PatternFill("solid", fgColor="FFE2EFDA"),
    PatternFill("solid", fgColor="FFFCE4D6"),
    PatternFill("solid", fgColor="FFE4DFEC"),
    PatternFill("solid", fgColor="FFFFF2CC"),
    PatternFill("solid", fgColor="FFD0E0E3"),
    PatternFill("solid", fgColor="FFCFE2F3"),
    PatternFill("solid", fgColor="FFD9EAD3"),
    PatternFill("solid", fgColor="FFF9CB9C"),
    PatternFill("solid", fgColor="FFC9DAF8"),
    PatternFill("solid", fgColor="FFD5A6BD"),
    PatternFill("solid", fgColor="FFB6D7A8"),
    PatternFill("solid", fgColor="FF9FC5E8"),
    PatternFill("solid", fgColor="FFB4A7D6"),
    PatternFill("solid", fgColor="FFF6B26B"),
    PatternFill("solid", fgColor="FF76A5AF"),
    PatternFill("solid", fgColor="FFA2C4C9"),
    PatternFill("solid", fgColor="FFD0E8E2"),
    PatternFill("solid", fgColor="FFE8E0F0"),
    PatternFill("solid", fgColor="FFFFE6CC"),
)
_DISABLED_STATUS = "已禁用"


def _palette_index_for_key(key: str) -> int:
    """按分组名稳定映射到调色板，避免「第 N 个分组」顺序导致周期性撞色。"""
    return zlib.adler32(key.encode("utf-8")) % len(_GROUP_FILLS)


def _group_fill(
    group: str,
    cache: dict[str, PatternFill],
    occupied: dict[int, str],
) -> PatternFill | None:
    key = (group or "").strip()
    if not key:
        return None
    if key in cache:
        return cache[key]
    idx = _palette_index_for_key(key)
    if idx in occupied and occupied[idx] != key:
        for step in range(1, len(_GROUP_FILLS)):
            alt = (idx + step) % len(_GROUP_FILLS)
            if alt not in occupied:
                idx = alt
                break
    occupied[idx] = key
    cache[key] = _GROUP_FILLS[idx]
    return cache[key]


def _resolve_color_column(headers: list[str], color_by: str | None) -> int | None:
    if color_by and color_by in headers:
        return headers.index(color_by) + 1
    if "分组" in headers:
        return headers.index("分组") + 1
    return None


def _table_cell_border() -> Border:
    """全表网格线：外框与内框统一黑色细线。"""
    return Border(left=_GRID_SIDE, right=_GRID_SIDE, top=_GRID_SIDE, bottom=_GRID_SIDE)


def _row_is_disabled(ws, row_idx: int, status_col: int | None) -> bool:
    if not status_col:
        return False
    return str(ws.cell(row=row_idx, column=status_col).value or "").strip() == _DISABLED_STATUS


def channel_remark_matches_base_url(base_url: Any, remark: Any) -> bool:
    """备注须含与 base_url 完全相同的一行；两者为空视为无需核对。"""
    source = str(base_url or "").strip()
    note = str(remark or "")
    if not source:
        return not bool(note.strip())
    return source in {line.strip() for line in note.splitlines() if line.strip()}


def apply_channel_remark_mismatch_style(ws) -> None:
    """仅用紫色字体标出未同步的备注，不覆盖原有底色或 base_url 样式。"""
    headers = [str(cell.value or "").strip() for cell in ws[1]]
    if "base_url" not in headers or "备注" not in headers:
        return
    base_col = headers.index("base_url") + 1
    remark_col = headers.index("备注") + 1
    for row_idx in range(2, (ws.max_row or 1) + 1):
        base_url = ws.cell(row=row_idx, column=base_col).value
        remark = ws.cell(row=row_idx, column=remark_col).value
        if channel_remark_matches_base_url(base_url, remark):
            continue
        cell = ws.cell(row=row_idx, column=remark_col)
        font = copy(cell.font)
        font.color = _REMARK_MISMATCH_FONT_COLOR
        cell.font = font


def _apply_sheet_style(ws, headers: list[str], *, color_by: str | None = None) -> None:
    """表头样式、全表黑框、按分组着色；「已禁用」行字体标红（底色不含红色）。"""
    max_col = len(headers)
    max_row = ws.max_row or 0
    if max_row < 1 or max_col < 1:
        return

    color_col = _resolve_color_column(headers, color_by)
    fallback_col = headers.index("项目环境") + 1 if "项目环境" in headers else None
    status_col = headers.index("状态") + 1 if "状态" in headers else None
    group_cache: dict[str, PatternFill] = {}
    occupied_indices: dict[int, str] = {}

    for row_idx in range(1, max_row + 1):
        is_header = row_idx == 1
        row_fill: PatternFill | None = None
        if not is_header and color_col:
            color_value = str(ws.cell(row=row_idx, column=color_col).value or "").strip()
            if not color_value and fallback_col:
                color_value = str(ws.cell(row=row_idx, column=fallback_col).value or "").strip()
            row_fill = _group_fill(color_value, group_cache, occupied_indices)

        disabled = (not is_header) and _row_is_disabled(ws, row_idx, status_col)
        data_font = _DISABLED_FONT if disabled else _DATA_FONT

        for col_idx in range(1, max_col + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = _table_cell_border()
            cell.alignment = Alignment(vertical="center", wrap_text=col_idx > 1)
            if is_header:
                cell.fill = _HEADER_FILL
                cell.font = _HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cell.font = data_font
                if row_fill is not None:
                    cell.fill = row_fill

    ws.freeze_panes = "A2"
    for col_idx, header in enumerate(headers, start=1):
        max_len = len(str(header))
        for row_idx in range(2, min(ws.max_row + 1, 200)):
            value = ws.cell(row=row_idx, column=col_idx).value
            if value is not None:
                max_len = max(max_len, min(len(str(value)), 48))
        ws.column_dimensions[get_column_letter(col_idx)].width = max(max_len + 2, 10)
    apply_channel_remark_mismatch_style(ws)


def _ensure_sheet(wb, title: str, headers: list[str]):
    if title in wb.sheetnames:
        return wb[title]
    ws = wb.create_sheet(title)
    ws.append(headers)
    return ws


def _read_sheet_rows(ws, headers: list[str]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    header_row = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    mapping = {str(name): idx for idx, name in enumerate(header_row) if name}
    for row_cells in ws.iter_rows(min_row=2, values_only=True):
        if not row_cells or not any(row_cells):
            continue
        item: dict[str, Any] = {}
        for header in headers:
            idx = mapping.get(header)
            value = row_cells[idx] if idx is not None and idx < len(row_cells) else ""
            item[header] = "" if value is None else value
        if str(item.get("名称") or "").strip():
            if headers == TOKEN_HEADERS:
                rows.append(normalize_token_row(item))
            else:
                rows.append(item)
    return rows


def _row_cell_values(headers: list[str], row: dict[str, Any]) -> list[str]:
    if headers == TOKEN_HEADERS:
        normalized = normalize_token_row(row)
        return [normalized[header] for header in headers]
    return ["" if row.get(header) is None else str(row.get(header, "")) for header in headers]


def _write_sheet(
    ws,
    headers: list[str],
    rows: list[dict[str, Any]],
    *,
    color_by: str | None = None,
) -> None:
    ws.delete_rows(1, ws.max_row or 1)
    for col_idx, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col_idx, value=header)
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, value in enumerate(_row_cell_values(headers, row), start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    _apply_sheet_style(ws, headers, color_by=color_by)


def load_workbook_data(path: Path) -> tuple[dict[str, list[dict[str, Any]]], dict[str, list[dict[str, Any]]]]:
    channel_env_rows: dict[str, list[dict[str, Any]]] = {env: [] for env in ENVS}
    token_env_rows: dict[str, list[dict[str, Any]]] = {env: [] for env in ENVS}
    if not path.exists():
        return channel_env_rows, token_env_rows

    wb = load_workbook(path)
    for env in ENVS:
        ch_title = SHEET_CHANNELS[env]
        if ch_title in wb.sheetnames:
            channel_env_rows[env] = _read_sheet_rows(wb[ch_title], CHANNEL_HEADERS)
        tok_title = SHEET_TOKENS[env]
        if tok_title in wb.sheetnames:
            token_env_rows[env] = _read_sheet_rows(wb[tok_title], TOKEN_HEADERS)
    wb.close()
    return channel_env_rows, token_env_rows


def save_workbook_data(
    path: Path,
    *,
    channel_env_rows: dict[str, list[dict[str, Any]]],
    token_env_rows: dict[str, list[dict[str, Any]]],
    channel_dup_source: dict[str, list[dict[str, Any]]] | None = None,
    token_dup_source: dict[str, list[dict[str, Any]]] | None = None,
    channel_headers: list[str] | None = None,
    dup_channel_headers: list[str] | None = None,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        wb = load_workbook(path)
    else:
        wb = Workbook()
        default = wb.active
        wb.remove(default)

    ch_headers = list(channel_headers or CHANNEL_HEADERS)
    dup_ch_headers = list(dup_channel_headers or DUP_CHANNEL_HEADERS)

    for env in ENVS:
        _write_sheet(
            _ensure_sheet(wb, SHEET_CHANNELS[env], ch_headers),
            ch_headers,
            channel_env_rows.get(env, []),
            color_by="分组",
        )
        _write_sheet(
            _ensure_sheet(wb, SHEET_TOKENS[env], TOKEN_HEADERS),
            TOKEN_HEADERS,
            token_env_rows.get(env, []),
            color_by="分组",
        )

    ch_source = channel_dup_source if channel_dup_source is not None else channel_env_rows
    tk_source = token_dup_source if token_dup_source is not None else token_env_rows
    dup_channels = build_duplicate_channels(ch_source)
    dup_tokens = build_duplicate_tokens(tk_source)
    _write_sheet(
        _ensure_sheet(wb, SHEET_DUP_CHANNELS, dup_ch_headers),
        dup_ch_headers,
        dup_channels,
        color_by="模型列表",
    )
    _write_sheet(
        _ensure_sheet(wb, SHEET_DUP_TOKENS, DUP_TOKEN_HEADERS),
        DUP_TOKEN_HEADERS,
        dup_tokens,
        color_by="分组",
    )

    wb.save(path)
    wb.close()


def create_empty_workbook(path: Path) -> None:
    """写入仅含表头的空工作簿（测试或手动初始化用；Web 采集不再自动调用）。"""
    save_workbook_data(path, channel_env_rows={env: [] for env in ENVS}, token_env_rows={env: [] for env in ENVS})


def ensure_workbook(path: Path) -> bool:
    """若 Excel 不存在则创建空模板，返回是否新建。

    Web 数据汇总已不再自动创建空文件；保留供测试或脚本显式调用。
    """
    if path.exists():
        return False
    create_empty_workbook(path)
    return True


def workbook_summary(path: Path) -> dict[str, Any]:
    """读取各工作表行数统计。"""
    channel_env_rows, token_env_rows = load_workbook_data(path)
    return {
        "channels": {env: len(channel_env_rows.get(env, [])) for env in ENVS},
        "tokens": {env: len(token_env_rows.get(env, [])) for env in ENVS},
        "duplicate_channels": len(build_duplicate_channels(channel_env_rows)),
        "duplicate_tokens": len(build_duplicate_tokens(token_env_rows)),
    }
