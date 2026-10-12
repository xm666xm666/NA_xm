# -*- coding: utf-8 -*-
"""上游信息 · 火山云采集（与 AI 网关数据汇总完全独立）。"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from .deploy_time import sort_rows_by_deploy_time_desc
from .volc_openapi import VolcOpenAPIClient, VolcOpenAPIError

_CN_TZ = timezone(timedelta(hours=8))

UPSTREAM_HEADERS = [
    "接入点名称",
    "部署名",
    "模型服务",
    "API Key",
    "TPM",
    "RPM",
    "项目方",
    "联系人",
    "API Key汇总",
    "部署时间",
]
COL_API_KEY = "API Key"
COL_API_KEY_CATALOG = "API Key汇总"


@dataclass
class UpstreamCollectStats:
    projects_ok: int = 0
    projects_failed: int = 0
    endpoints: int = 0
    project_results: list[dict[str, Any]] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    rows: list[dict[str, Any]] = field(default_factory=list)
    token_overlay: int = 0
    token_overlay_file: str = ""


def format_model_service(endpoint: dict[str, Any]) -> str:
    ref = endpoint.get("ModelReference") or {}
    if not isinstance(ref, dict):
        return ""
    foundation = ref.get("FoundationModel") or {}
    if isinstance(foundation, dict):
        name = str(foundation.get("Name") or "").strip()
        version = str(foundation.get("ModelVersion") or "").strip()
        if name and version:
            return f"{name}｜{version}"
        return name or version
    custom_id = str(ref.get("CustomModelId") or "").strip()
    return custom_id


def foundation_model_name(endpoint: dict[str, Any]) -> str:
    ref = endpoint.get("ModelReference") or {}
    if not isinstance(ref, dict):
        return ""
    foundation = ref.get("FoundationModel") or {}
    if isinstance(foundation, dict):
        return str(foundation.get("Name") or "").strip()
    return ""


_ENV_GROUPS: dict[str, tuple[str, ...]] = {
    "prd": ("prd", "prod"),
    "uat": ("uat", "dev"),
    "test": ("test", "demo"),
}
_SPECIFIC_TO_GROUP = {
    "prd": "prd",
    "prod": "prd",
    "uat": "uat",
    "dev": "uat",
    "test": "test",
    "demo": "test",
}
# 先匹配更完整的词，避免 prod/prd、demo/test 互相误伤
_SPECIFIC_WORDS = ("prod", "prd", "uat", "demo", "test", "dev")
# 无对应环境 Key 时：test 用 uat/dev；不要落到 prd，也不要让 uat/dev 去抢 test Key
_FALLBACK_GROUPS: dict[str, tuple[str, ...]] = {
    "test": ("uat",),
}
_TOKEN_UUID_RE = re.compile(
    r"[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}",
    re.I,
)
_REAL_TOKEN_FILENAME = "火山真实数据.xlsx"


def _normalize_env_text(text: str) -> str:
    return (text or "").strip().lower().replace("_", "-")


def detect_env_specific(text: str) -> str:
    """识别具体环境词：prd / prod / uat / dev / test / demo。"""
    raw = _normalize_env_text(text)
    if not raw:
        return ""
    for word in _SPECIFIC_WORDS:
        if re.search(rf"(^|-){re.escape(word)}($|-)", raw):
            return word
    for word in _SPECIFIC_WORDS:
        if word in raw:
            return word
    return ""


def detect_env_tag(text: str) -> str:
    """从接入点/API Key 名称识别分组：prd / uat / test。

    - prd：``prd`` / ``prod``
    - uat：``uat`` / ``dev``（同项目同时有 uat Key 和 dev Key 时改为精确对应）
    - test：``test`` / ``demo``
    """
    return _SPECIFIC_TO_GROUP.get(detect_env_specific(text), "")


def _score_name_for_aliases(key_name: str, aliases: tuple[str, ...]) -> int:
    name = _normalize_env_text(key_name)
    if not name or not aliases:
        return 0
    score = 0
    for alias in aliases:
        exacts = (
            f"token-{alias}",
            f"{alias}-key",
            f"{alias}-token",
            f"key-{alias}",
            alias,
        )
        if name in exacts or name.endswith(f"-{alias}") or name.startswith(f"{alias}-"):
            score = max(score, 100)
        if alias in name:
            score = max(score, 50)
    return score


def score_api_key_for_env(env: str, key_name: str) -> int:
    """环境分组与 API Key 名称的匹配分；0 表示不匹配。"""
    if not env:
        return 0
    aliases = _ENV_GROUPS.get(env, (env,))
    return _score_name_for_aliases(key_name, aliases)


def group_aliases_are_split(api_keys: list[dict[str, Any]], group: str) -> bool:
    """同组两个别名是否都作为独立 Key 出现（如同时有 token-uat 与 token-dev）。"""
    aliases = set(_ENV_GROUPS.get(group) or ())
    if len(aliases) < 2:
        return False
    found: set[str] = set()
    for item in api_keys:
        spec = detect_env_specific(str(item.get("Name") or ""))
        if spec in aliases:
            found.add(spec)
        if found >= aliases:
            return True
    return found >= aliases


def _api_key_sort_id(item: dict[str, Any]) -> int:
    raw = item.get("Id")
    try:
        return int(raw)
    except (TypeError, ValueError):
        return 0


def pick_default_project_key(
    api_keys: list[dict[str, Any]],
    raw_key_by_id: dict[Any, str] | None = None,
) -> dict[str, Any] | None:
    """项目内 Key 对接入点通用：优先已有明文的最老一把（Id 最小），否则最老一把。"""
    if not api_keys:
        return None
    raw = raw_key_by_id or {}
    with_plain = [
        item
        for item in api_keys
        if item.get("Id") is not None and str(raw.get(item.get("Id")) or "").strip()
    ]
    pool = with_plain or list(api_keys)
    return min(pool, key=_api_key_sort_id)


def env_groups_in_keys(api_keys: list[dict[str, Any]]) -> set[str]:
    groups: set[str] = set()
    for item in api_keys:
        spec = detect_env_specific(str(item.get("Name") or ""))
        group = _SPECIFIC_TO_GROUP.get(spec, "")
        if group:
            groups.add(group)
    return groups


def has_generic_api_key(api_keys: list[dict[str, Any]]) -> bool:
    return any(not detect_env_specific(str(item.get("Name") or "")) for item in api_keys)


def is_partial_env_split(api_keys: list[dict[str, Any]]) -> bool:
    """有通用 Key、且环境 Key 只覆盖一个分组（如 clinai：一把旧 Key + 一把新 test Key）。"""
    return has_generic_api_key(api_keys) and len(env_groups_in_keys(api_keys)) <= 1


def format_deploy_time(value: Any) -> str:
    """统一写成 YYYY-MM-DD HH:MM:SS（东八区）。"""
    if value in (None, ""):
        return ""
    if isinstance(value, datetime):
        dt = value
        if dt.tzinfo is not None:
            dt = dt.astimezone(_CN_TZ)
        return dt.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, (int, float)):
        ts = float(value)
        if ts > 1e12:
            ts /= 1000.0
        if ts <= 0:
            return ""
        return datetime.fromtimestamp(ts, _CN_TZ).strftime("%Y-%m-%d %H:%M:%S")
    text = str(value).strip()
    if not text:
        return ""
    if re.fullmatch(r"\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}", text):
        return text
    if re.fullmatch(r"\d{10,13}", text):
        return format_deploy_time(int(text))
    iso = text.replace("Z", "+00:00")
    try:
        dt = datetime.fromisoformat(iso)
    except ValueError:
        return ""
    if dt.tzinfo is not None:
        dt = dt.astimezone(_CN_TZ)
    return dt.strftime("%Y-%m-%d %H:%M:%S")


def format_endpoint_deploy_time(endpoint: dict[str, Any]) -> str:
    raw = endpoint.get("CreateTime") or endpoint.get("create_time") or endpoint.get("CreatedAt")
    formatted = format_deploy_time(raw)
    if formatted:
        return formatted
    stamp = endpoint_sort_stamp(str(endpoint.get("Id") or ""))
    if len(stamp) >= 14 and stamp.isdigit():
        return (
            f"{stamp[0:4]}-{stamp[4:6]}-{stamp[6:8]} "
            f"{stamp[8:10]}:{stamp[10:12]}:{stamp[12:14]}"
        )
    return ""


def endpoint_sort_stamp(endpoint_id: str) -> str:
    found = re.search(r"ep-(\d{8,14})", str(endpoint_id or ""), re.I)
    if not found:
        return ""
    return found.group(1).ljust(14, "0")


def is_newest_endpoint_in_env(
    endpoint_name: str,
    endpoint_id: str,
    sibling_endpoints: list[tuple[str, str]] | None,
    group: str,
) -> bool:
    """同环境接入点里，当前是否为部署名时间戳最新的一条。"""
    if not group or not sibling_endpoints:
        return False
    self_stamp = endpoint_sort_stamp(endpoint_id)
    if not self_stamp:
        return False
    for _name, other_id in sibling_endpoints:
        other_group = _SPECIFIC_TO_GROUP.get(detect_env_specific(_name), "")
        if other_group != group:
            continue
        other_stamp = endpoint_sort_stamp(other_id)
        if not other_stamp:
            continue
        if other_stamp > self_stamp:
            return False
        if other_stamp == self_stamp and str(other_id) > str(endpoint_id):
            return False
    return True


def _best_key_for_aliases(
    api_keys: list[dict[str, Any]],
    aliases: tuple[str, ...],
) -> dict[str, Any] | None:
    best: dict[str, Any] | None = None
    best_score = 0
    best_id = 10**18
    for item in api_keys:
        score = _score_name_for_aliases(str(item.get("Name") or ""), aliases)
        if score <= 0:
            continue
        kid = _api_key_sort_id(item)
        if score > best_score or (score == best_score and kid < best_id):
            best_score = score
            best = item
            best_id = kid
    return best


def match_api_key_for_endpoint(
    endpoint_name: str,
    api_keys: list[dict[str, Any]],
    raw_key_by_id: dict[Any, str] | None = None,
    *,
    endpoint_id: str = "",
    sibling_endpoints: list[tuple[str, str]] | None = None,
) -> dict[str, Any] | None:
    """按 prd/uat/test 将接入点匹配到项目下 API Key。

    默认 uat 与 dev 互通（prd/prod、test/demo 同理）。
    极端情况：同项目同时存在 uat Key 与 dev Key 时，uat 只对 uat、dev 只对 dev。
    对不上或接入点无环境词时，默认用项目里一把 Key（项目内通用）。
    无 test Key 时 test 接入点改用 uat/dev Key，避免误落到 prd。
    只有一把环境 Key（其余为通用 Key）时：该环境只把最新接入点挂到这把 Key，旧接入点仍用通用 Key。
    """
    if not api_keys:
        return None
    if len(api_keys) == 1:
        return api_keys[0]
    specific = detect_env_specific(endpoint_name)
    group = _SPECIFIC_TO_GROUP.get(specific, "")
    default = pick_default_project_key(api_keys, raw_key_by_id)
    if group:
        aliases = _ENV_GROUPS[group]
        if specific and group_aliases_are_split(api_keys, group):
            aliases = (specific,)
        hit = _best_key_for_aliases(api_keys, aliases)
        if hit is None:
            for fallback in _FALLBACK_GROUPS.get(group, ()):
                hit = _best_key_for_aliases(api_keys, _ENV_GROUPS[fallback])
                if hit is not None:
                    break
        if hit is not None:
            if is_partial_env_split(api_keys):
                if is_newest_endpoint_in_env(
                    endpoint_name, endpoint_id, sibling_endpoints, group
                ):
                    return hit
                return default
            return hit
    return default


def normalize_token_uuid(token: str) -> str:
    """去掉 ark- 前缀后取出 UUID，用于和真实表对齐。"""
    text = str(token or "").strip()
    if not text or text.lower() in {"nan", "none"}:
        return ""
    text = re.sub(r"^ark-", "", text, flags=re.I)
    found = _TOKEN_UUID_RE.search(text)
    return found.group(0).lower() if found else text.lower()


def find_real_token_reference(*directories: Path) -> Path | None:
    """查找「火山真实数据.xlsx」，用作 token 对齐基准。"""
    for directory in directories:
        if directory is None:
            continue
        root = Path(directory)
        if not root.is_dir():
            continue
        preferred = root / _REAL_TOKEN_FILENAME
        if preferred.is_file():
            return preferred
        extras = sorted(
            (
                path
                for path in root.glob("*真实数据*.xlsx")
                if "比对" not in path.name
            ),
            key=lambda path: path.stat().st_mtime,
            reverse=True,
        )
        if extras:
            return extras[0]
    return None


def format_api_key_catalog(
    api_keys: list[dict[str, Any]],
    raw_key_by_id: dict[Any, str] | None = None,
) -> str:
    """项目下全部 API Key，对齐控制台「查看 API Key」：每行「名称 | API Key」。"""
    raw = raw_key_by_id or {}
    lines: list[str] = []
    for item in sorted(api_keys, key=_api_key_sort_id):
        name = str(item.get("Name") or "").strip() or "(未命名)"
        key_id = item.get("Id")
        value = ""
        if key_id is not None:
            value = str(raw.get(key_id) or "").strip()
        lines.append(f"{name} | {value}".rstrip() if value else name)
    return "\n".join(lines)


def _row_api_key(row: dict[str, Any]) -> str:
    return str(row.get(COL_API_KEY) or row.get("token") or "").strip()


def _header_api_key_index(names: list[str]) -> int:
    for candidate in (COL_API_KEY, "token", "Token"):
        if candidate in names:
            return names.index(candidate)
    raise ValueError("missing API Key column")


def load_real_tokens_by_deploy_id(path: Path) -> dict[str, str]:
    """从真实数据工作簿读取 部署名 → API Key。兼容旧列名 token。"""
    from openpyxl import load_workbook

    mapping: dict[str, str] = {}
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        title = "火山云" if "火山云" in workbook.sheetnames else (
            workbook.sheetnames[0] if workbook.sheetnames else ""
        )
        if not title:
            return mapping
        rows = workbook[title].iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            return mapping
        names = [str(cell or "").strip() for cell in header]
        try:
            ep_idx = names.index("部署名")
            token_idx = _header_api_key_index(names)
        except ValueError:
            return mapping
        for row in rows:
            if not row or ep_idx >= len(row):
                continue
            deploy_id = str(row[ep_idx] or "").strip()
            token = ""
            if token_idx < len(row):
                token = str(row[token_idx] or "").strip()
            if token.lower() in {"nan", "none"}:
                token = ""
            if deploy_id and token:
                mapping[deploy_id] = token
    finally:
        workbook.close()
    return mapping


def apply_real_token_overlay(
    rows: list[dict[str, Any]],
    real_by_deploy_id: dict[str, str],
) -> int:
    """按部署名把 API Key 改成真实表里的值；返回改写条数。"""
    if not rows or not real_by_deploy_id:
        return 0
    changed = 0
    for row in rows:
        deploy_id = str(row.get("部署名") or "").strip()
        real = (real_by_deploy_id.get(deploy_id) or "").strip()
        if not real:
            continue
        current = _row_api_key(row)
        if current != real:
            changed += 1
        row[COL_API_KEY] = real
        if "token" in row and COL_API_KEY != "token":
            row.pop("token", None)
    return changed


def resolve_endpoint_token(
    endpoint_name: str,
    api_keys: list[dict[str, Any]],
    raw_key_by_id: dict[Any, str],
    *,
    endpoint_id: str = "",
    sibling_endpoints: list[tuple[str, str]] | None = None,
) -> tuple[str, str]:
    """返回 (token明文或空, 匹配到的API Key名称)。项目内对不上时用默认一把。"""
    matched = match_api_key_for_endpoint(
        endpoint_name,
        api_keys,
        raw_key_by_id,
        endpoint_id=endpoint_id,
        sibling_endpoints=sibling_endpoints,
    )
    if not matched:
        return "", ""
    key_id = matched.get("Id")
    name = str(matched.get("Name") or "").strip()
    token = ""
    if key_id is not None and key_id in raw_key_by_id:
        token = str(raw_key_by_id.get(key_id) or "").strip()
    if not token:
        fallback = pick_default_project_key(api_keys, raw_key_by_id)
        if fallback is not None:
            key_id = fallback.get("Id")
            name = str(fallback.get("Name") or "").strip()
            if key_id is not None and key_id in raw_key_by_id:
                token = str(raw_key_by_id.get(key_id) or "").strip()
    return token, name


_GENERIC_PROJECT_OWNERS = frozenset(
    {
        "",
        "n/a",
        "na",
        "none",
        "null",
        "default",
        "默认项目",
        "默认",
        "账号全部资源",
        "全部资源",
    }
)


def extract_project_owner(project_name: str, display_name: str = "") -> str:
    """从火山项目 DisplayName 取出负责人。

    控制台约定为 ``项目名-负责人``（如 ``phivb-范亚浩``、``smo-rpm-金亥其``）。
    没有负责人（DisplayName 为空、等于项目名、或仅为「默认项目」等）时返回 ``N/A``。
    """
    name = str(project_name or "").strip()
    display = str(display_name or "").strip()
    owner = ""
    prefix = f"{name}-" if name else ""
    if prefix and display.startswith(prefix):
        owner = display[len(prefix) :].strip()
    elif prefix and display.lower().startswith(prefix.lower()):
        owner = display[len(prefix) :].strip()
    elif display and display != name:
        owner = display
    if owner.lower() in _GENERIC_PROJECT_OWNERS or owner == name:
        return "N/A"
    return owner or "N/A"


def format_project_party(project_name: str, display_name: str = "") -> str:
    """项目方列：``项目名-负责人``；无负责人时为 ``项目名-N/A``。"""
    name = str(project_name or "").strip()
    owner = extract_project_owner(name, display_name)
    if not name:
        return owner
    return f"{name}-{owner}"


def _positive_rate(value: Any) -> int | None:
    """有效限流为正整数；缺省 / -1（不限或未设置）视为空。"""
    if value is None or value == "":
        return None
    try:
        num = int(value)
    except (TypeError, ValueError):
        return None
    return num if num > 0 else None


def build_model_rate_map(items: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    """FoundationModelName -> CurrentRateLimit（优先）/ DefaultRateLimit。"""
    mapping: dict[str, dict[str, Any]] = {}
    for item in items:
        name = str(item.get("FoundationModelName") or "").strip()
        if not name:
            continue
        current = item.get("CurrentRateLimit")
        default = item.get("DefaultRateLimit")
        if isinstance(current, dict):
            mapping[name] = current
        elif isinstance(default, dict):
            mapping[name] = default
    return mapping


def resolve_rpm_tpm(
    endpoint: dict[str, Any],
    *,
    model_rate_map: dict[str, dict[str, Any]] | None = None,
) -> tuple[Any, Any]:
    """解析 RPM/TPM。

    优先级：
    1. 接入点自带 RateLimit（Rpm/Tpm > 0）——自定义限流
    2. ListModelRateLimit 按基础模型名匹配的 CurrentRateLimit
    """
    ep_rate = endpoint.get("RateLimit")
    if isinstance(ep_rate, dict):
        rpm = _positive_rate(ep_rate.get("Rpm"))
        tpm = _positive_rate(ep_rate.get("Tpm"))
        if rpm is not None or tpm is not None:
            return (
                "" if rpm is None else rpm,
                "" if tpm is None else tpm,
            )

    model_name = foundation_model_name(endpoint)
    model_rate = (model_rate_map or {}).get(model_name) if model_name else None
    if isinstance(model_rate, dict):
        rpm = _positive_rate(model_rate.get("Rpm"))
        tpm = _positive_rate(model_rate.get("Tpm"))
        return (
            "" if rpm is None else rpm,
            "" if tpm is None else tpm,
        )
    return "", ""


def endpoint_to_row(
    endpoint: dict[str, Any],
    *,
    project_name: str,
    token: str = "",
    api_key: str = "",
    api_key_catalog: str = "",
    contact: str = "",
    display_name: str = "",
    model_rate_map: dict[str, dict[str, Any]] | None = None,
) -> dict[str, Any]:
    rpm, tpm = resolve_rpm_tpm(endpoint, model_rate_map=model_rate_map)
    raw_name = str(endpoint.get("ProjectName") or project_name or "").strip()
    party = format_project_party(raw_name, display_name)
    owner = extract_project_owner(raw_name, display_name)
    return {
        "接入点名称": str(endpoint.get("Name") or "").strip(),
        "部署名": str(endpoint.get("Id") or "").strip(),
        "模型服务": format_model_service(endpoint),
        COL_API_KEY: (api_key or token or "").strip(),
        "TPM": tpm,
        "RPM": rpm,
        "项目方": party,
        "联系人": str(contact or "").strip() or owner,
        COL_API_KEY_CATALOG: api_key_catalog,
        "部署时间": format_endpoint_deploy_time(endpoint),
    }


def test_volc_connection(
    access_key_id: str,
    secret_access_key: str,
    *,
    region: str = "cn-beijing",
) -> dict[str, Any]:
    client = VolcOpenAPIClient(access_key_id, secret_access_key, region=region)
    projects = client.list_projects()
    names = []
    for item in projects:
        name = str(item.get("ProjectName") or "").strip()
        if not name:
            continue
        names.append(format_project_party(name, str(item.get("DisplayName") or "")))
    return {
        "project_count": len(names),
        "projects": names,
        "region": region,
    }


def collect_volc_upstream(
    access_key_id: str,
    secret_access_key: str,
    *,
    region: str = "cn-beijing",
    project_names: list[str] | None = None,
    include_api_keys: bool = True,
    real_token_by_ep: dict[str, str] | None = None,
    real_token_file: str = "",
) -> UpstreamCollectStats:
    stats = UpstreamCollectStats()
    client = VolcOpenAPIClient(access_key_id, secret_access_key, region=region)

    try:
        all_projects = client.list_projects()
    except VolcOpenAPIError as exc:
        stats.errors.append(str(exc))
        return stats
    except Exception as exc:  # noqa: BLE001
        stats.errors.append(f"ListProjects 失败: {exc}")
        return stats

    wanted = {str(x).strip() for x in (project_names or []) if str(x).strip()}
    projects = []
    for item in all_projects:
        name = str(item.get("ProjectName") or "").strip()
        if not name:
            continue
        if wanted and name not in wanted:
            continue
        projects.append(item)

    if not projects:
        stats.errors.append("未找到可采集的项目（请检查 IAM 权限或项目筛选）")
        return stats

    model_rate_map: dict[str, dict[str, Any]] = {}
    try:
        model_rate_map = build_model_rate_map(client.list_model_rate_limits())
    except VolcOpenAPIError as exc:
        stats.errors.append(f"ListModelRateLimit: {exc}（RPM/TPM 可能为空）")
    except Exception as exc:  # noqa: BLE001
        stats.errors.append(f"ListModelRateLimit 失败: {exc}（RPM/TPM 可能为空）")

    for project in projects:
        name = str(project.get("ProjectName") or "").strip()
        display = str(project.get("DisplayName") or "").strip()
        label = format_project_party(name, display)
        try:
            endpoints = client.list_endpoints(name)
            api_keys: list[dict[str, Any]] = []
            raw_key_by_id: dict[Any, str] = {}
            if include_api_keys:
                try:
                    api_keys = client.list_api_keys(name)
                except VolcOpenAPIError as exc:
                    stats.errors.append(f"[{name}] ListApiKeys: {exc}")
                for key_item in api_keys:
                    key_id = key_item.get("Id")
                    if key_id is None or key_id in raw_key_by_id:
                        continue
                    try:
                        raw_key_by_id[key_id] = client.get_raw_api_key(key_id)
                    except VolcOpenAPIError as exc:
                        stats.errors.append(
                            f"[{name}] GetRawApiKey({key_item.get('Name') or key_id}): {exc}"
                        )
                        raw_key_by_id[key_id] = ""
            catalog = format_api_key_catalog(api_keys, raw_key_by_id)
            rows: list[dict[str, Any]] = []
            missing_token = 0
            sibling_endpoints = [
                (str(item.get("Name") or ""), str(item.get("Id") or ""))
                for item in endpoints
                if str(item.get("Id") or "").strip() or str(item.get("Name") or "").strip()
            ]
            for ep in endpoints:
                if not (
                    str(ep.get("Id") or "").strip() or str(ep.get("Name") or "").strip()
                ):
                    continue
                token, key_name = resolve_endpoint_token(
                    str(ep.get("Name") or ""),
                    api_keys,
                    raw_key_by_id,
                    endpoint_id=str(ep.get("Id") or ""),
                    sibling_endpoints=sibling_endpoints,
                )
                if include_api_keys and not token:
                    missing_token += 1
                rows.append(
                    endpoint_to_row(
                        ep,
                        project_name=name,
                        display_name=display,
                        api_key=token,
                        api_key_catalog=catalog,
                        model_rate_map=model_rate_map,
                    )
                )
            stats.rows.extend(rows)
            stats.endpoints += len(rows)
            stats.projects_ok += 1
            msg = f"采集 {len(rows)} 条接入点"
            if include_api_keys and api_keys:
                filled = len(rows) - missing_token
                msg += f"；API Key 已填 {filled}/{len(rows)}"
                if missing_token:
                    msg += "（仍空的是项目下没有可用 Key 明文）"
            stats.project_results.append(
                {
                    "project": name,
                    "display_name": display,
                    "label": label,
                    "success": True,
                    "endpoints": len(rows),
                    "message": msg,
                }
            )
        except VolcOpenAPIError as exc:
            stats.projects_failed += 1
            stats.errors.append(f"[{name}] {exc}")
            stats.project_results.append(
                {
                    "project": name,
                    "display_name": display,
                    "label": label,
                    "success": False,
                    "endpoints": 0,
                    "message": str(exc),
                }
            )
        except Exception as exc:  # noqa: BLE001
            stats.projects_failed += 1
            msg = f"[{name}] {exc}"
            stats.errors.append(msg)
            stats.project_results.append(
                {
                    "project": name,
                    "display_name": display,
                    "label": label,
                    "success": False,
                    "endpoints": 0,
                    "message": str(exc),
                }
            )

    stats.rows = sort_rows_by_deploy_time_desc(stats.rows)
    if include_api_keys and not real_token_by_ep:
        excel_dir = Path(__file__).resolve().parents[2] / "excel"
        ref = find_real_token_reference(excel_dir)
        if ref is not None:
            real_token_by_ep = load_real_tokens_by_deploy_id(ref)
            real_token_file = real_token_file or ref.name
    if include_api_keys and real_token_by_ep:
        stats.token_overlay = apply_real_token_overlay(stats.rows, real_token_by_ep)
        stats.token_overlay_file = real_token_file
    return stats
