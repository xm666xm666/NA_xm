"""网站渠道与令牌数据采集、去重与 Excel 汇总。"""

from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

from .api_client import NewAPIClient, SecureVerificationError, authenticate_client, coerce_verify_ssl, refresh_auth_session_cache
from .channel_types import load_bundled_types
from .rate_budget import (
    assert_can_proceed,
    note_rate_limited,
    remaining,
    seconds_until_reset,
)
from .token_api_fields import read_model_limits
from .totp import has_otp_source, pick_totp_secret, resolve_otp_code
from .url_utils import normalize_na_base_url

ENVS = ("test", "uat", "prod")
EXCEL_STEM = "ai网关令牌及渠道列表"
DEFAULT_EXCEL_NAME = f"{EXCEL_STEM}.xlsx"  # 无时间戳的旧文件名，兼容历史文件

CHANNEL_HEADERS = [
    "名称", "分组", "类型", "状态", "响应时间", "模型列表", "模型重定向",
    "base_url", "备注", "部署时间",
]
CHANNEL_HEADERS_WITH_KEY = [
    "名称", "分组", "类型", "状态", "响应时间", "模型列表", "模型重定向",
    "base_url", "备注", "密钥", "部署时间",
]
TOKEN_HEADERS = ["名称", "项目环境", "状态", "剩余额度", "分组", "可用模型", "包含渠道", "IP限制", "创建时间", "过期时间"]
BATCH_CONFIG_HEADER = "批量配置"
DUP_CHANNEL_HEADERS = ["名称", "类型", "状态", "模型列表", "模型重定向", "base_url", "备注", "部署时间", "来源", "重复序号"]
DUP_CHANNEL_HEADERS_WITH_KEY = [
    "名称", "类型", "状态", "模型列表", "模型重定向", "base_url", "备注", "密钥", "部署时间", "来源", "重复序号",
]
DUP_TOKEN_HEADERS = [
    "名称", "项目环境", "状态", "剩余额度", "分组", "可用模型", "包含渠道",
    "IP限制", "创建时间", "过期时间", "重复序号", "来源",
]

SHEET_CHANNELS = {env: f"{env}-渠道管理" for env in ENVS}
SHEET_TOKENS = {env: f"{env}-令牌管理" for env in ENVS}
SHEET_DUP_CHANNELS = "重复模型"
SHEET_DUP_TOKENS = "重复模型相关令牌"


@dataclass
class SiteConfig:
    base_url: str
    username: str = ""
    password: str = ""
    env: str = ""
    label: str = ""
    otp_code: str = ""
    totp_secret: str = ""
    access_token: str = ""
    user_id: int | None = None
    verify_ssl: bool = True


@dataclass
class CollectStats:
    sites_ok: int = 0
    sites_failed: int = 0
    channels_added: dict[str, int] = field(default_factory=dict)
    tokens_added: dict[str, int] = field(default_factory=dict)
    duplicate_channels: int = 0
    duplicate_tokens: int = 0
    channels_skipped_duplicate: int = 0
    tokens_skipped_duplicate: int = 0
    channels_skipped_by_env: dict[str, int] = field(default_factory=dict)
    tokens_skipped_by_env: dict[str, int] = field(default_factory=dict)
    channels_in_sheet: dict[str, int] = field(default_factory=dict)
    tokens_in_sheet: dict[str, int] = field(default_factory=dict)
    site_results: list[dict[str, Any]] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    include_keys: bool = False
    keys_fetched: int = 0
    keys_failed: int = 0


def _build_type_label_map() -> dict[int, str]:
    mapping: dict[int, str] = {}
    for item in load_bundled_types():
        value = item.get("value")
        label = item.get("label")
        if value is not None and label:
            mapping[int(value)] = str(label)
    return mapping


def detect_env(base_url: str, explicit: str = "") -> str:
    env = (explicit or "").strip().lower()
    if env in ENVS:
        return env
    text = (base_url or "").lower()
    if re.search(r"[\-/_.]uat[\-/_.]|uat[\-/_.]", text):
        return "uat"
    if re.search(r"[\-/_.]prod[\-/_.]|prod[\-/_.]|[\-/_.]prd[\-/_.]", text):
        return "prod"
    if re.search(r"[\-/_.]test[\-/_.]|test[\-/_.]", text):
        return "test"
    return "test"


def _env_sites_path() -> Path:
    return Path(__file__).resolve().parents[2] / "env" / "sites.json"


def load_raw_env_sites() -> dict[str, Any]:
    path = _env_sites_path()
    if not path.exists():
        return {}
    try:
        with path.open(encoding="utf-8") as f:
            data = json.load(f)
    except (json.JSONDecodeError, OSError):
        return {}
    return data if isinstance(data, dict) else {}


def match_env_site_dict(base_url: str, env: str = "") -> dict[str, Any] | None:
    data = load_raw_env_sites()
    if not data:
        return None
    wanted_env = (env or "").strip().lower()
    if wanted_env in ENVS and isinstance(data.get(wanted_env), dict):
        return data[wanted_env]
    target = normalize_na_base_url(base_url)
    if target:
        for key in ENVS:
            raw = data.get(key)
            if not isinstance(raw, dict):
                continue
            if normalize_na_base_url(str(raw.get("base_url") or "")) == target:
                return raw
    guessed = detect_env(base_url, wanted_env)
    raw = data.get(guessed)
    return raw if isinstance(raw, dict) else None


def _blank(value: Any) -> bool:
    return value is None or str(value).strip() == ""


def enrich_site_from_env(site: SiteConfig) -> SiteConfig:
    """sites.json 预置的账号密码 / totp_secret 在请求未带时自动补上。"""
    raw = match_env_site_dict(site.base_url, site.env)
    if not raw:
        return site
    if _blank(site.username) and raw.get("username"):
        site.username = str(raw.get("username") or "").strip()
    if _blank(site.password) and raw.get("password") not in (None, ""):
        site.password = str(raw.get("password") or "")
    if _blank(site.access_token) and raw.get("access_token"):
        site.access_token = str(raw.get("access_token") or "").strip()
    if site.user_id in (None, 0) and raw.get("user_id") not in (None, ""):
        try:
            site.user_id = int(raw.get("user_id"))
        except (TypeError, ValueError):
            pass
    if not str(site.totp_secret or "").strip():
        site.totp_secret = pick_totp_secret(mapping=raw)
    if not str(site.otp_code or "").strip():
        site.otp_code = str(raw.get("otp_code") or raw.get("verification_code") or "").strip()
    return site


def enrich_auth_from_env(auth: dict[str, Any], *, base_url: str = "", env: str = "") -> dict[str, Any]:
    """页面未填的账号密码 / Token / TOTP，按地址从 env/sites.json 补上。"""
    out = dict(auth or {})
    raw = match_env_site_dict(base_url or str(out.get("base_url") or ""), env or str(out.get("env") or ""))
    if raw:
        if _blank(out.get("username")) and raw.get("username"):
            out["username"] = str(raw.get("username") or "").strip()
        if _blank(out.get("password")) and raw.get("password") not in (None, ""):
            out["password"] = raw.get("password")
        if _blank(out.get("access_token")) and _blank(out.get("password")) and raw.get("access_token"):
            out["access_token"] = str(raw.get("access_token") or "").strip()
        if out.get("user_id") in (None, "", 0) and raw.get("user_id") not in (None, ""):
            out["user_id"] = raw.get("user_id")
        if _blank(out.get("default_key")) and raw.get("default_key"):
            out["default_key"] = str(raw.get("default_key") or "").strip()
    return enrich_auth_totp(out, base_url=base_url or str(out.get("base_url") or ""), env=env or str(out.get("env") or ""))


def enrich_auth_totp(auth: dict[str, Any], *, base_url: str = "", env: str = "") -> dict[str, Any]:
    """登录请求未带 TOTP 密钥时，按地址从 env/sites.json 补上。"""
    out = dict(auth or {})
    if pick_totp_secret(out.get("totp_secret"), out.get("otp_secret"), out.get("otp_code"), out.get("verification_code")):
        if not str(out.get("totp_secret") or "").strip():
            out["totp_secret"] = pick_totp_secret(
                out.get("otp_secret"),
                out.get("otp_code"),
                out.get("verification_code"),
            )
        return out
    raw = match_env_site_dict(base_url, env)
    if not raw:
        return out
    secret = pick_totp_secret(mapping=raw)
    if secret:
        out["totp_secret"] = secret
    if not str(out.get("otp_code") or "").strip():
        env_otp = str(raw.get("otp_code") or raw.get("verification_code") or "").strip()
        if env_otp:
            out["otp_code"] = env_otp
    return out


def site_has_otp_source(site: SiteConfig) -> bool:
    return has_otp_source(site.otp_code, totp_secret=site.totp_secret)


def _format_channel_status(status: Any) -> str:
    return "已启用" if int(status or 0) == 1 else "已禁用"


def format_channel_response_time(value: Any) -> str:
    """NewAPI 的 response_time 为毫秒；0 或缺失表示未测试。"""
    try:
        milliseconds = int(value)
    except (TypeError, ValueError):
        return "未测试"
    if milliseconds <= 0:
        return "未测试"
    return f"{milliseconds / 1000.0:.2f} 秒"


def _format_token_status(status: Any) -> str:
    return "已启用" if int(status or 0) == 1 else "已禁用"


def _format_timestamp(value: Any) -> str:
    try:
        ts = int(value)
    except (TypeError, ValueError):
        return ""
    if ts <= 0:
        return ""
    return datetime.fromtimestamp(ts).strftime("%Y-%m-%d %H:%M:%S")


def _format_expire_time(value: Any) -> str:
    try:
        ts = int(value)
    except (TypeError, ValueError):
        return ""
    if ts < 0:
        return "永不过期"
    if ts == 0:
        return ""
    return datetime.fromtimestamp(ts).strftime("%Y-%m-%d %H:%M:%S")


def _format_remain_quota(token: dict[str, Any]) -> str:
    if token.get("unlimited_quota"):
        return "无限额度"
    remain = token.get("remain_quota")
    used = token.get("used_quota")
    try:
        remain_int = int(remain)
    except (TypeError, ValueError):
        remain_int = None
    try:
        used_int = int(used)
    except (TypeError, ValueError):
        used_int = None
    if remain_int is not None and used_int is not None:
        total = remain_int + used_int
        return f"{remain_int}/{total}"
    if remain_int is not None:
        return str(remain_int)
    return str(remain or "")


def _format_model_mapping(raw: Any) -> str:
    if not raw:
        return ""
    try:
        mapping = json.loads(raw) if isinstance(raw, str) else raw
    except (json.JSONDecodeError, TypeError):
        return str(raw)
    if not isinstance(mapping, dict):
        return ""
    parts: list[str] = []
    for src, dst in mapping.items():
        src_text = str(src).strip()
        dst_text = str(dst).strip()
        if src_text and dst_text:
            parts.append(f"{src_text}->{dst_text}")
    return ", ".join(parts)


def _format_models_list(raw: Any) -> str:
    text = str(raw or "").strip()
    if not text:
        return ""
    parts = [part.strip() for part in text.split(",") if part.strip()]
    return ", ".join(parts)


def _channel_names_for_token(token: dict[str, Any], channels: list[dict[str, Any]]) -> str:
    _enabled, limits = read_model_limits(token)
    if not limits:
        return "无限制"
    models = {part.strip() for part in limits.split(",") if part.strip()}
    if not models:
        return "无限制"

    group = str(token.get("group") or "").strip()
    matched: list[str] = []
    for channel in channels:
        ch_group = str(channel.get("group") or "").strip()
        if group and group not in ch_group.split(","):
            continue
        name = str(channel.get("name") or "").strip()
        if not name:
            continue
        ch_models = str(channel.get("models") or "")
        if any(model in ch_models for model in models) or name in models:
            matched.append(name)
    if matched:
        return ", ".join(sorted(set(matched)))
    return ", ".join(sorted(models))


def _format_available_models(token: dict[str, Any]) -> str:
    _enabled, limits = read_model_limits(token)
    if not limits:
        return "无限制"
    parts = [part.strip() for part in limits.split(",") if part.strip()]
    return ", ".join(parts) if parts else "无限制"


def _format_ip_limit(raw: Any) -> str:
    if raw is None:
        return ""
    text = str(raw).strip()
    return text if text else "无限制"


def normalize_token_row(row: dict[str, Any] | None) -> dict[str, str]:
    """补齐令牌行全部列，空值保留为空字符串（便于 Excel 留空）。"""
    source = row or {}
    normalized: dict[str, str] = {}
    for header in TOKEN_HEADERS:
        value = source.get(header, "")
        normalized[header] = "" if value is None else str(value)
    return normalized


def _enrich_token_for_summary(client: NewAPIClient, token: dict[str, Any]) -> dict[str, Any]:
    """列表接口字段不全时，按需拉取详情补全。"""
    record = dict(token)
    token_id = record.get("id")
    if token_id is None:
        return record

    _enabled, limits = read_model_limits(record)
    need_detail = (
        not limits
        or record.get("created_time") in (None, "", 0)
        or record.get("allow_ips") is None
        or "expired_time" not in record
    )
    if not need_detail:
        return record

    try:
        detail = client.get_token(int(token_id))
    except Exception:  # noqa: BLE001
        return record

    for field in (
        "created_time",
        "expired_time",
        "allow_ips",
        "remain_quota",
        "used_quota",
        "unlimited_quota",
        "group",
        "status",
        "model_limits",
        "model_limits_enabled",
        "token_model_limit",
        "token_model_limit_enabled",
    ):
        if field not in detail:
            continue
        current = record.get(field)
        if field in {"model_limits", "token_model_limit"}:
            if not read_model_limits(record)[1] and detail.get(field):
                record[field] = detail[field]
            continue
        if current in (None, "", 0) and detail.get(field) not in (None, ""):
            record[field] = detail[field]
    return record


def channel_to_row(
    channel: dict[str, Any],
    type_labels: dict[int, str],
    *,
    api_key: str | None = None,
) -> dict[str, Any]:
    type_id = int(channel.get("type") or 0)
    channel_id = int(channel.get("id") or 0)
    name = str(channel.get("name") or "").strip()
    if not name and channel_id:
        name = f"(未命名#{channel_id})"
    row = {
        "名称": name,
        "分组": str(channel.get("group") or "").strip(),
        "类型": type_labels.get(type_id, str(type_id)),
        "状态": _format_channel_status(channel.get("status")),
        "响应时间": format_channel_response_time(channel.get("response_time")),
        "模型列表": _format_models_list(channel.get("models")),
        "模型重定向": _format_model_mapping(channel.get("model_mapping")),
        "base_url": str(channel.get("base_url") or "").strip(),
        "备注": str(channel.get("remark") or "").strip(),
        "部署时间": _format_timestamp(channel.get("created_time")),
        "_sort_id": channel_id,
    }
    if api_key is not None:
        row["密钥"] = str(api_key or "").strip()
    return row


def token_to_row(
    token: dict[str, Any],
    channels: list[dict[str, Any]],
) -> dict[str, str]:
    group = str(token.get("group") or "").strip()
    token_id = int(token.get("id") or 0)
    name = str(token.get("name") or "").strip()
    if not name and token_id:
        name = f"(未命名#{token_id})"
    row = normalize_token_row(
        {
            "名称": name,
            "项目环境": group,
            "状态": _format_token_status(token.get("status")),
            "剩余额度": _format_remain_quota(token),
            "分组": group,
            "可用模型": _format_available_models(token),
            "包含渠道": _channel_names_for_token(token, channels),
            "IP限制": _format_ip_limit(token.get("allow_ips")),
            "创建时间": _format_timestamp(token.get("created_time")),
            "过期时间": _format_expire_time(token.get("expired_time")),
        }
    )
    if token_id:
        row["_sort_id"] = token_id
    return row


# 上游 CriticalRateLimit：登录、/api/verify、/key 共用；次数以服务端 429 为准
_KEY_FETCH_INTERVAL_SEC = 1.0


def _site_auth_payload(site: SiteConfig, *, otp: str = "") -> dict[str, Any]:
    token = str(site.access_token or "").strip()
    if token:
        uid = int(site.user_id or 0)
        if not uid:
            raise ValueError("使用 Access Token 时需填写 user_id")
        return {
            "access_token": token,
            "user_id": uid,
            "otp_code": otp,
            "totp_secret": str(site.totp_secret or "").strip(),
            "skip_session_cache": False,
            "verify_ssl": bool(site.verify_ssl),
        }
    if not site.username:
        raise ValueError("站点缺少用户名（或配置 Access Token + user_id）")
    if not str(site.password or ""):
        raise ValueError("站点缺少密码（或配置 Access Token + user_id）")
    return {
        "username": site.username,
        "password": site.password,
        "otp_code": otp,
        "totp_secret": str(site.totp_secret or "").strip(),
        "skip_session_cache": False,
        "verify_ssl": bool(site.verify_ssl),
    }


def _fetch_channel_keys(
    client: NewAPIClient,
    channels_raw: list[dict[str, Any]],
    *,
    otp_code: str,
    totp_secret: str = "",
) -> tuple[dict[int, str], list[str], dict[str, Any]]:
    """强制 /api/verify 后逐个拉密钥，直到全部完成或上游 429。"""
    import time as _time

    code = resolve_otp_code(otp_code, totp_secret=totp_secret)
    if not code:
        raise ValueError("导出渠道密钥需提供二次验证码（认证器 6 位、备用码，或 totp_secret）")
    if not client.is_root():
        raise RuntimeError(
            f"获取渠道密钥需要 Root 权限（当前 role={client.role}）。"
            "请使用 Root 账号的用户名密码或 Access Token 采集。"
        )

    targets: list[int] = []
    for item in channels_raw:
        try:
            cid = int(item.get("id") or 0)
        except (TypeError, ValueError):
            continue
        if cid > 0:
            targets.append(cid)

    total = len(targets)
    warnings: list[str] = []
    empty_meta = {
        "keys_total": 0,
        "keys_pending": 0,
        "stopped_early": False,
        "stop_reason": "done",
        "can_retry_soon": False,
        "wait_minutes": 0,
        "quota_remaining": 0,
    }
    if total == 0:
        return {}, warnings, empty_meta

    # 刚收到 429 则先等窗口，避免再空耗 login/verify
    assert_can_proceed(client.base_url, need=1)

    # 密码登录：去掉 Authorization，让 verify/key 都落在同一 cookie session（官方前端同款）
    client.prefer_cookie_session()
    client.ensure_secure_verified(code, force=True)

    keys: dict[int, str] = {}
    stopped_early = False
    stop_reason = "done"
    reverified_once = False
    for index, cid in enumerate(targets):
        if index > 0:
            _time.sleep(_KEY_FETCH_INTERVAL_SEC)
        try:
            keys[cid] = client.fetch_channel_key(cid)
        except SecureVerificationError as exc:
            if not reverified_once:
                reverified_once = True
                warnings.append(f"拉密钥遇到验证失效，正在用同一验证码强制重验后重试…（{exc}）")
                try:
                    client.prefer_cookie_session()
                    client.ensure_secure_verified(code, force=True)
                    keys[cid] = client.fetch_channel_key(cid)
                    continue
                except Exception as retry_exc:  # noqa: BLE001
                    stop_reason = "verify_expired"
                    warnings.append(
                        f"强制重验后仍失败（已获取 {len(keys)}/{total}）：{retry_exc}。"
                        "请换一个新的 6 位验证码后立即再点「继续补全密钥」。"
                    )
                    stopped_early = True
                    break
            stop_reason = "verify_expired"
            warnings.append(
                f"安全验证未生效，已获取 {len(keys)}/{total}。上游：{exc}。"
                "请换新验证码后立即再试。"
            )
            stopped_early = True
            break
        except Exception as exc:  # noqa: BLE001
            text = str(exc)
            is_rate = "429" in text or "限流" in text or "Too Many Requests" in text
            if is_rate:
                note_rate_limited(client.base_url)
                stop_reason = "rate_limit"
                wait = seconds_until_reset(client.base_url)
                mins = max(1, (wait + 59) // 60)
                warnings.append(
                    f"上游返回 429 限流，已停止继续拉取。"
                    f"本批已获取 {len(keys)}/{total}，剩余约 {total - len(keys)}。"
                    f"请等待约 {mins} 分钟后再「继续补全密钥」。"
                )
                stopped_early = True
                break
            warnings.append(f"渠道#{cid} 密钥失败: {exc}")

    pending_after = max(0, total - len(keys))
    rem_after = remaining(client.base_url)
    can_retry_soon = stop_reason == "verify_expired"
    wait_minutes = 0
    if stop_reason == "rate_limit" and pending_after > 0:
        wait_minutes = max(1, (seconds_until_reset(client.base_url) + 59) // 60)

    meta = {
        "keys_total": total,
        "keys_pending": pending_after,
        "stopped_early": stopped_early or pending_after > 0,
        "stop_reason": stop_reason,
        "can_retry_soon": can_retry_soon,
        "wait_minutes": wait_minutes,
        "quota_remaining": rem_after,
    }
    return keys, warnings, meta



def _authenticate_for_key_export(client: NewAPIClient, site: SiteConfig, otp: str) -> dict[str, Any]:
    """登录尽量不消耗验证码；验证码留给强制 /api/verify。

    补密钥时跳过磁盘会话缓存，避免沿用带过期 secure_verified_at 的旧 cookie
    （上游会直接返回 VERIFICATION_EXPIRED）。
    """
    auth_no_otp = {
        **_site_auth_payload(site, otp=""),
        "skip_session_cache": True,
    }
    try:
        authenticate_client(client, auth_no_otp)
        return auth_no_otp
    except RuntimeError as exc:
        text = str(exc)
        needs_login_2fa = "两步验证" in text or "2fa" in text.lower() or "2FA" in text
        if not needs_login_2fa or not site_has_otp_source(site):
            raise
        auth_with_otp = {
            **_site_auth_payload(site, otp=otp),
            "skip_session_cache": True,
        }
        authenticate_client(client, auth_with_otp)
        return auth_with_otp


def fetch_site_data(
    site: SiteConfig,
    *,
    type_labels: dict[int, str],
    include_keys: bool = False,
) -> tuple[str, list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    base_url = normalize_na_base_url(site.base_url)
    if not base_url:
        raise ValueError(f"站点地址无效: {site.base_url}")

    env = detect_env(base_url, site.env)
    client = NewAPIClient(base_url, verify_ssl=bool(site.verify_ssl))
    otp = str(site.otp_code or "").strip()
    if include_keys:
        auth = _authenticate_for_key_export(client, site, otp)
    else:
        auth = _site_auth_payload(site, otp=otp)
        authenticate_client(client, auth)

    channels_raw = client.list_all_channels()
    tokens_raw = client.list_all_tokens()
    tokens_raw = [_enrich_token_for_summary(client, item) for item in tokens_raw]

    key_map: dict[int, str] = {}
    key_warnings: list[str] = []
    key_meta: dict[str, Any] = {"keys_total": 0, "keys_pending": 0, "stopped_early": False}
    if include_keys:
        key_map, key_warnings, key_meta = _fetch_channel_keys(
            client, channels_raw, otp_code=otp, totp_secret=str(site.totp_secret or "")
        )
        refresh_auth_session_cache(client, auth)

    channels = []
    for item in channels_raw:
        try:
            cid = int(item.get("id") or 0)
        except (TypeError, ValueError):
            cid = 0
        channels.append(
            channel_to_row(
                item,
                type_labels,
                api_key=(key_map.get(cid, "") if include_keys else None),
            )
        )
    tokens = [token_to_row(item, channels_raw) for item in tokens_raw]
    channels = [row for row in channels if row.get("名称") or row.get("_sort_id")]
    tokens = [row for row in tokens if row.get("名称") or row.get("_sort_id")]
    meta = {
        "keys_fetched": len(key_map),
        "keys_failed": max(0, int(key_meta.get("keys_total") or 0) - len(key_map)) if include_keys else 0,
        "keys_total": int(key_meta.get("keys_total") or 0),
        "keys_pending": int(key_meta.get("keys_pending") or 0),
        "stopped_early": bool(key_meta.get("stopped_early")),
        "key_warnings": key_warnings,
    }
    return env, channels, tokens, meta


def _read_channel_sheet_with_keys(path: Path, env: str) -> tuple[list[str], list[dict[str, Any]]]:
    from openpyxl import load_workbook as _load_wb

    title = SHEET_CHANNELS[env]
    wb = _load_wb(path)
    if title not in wb.sheetnames:
        wb.close()
        raise ValueError(f"Excel 中无工作表 {title}")
    ws = wb[title]
    header_cells = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    headers = [str(h) for h in header_cells if h is not None]
    rows: list[dict[str, Any]] = []
    for row_cells in ws.iter_rows(min_row=2, values_only=True):
        if not row_cells or not any(row_cells):
            continue
        item = {
            headers[i]: ("" if (i >= len(row_cells) or row_cells[i] is None) else row_cells[i])
            for i in range(len(headers))
        }
        if str(item.get("名称") or "").strip():
            rows.append(item)
    wb.close()
    return headers, rows


def _try_read_channel_sheet_with_keys(path: Path, env: str) -> tuple[list[str], list[dict[str, Any]]]:
    try:
        return _read_channel_sheet_with_keys(path, env)
    except ValueError:
        return list(CHANNEL_HEADERS_WITH_KEY), []


def scan_channel_key_status(path: Path | None) -> dict[str, Any]:
    """统计各环境渠道密钥已获取 / 仍缺，供③⑥ 日志与「继续补全」按钮。"""
    empty_env = {
        "channels": 0,
        "keys_filled": 0,
        "keys_missing": 0,
        "has_key_column": False,
        "complete": False,
        "empty": True,
    }
    if not path or not Path(path).exists():
        return {
            "exists": False,
            "envs": {env: dict(empty_env) for env in ENVS},
            "pending_total": 0,
            "filled_total": 0,
            "channels_total": 0,
            "complete": False,
            "has_key_column": False,
            "can_continue": False,
        }
    src = Path(path)
    envs: dict[str, dict[str, Any]] = {}
    pending_total = 0
    filled_total = 0
    channels_total = 0
    has_key_column = False
    for env in ENVS:
        headers, rows = _try_read_channel_sheet_with_keys(src, env)
        named = [row for row in rows if str(row.get("名称") or "").strip()]
        has_col = "密钥" in headers
        if has_col:
            has_key_column = True
        filled = sum(1 for row in named if str(row.get("密钥") or "").strip()) if has_col else 0
        missing = max(0, len(named) - filled)
        envs[env] = {
            "channels": len(named),
            "keys_filled": filled,
            "keys_missing": missing,
            "has_key_column": has_col,
            "complete": bool(named) and missing == 0,
            "empty": not named,
        }
        pending_total += missing
        filled_total += filled
        channels_total += len(named)
    return {
        "exists": True,
        "envs": envs,
        "pending_total": pending_total,
        "filled_total": filled_total,
        "channels_total": channels_total,
        "complete": channels_total > 0 and pending_total == 0,
        "has_key_column": has_key_column,
        "can_continue": pending_total > 0,
    }


def _load_all_env_rows_from_excel(
    src: Path,
) -> tuple[dict[str, list[dict[str, Any]]], dict[str, list[dict[str, Any]]]]:
    """读出各环境渠道（尽量保留密钥列）与令牌，供写回同一份 Excel。"""
    channel_env_rows: dict[str, list[dict[str, Any]]] = {e: [] for e in ENVS}
    token_env_rows: dict[str, list[dict[str, Any]]] = {e: [] for e in ENVS}
    loaded_ch, loaded_tk = load_workbook_data(src)
    for e in ENVS:
        token_env_rows[e] = loaded_tk.get(e, [])
        headers, env_rows = _try_read_channel_sheet_with_keys(src, e)
        for row in env_rows:
            row.setdefault("密钥", "")
        if env_rows:
            channel_env_rows[e] = env_rows
        else:
            fallback = loaded_ch.get(e, [])
            for row in fallback:
                row.setdefault("密钥", "")
            channel_env_rows[e] = fallback
    return channel_env_rows, token_env_rows


def _save_summary_workbook_with_keys(
    path: Path,
    channel_env_rows: dict[str, list[dict[str, Any]]],
    token_env_rows: dict[str, list[dict[str, Any]]],
) -> None:
    save_workbook_data(
        path,
        channel_env_rows=channel_env_rows,
        token_env_rows=token_env_rows,
        channel_dup_source=channel_env_rows,
        token_dup_source=token_env_rows,
        channel_headers=CHANNEL_HEADERS_WITH_KEY,
        dup_channel_headers=DUP_CHANNEL_HEADERS_WITH_KEY,
    )


def _bootstrap_empty_env_sheet(
    *,
    site: SiteConfig,
    src: Path,
    env: str,
) -> tuple[dict[str, Any], Path]:
    """该环境渠道表为空时：从站点拉取渠道+令牌+密钥，写入已有工作簿，不覆盖其它环境。"""
    type_labels = _build_type_label_map()
    fetched_env, channels, tokens, meta = fetch_site_data(
        site,
        type_labels=type_labels,
        include_keys=True,
    )
    env = fetched_env or env
    channel_env_rows, token_env_rows = _load_all_env_rows_from_excel(src)
    channel_env_rows[env] = _sort_rows_by_id(list(channels))
    token_env_rows[env] = _sort_rows_by_id(list(tokens))
    _save_summary_workbook_with_keys(src, channel_env_rows, token_env_rows)

    filled = int(meta.get("keys_fetched") or 0)
    pending = int(meta.get("keys_pending") or 0)
    warnings = list(meta.get("key_warnings") or [])
    msg = (
        f"[{env}] 工作表原为空，已从站点写入渠道 {len(channels)}、令牌 {len(tokens)}，"
        f"补全密钥 {filled}"
    )
    if pending:
        msg += f"，仍缺 {pending}"
    return {
        "env": env,
        "keys_fetched": filled,
        "keys_pending": pending,
        "keys_total_missing": pending,
        "bootstrapped": True,
        "channels_written": len(channels),
        "tokens_written": len(tokens),
        "warnings": warnings,
        "message": msg,
    }, src


def fill_missing_channel_keys(
    *,
    site: SiteConfig,
    root_dir: Path,
    excel_path: Path | None = None,
) -> tuple[dict[str, Any], Path]:
    """对指定 Excel（默认③最新网关汇总）补全某环境渠道密钥。

    该环境渠道表已有数据时，只补空密钥行。
    表为空（或尚无该环境工作表）时，从站点拉取渠道、令牌和密钥，写入已有工作簿，不覆盖其它环境。
    ⑥ 传入总览文件路径时，只改渠道/令牌表，保留火山云 / Azure / 总览目录。
    """
    otp = str(site.otp_code or "").strip()
    if not site_has_otp_source(site):
        raise ValueError("继续补全密钥须填写二次验证码，或在 env/sites.json 预置 totp_secret")
    # 尽早校验账号，避免无 Excel 时才暴露缺密码
    _site_auth_payload(site, otp="")

    src = Path(excel_path) if excel_path else latest_excel_path(root_dir)
    if src is None or not src.exists():
        src = build_timestamped_path(root_dir)
    if not src.exists():
        raise FileNotFoundError("未找到 Excel，请先「一键采集」或「一键生成总览表」后再补全密钥")

    env = detect_env(normalize_na_base_url(site.base_url), site.env)
    headers, rows = _try_read_channel_sheet_with_keys(src, env)

    # 旧表无「密钥」列时自动补空列，便于在已有渠道数据上补拉
    if "密钥" not in headers:
        headers = list(CHANNEL_HEADERS_WITH_KEY)
        for row in rows:
            row.setdefault("密钥", "")

    for row in rows:
        row.setdefault("密钥", "")

    if not rows:
        return _bootstrap_empty_env_sheet(site=site, src=src, env=env)

    missing_names = [
        str(r.get("名称") or "").strip()
        for r in rows
        if str(r.get("名称") or "").strip() and not str(r.get("密钥") or "").strip()
    ]
    if not missing_names:
        return {
            "env": env,
            "keys_fetched": 0,
            "keys_pending": 0,
            "message": f"[{env}] 没有缺失的密钥",
            "warnings": [],
        }, src

    client = NewAPIClient(normalize_na_base_url(site.base_url), verify_ssl=bool(site.verify_ssl))
    auth = _authenticate_for_key_export(client, site, otp)
    # 认证成功后把 cookie/token 会话再落盘，供下一轮少登录
    channels_raw = client.list_all_channels()
    name_to_ch = {
        str(ch.get("name") or "").strip(): ch
        for ch in channels_raw
        if str(ch.get("name") or "").strip()
    }
    subset = [name_to_ch[n] for n in missing_names if n in name_to_ch]
    unmatched = [n for n in missing_names if n not in name_to_ch]
    warnings: list[str] = []
    if unmatched:
        sample = "、".join(unmatched[:5])
        more = f" 等 {len(unmatched)} 个" if len(unmatched) > 5 else ""
        warnings.append(f"Excel 中有 {len(unmatched)} 个渠道名称在站点上未找到：{sample}{more}")
    if not subset:
        raise ValueError(
            f"[{env}] 有 {len(missing_names)} 条空密钥，但无法在站点匹配到同名渠道。"
            "请先对该环境重新「一键采集」同步渠道名称后再补全。"
        )

    key_map, fetch_warnings, key_meta = _fetch_channel_keys(
        client, subset, otp_code=otp, totp_secret=str(site.totp_secret or "")
    )
    refresh_auth_session_cache(client, auth)
    warnings.extend(fetch_warnings)
    id_to_name = {
        int(ch.get("id") or 0): str(ch.get("name") or "").strip()
        for ch in subset
        if int(ch.get("id") or 0) > 0
    }
    name_to_key = {
        id_to_name[cid]: key for cid, key in key_map.items() if id_to_name.get(cid)
    }

    filled = 0
    for row in rows:
        name = str(row.get("名称") or "").strip()
        if name in name_to_key and not str(row.get("密钥") or "").strip():
            row["密钥"] = name_to_key[name]
            filled += 1

    # 读出全部环境（尽量保留已有密钥列）
    channel_env_rows, token_env_rows = _load_all_env_rows_from_excel(src)
    channel_env_rows[env] = rows

    # 补全密钥写回同一份最新 Excel，避免每次新建导致密钥分散在多份文件里
    save_path = src
    _save_summary_workbook_with_keys(save_path, channel_env_rows, token_env_rows)
    pending = max(0, len(missing_names) - filled)
    stopped = bool(key_meta.get("stopped_early"))
    stop_reason = str(key_meta.get("stop_reason") or "")
    wait_minutes = int(key_meta.get("wait_minutes") or 0)
    can_retry_soon = bool(key_meta.get("can_retry_soon"))
    msg = f"[{env}] 补全密钥 {filled}/{len(missing_names)}"
    if pending:
        if stop_reason == "verify_expired" or can_retry_soon:
            msg += f"，仍缺 {pending}（可立即换新验证码再续，不必等 20 分钟）"
        elif stop_reason in {"rate_limit", "budget_exhausted"} and wait_minutes > 0:
            msg += f"，仍缺 {pending}（请约 {wait_minutes} 分钟后再续）"
        elif stopped or filled:
            rem = int(key_meta.get("quota_remaining") or 0)
            if rem >= 1:
                msg += f"，仍缺 {pending}（配额仍有剩余，可立即再点继续补全）"
            else:
                msg += f"，仍缺 {pending}（请约 20 分钟后再续）"
        else:
            msg += f"，仍缺 {pending}"
    return {
        "env": env,
        "keys_fetched": filled,
        "keys_pending": pending,
        "keys_total_missing": len(missing_names),
        "stop_reason": stop_reason,
        "can_retry_soon": can_retry_soon,
        "wait_minutes": wait_minutes,
        "warnings": warnings,
        "message": msg,
    }, save_path


def _duplicate_name_key(row: dict[str, Any]) -> str:
    """跨环境重复检测：按名称。"""
    return str(row.get("名称") or "").strip()


def _merge_row_key(row: dict[str, Any]) -> str:
    """同环境合并：按 ID 唯一，避免同名渠道被覆盖。"""
    sort_id = row.get("_sort_id")
    try:
        cid = int(sort_id)
    except (TypeError, ValueError):
        cid = 0
    if cid > 0:
        return f"id:{cid}"
    name = str(row.get("名称") or "").strip()
    return f"name:{name}" if name else ""


def _sort_rows_by_id(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return sorted(rows, key=lambda row: int(row.get("_sort_id") or 0))


def _merge_rows(existing: list[dict[str, Any]], incoming: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], int]:
    order: list[str] = []
    index: dict[str, dict[str, Any]] = {}
    added = 0
    for row in existing:
        key = _merge_row_key(row)
        if not key:
            continue
        if key not in index:
            order.append(key)
        index[key] = row
    for row in incoming:
        key = _merge_row_key(row)
        if not key:
            continue
        if key not in index:
            order.append(key)
            added += 1
        index[key] = row
    merged = [index[key] for key in order]
    return _sort_rows_by_id(merged), added


def _cross_env_duplicate_names(env_rows: dict[str, list[dict[str, Any]]]) -> set[str]:
    """跨 test/uat/prod 同名记录。"""
    name_envs: dict[str, set[str]] = {}
    for env, rows in env_rows.items():
        for row in rows:
            name = _duplicate_name_key(row)
            if not name:
                continue
            name_envs.setdefault(name, set()).add(env)
    return {name for name, envs in name_envs.items() if len(envs) >= 2}


def _exclude_duplicate_names(
    env_rows: dict[str, list[dict[str, Any]]],
    duplicate_names: set[str],
) -> tuple[dict[str, list[dict[str, Any]]], int]:
    """从各环境主表中剔除跨环境重复项。"""
    if not duplicate_names:
        return env_rows, 0

    filtered: dict[str, list[dict[str, Any]]] = {env: [] for env in ENVS}
    skipped = 0
    for env in ENVS:
        for row in env_rows.get(env, []):
            if _duplicate_name_key(row) in duplicate_names:
                skipped += 1
                continue
            filtered[env].append(row)
    return filtered, skipped


def build_duplicate_channels(env_rows: dict[str, list[dict[str, Any]]]) -> list[dict[str, Any]]:
    name_to_sources: dict[str, list[str]] = {}
    name_to_row: dict[str, dict[str, Any]] = {}
    for env in ENVS:
        sheet = SHEET_CHANNELS[env]
        for row in env_rows.get(env, []):
            name = _duplicate_name_key(row)
            if not name:
                continue
            name_to_sources.setdefault(name, []).append(sheet)
            name_to_row.setdefault(name, row)

    duplicates: list[dict[str, Any]] = []
    serial = 1
    for name in sorted(name_to_sources.keys()):
        sources = sorted(set(name_to_sources[name]))
        if len(sources) < 2:
            continue
        base = dict(name_to_row[name])
        for source in sources:
            item = dict(base)
            item["来源"] = source
            item["重复序号"] = serial
            duplicates.append(item)
        serial += 1
    return duplicates


def build_duplicate_tokens(env_rows: dict[str, list[dict[str, Any]]]) -> list[dict[str, Any]]:
    name_to_sources: dict[str, list[str]] = {}
    name_to_row: dict[str, dict[str, Any]] = {}
    for env in ENVS:
        sheet = SHEET_TOKENS[env]
        for row in env_rows.get(env, []):
            name = _duplicate_name_key(row)
            if not name:
                continue
            name_to_sources.setdefault(name, []).append(sheet)
            name_to_row.setdefault(name, row)

    duplicates: list[dict[str, Any]] = []
    serial = 1
    for name in sorted(name_to_sources.keys()):
        sources = sorted(set(name_to_sources[name]))
        if len(sources) < 2:
            continue
        base = normalize_token_row(name_to_row[name])
        for source in sources:
            item = dict(base)
            item["来源"] = source
            item["重复序号"] = serial
            duplicates.append(item)
        serial += 1
    return duplicates


def _skipped_by_env(
    env_rows: dict[str, list[dict[str, Any]]],
    duplicate_names: set[str],
) -> dict[str, int]:
    counts: dict[str, int] = {env: 0 for env in ENVS}
    for env in ENVS:
        for row in env_rows.get(env, []):
            if _duplicate_name_key(row) in duplicate_names:
                counts[env] += 1
    return counts


def run_data_summary_collect(
    *,
    sites: list[SiteConfig],
    excel_path: Path,
    include_keys: bool = False,
    output_path: Path | None = None,
    staging: bool = False,
) -> tuple[CollectStats, Path]:
    stats = CollectStats(include_keys=bool(include_keys))
    type_labels = _build_type_label_map()
    root_dir = excel_path.parent.parent

    channel_env_rows: dict[str, list[dict[str, Any]]] = {env: [] for env in ENVS}
    token_env_rows: dict[str, list[dict[str, Any]]] = {env: [] for env in ENVS}
    touched_envs: set[str] = set()

    # 一键采集始终新建 Excel；同一次点击多站时通过 output_path 写入同一新文件
    session_path = output_path if (output_path and output_path.exists()) else None
    if session_path:
        try:
            for env in ENVS:
                try:
                    _hdr, rows = _read_channel_sheet_with_keys(session_path, env)
                    for row in rows:
                        row.setdefault("密钥", "")
                    channel_env_rows[env] = rows
                except ValueError:
                    pass
            prev_ch, prev_tk = load_workbook_data(session_path)
            for env in ENVS:
                if not channel_env_rows[env]:
                    channel_env_rows[env] = prev_ch.get(env, [])
                token_env_rows[env] = list(prev_tk.get(env, []))
        except Exception:  # noqa: BLE001
            pass

    for site in sites:
        label = site.label or site.base_url
        try:
            if include_keys and not site_has_otp_source(site):
                raise ValueError("含密钥采集须为该站点填写二次验证码，或在 env/sites.json 预置 totp_secret")
            env, channels, tokens, meta = fetch_site_data(
                site,
                type_labels=type_labels,
                include_keys=include_keys,
            )
            resolved_url = normalize_na_base_url(site.base_url)
            touched_envs.add(env)

            # 本站结果覆盖该环境（不与历史同环境合并，避免脏数据）
            channel_env_rows[env] = list(channels)
            token_env_rows[env] = list(tokens)
            ch_added = len(channels)
            tk_added = len(tokens)
            stats.sites_ok += 1
            stats.channels_added[env] = stats.channels_added.get(env, 0) + ch_added
            stats.tokens_added[env] = stats.tokens_added.get(env, 0) + tk_added
            stats.keys_fetched += int(meta.get("keys_fetched") or 0)
            stats.keys_failed += int(meta.get("keys_failed") or 0)
            for warn in meta.get("key_warnings") or []:
                stats.errors.append(f"{label}: {warn}")
            stats.site_results.append(
                {
                    "label": label,
                    "base_url": resolved_url,
                    "env": env,
                    "env_explicit": bool((site.env or "").strip()),
                    "channels_fetched": len(channels),
                    "tokens_fetched": len(tokens),
                    "channels_in_env": len(channel_env_rows[env]),
                    "tokens_in_env": len(token_env_rows[env]),
                    "keys_fetched": int(meta.get("keys_fetched") or 0),
                    "keys_failed": int(meta.get("keys_failed") or 0),
                    "keys_pending": int(meta.get("keys_pending") or 0),
                    "stopped_early": bool(meta.get("stopped_early")),
                }
            )
        except Exception as exc:  # noqa: BLE001
            stats.sites_failed += 1
            stats.errors.append(f"{label}: {exc}")

    dup_channels = build_duplicate_channels(channel_env_rows)
    dup_tokens = build_duplicate_tokens(token_env_rows)
    stats.duplicate_channels = len({row.get("名称") for row in dup_channels})
    stats.duplicate_tokens = len({row.get("名称") for row in dup_tokens})

    ch_dup_names = _cross_env_duplicate_names(channel_env_rows)
    tk_dup_names = _cross_env_duplicate_names(token_env_rows)
    stats.channels_skipped_duplicate = 0
    stats.tokens_skipped_duplicate = 0
    stats.channels_skipped_by_env = _skipped_by_env(channel_env_rows, ch_dup_names)
    stats.tokens_skipped_by_env = _skipped_by_env(token_env_rows, tk_dup_names)
    stats.channels_in_sheet = {env: len(channel_env_rows.get(env, [])) for env in ENVS}
    stats.tokens_in_sheet = {env: len(token_env_rows.get(env, [])) for env in ENVS}

    for env in ENVS:
        channel_env_rows[env] = _sort_rows_by_id(channel_env_rows.get(env, []))
        token_env_rows[env] = _sort_rows_by_id(token_env_rows.get(env, []))

    # 全部失败：不落盘，避免新建空/残缺文件
    if stats.sites_ok == 0:
        fallback = output_path if (output_path and output_path.exists()) else excel_path
        return stats, fallback

    # 任一环境含密钥列则整表用 WITH_KEY 表头
    use_key_headers = include_keys or any(
        "密钥" in row for rows in channel_env_rows.values() for row in rows
    )
    channel_headers = CHANNEL_HEADERS_WITH_KEY if use_key_headers else CHANNEL_HEADERS
    dup_channel_headers = DUP_CHANNEL_HEADERS_WITH_KEY if use_key_headers else DUP_CHANNEL_HEADERS

    # 一键采集：每次新建带时间戳文件（同一次多站共用 output_path；staging 时用 .part.xlsx，完成后再发布）
    if output_path:
        save_path = output_path
    elif staging:
        save_path = build_staging_path(root_dir)
    else:
        save_path = build_timestamped_path(root_dir)
    save_workbook_data(
        save_path,
        channel_env_rows=channel_env_rows,
        token_env_rows=token_env_rows,
        channel_dup_source=channel_env_rows,
        token_dup_source=token_env_rows,
        channel_headers=channel_headers,
        dup_channel_headers=dup_channel_headers,
    )
    return stats, save_path


def _classify_login_probe_error(exc: BaseException) -> str:
    """无验证码登录探测失败时的分类：need_otp / setup / rate_limited / error。"""
    text = str(exc or "")
    lower = text.lower()
    if "尚未绑定" in text:
        return "setup"
    if (
        "429" in text
        or "too many requests" in lower
        or "限流" in text
        or "配额不足" in text
        or "criticalratelimit" in lower
    ):
        return "rate_limited"
    if any(key in text for key in ("两步验证", "验证码", "二次验证", "备用码")) or "2fa" in lower:
        return "need_otp"
    return "error"


def inspect_site_auth(site: SiteConfig, *, use_site_otp: bool = False) -> dict[str, Any]:
    """试登录：默认不带手填验证码，判断该站是否需要登录两步验证。

    ``use_site_otp=True`` 时用站点已填 otp_code / totp_secret 完成认证并写入会话缓存。
    成功会写入会话缓存，后续采集可复用。已有有效会话时不会再打登录接口。
    """
    env = detect_env(site.base_url, site.env)
    label = (site.label or env).strip()
    row: dict[str, Any] = {
        "env": env,
        "label": label,
        "base_url": site.base_url,
        "ok": False,
        "needs_login_otp": False,
        "needs_verify_otp": False,
        "needs_otp_prompt": False,
        "session_ready": False,
        "has_totp_secret": bool(str(site.totp_secret or "").strip()),
        "collectable": False,
        "auth_mode": "access_token" if str(site.access_token or "").strip() else "password",
        "status": "error",
        "message": "",
    }
    try:
        otp = str(site.otp_code or "").strip() if use_site_otp else ""
        auth = _site_auth_payload(site, otp=otp)
    except ValueError as exc:
        row["message"] = str(exc)
        return row

    base_url = normalize_na_base_url(site.base_url) or site.base_url
    client = NewAPIClient(base_url, verify_ssl=bool(site.verify_ssl))
    try:
        authenticate_client(client, auth)
        row["ok"] = True
        row["status"] = "ready"
        row["collectable"] = True
        row["session_ready"] = True
        if use_site_otp and str(site.otp_code or "").strip():
            row["message"] = "认证成功"
        elif row.get("has_totp_secret"):
            row["message"] = "认证成功（已用预置 TOTP 或有效会话）"
        else:
            row["message"] = "登录无需验证码（或已有有效会话）"
        return row
    except Exception as exc:  # noqa: BLE001
        kind = _classify_login_probe_error(exc)
        text = str(exc)
        if kind == "need_otp":
            row["ok"] = True
            row["needs_login_otp"] = True
            row["needs_otp_prompt"] = True
            row["collectable"] = True
            row["status"] = "need_otp"
            row["message"] = "需要临时 6 位验证码（约 30 秒有效）"
            return row
        row["status"] = kind if kind in {"setup", "rate_limited"} else "error"
        row["message"] = text
        return row


def inspect_summary_sites_auth(
    sites: list[SiteConfig],
    *,
    include_keys: bool = False,
) -> list[dict[str, Any]]:
    """逐站审查登录 2FA；勾选含密钥时，可采集站点仍需二次验证码。"""
    items = [inspect_site_auth(site) for site in sites]
    if not include_keys:
        return items
    for item, site in zip(items, sites):
        if not item.get("collectable"):
            continue
        item["needs_verify_otp"] = True
        if site_has_otp_source(site):
            item["needs_otp_prompt"] = False
            item["has_totp_secret"] = True
            if item.get("status") == "ready":
                item["message"] = "登录无需验证码；导出密钥将用 TOTP 密钥自动生成验证码"
            elif item.get("status") == "need_otp":
                item["message"] = "登录与导出密钥将用 TOTP 密钥自动生成验证码"
            continue
        item["needs_otp_prompt"] = True
        if item.get("status") == "ready":
            item["message"] = "登录无需验证码；导出密钥仍需临时 6 位验证码（约 30 秒有效）"
        elif item.get("status") == "need_otp":
            item["message"] = "需要临时 6 位验证码（约 30 秒有效，登录与导出密钥共用）"
    return items


def login_site_auth(site: SiteConfig) -> dict[str, Any]:
    """用站点 otp_code / totp_secret 登录并写入会话缓存（工作台全局认证）。"""
    return inspect_site_auth(site, use_site_otp=True)


def parse_sites(raw_sites: list[dict[str, Any]]) -> list[SiteConfig]:
    sites: list[SiteConfig] = []
    for item in raw_sites or []:
        base_url = str(item.get("base_url") or "").strip()
        if not base_url:
            continue
        user_id_raw = item.get("user_id")
        user_id: int | None
        try:
            user_id = int(user_id_raw) if user_id_raw not in (None, "") else None
        except (TypeError, ValueError):
            user_id = None
        sites.append(
            SiteConfig(
                base_url=base_url,
                username=str(item.get("username") or "").strip(),
                password=str(item.get("password") or ""),
                env=str(item.get("env") or "").strip().lower(),
                label=str(item.get("label") or "").strip(),
                otp_code=str(item.get("otp_code") or item.get("verification_code") or "").strip(),
                totp_secret=pick_totp_secret(mapping=item),
                access_token=str(item.get("access_token") or "").strip(),
                user_id=user_id,
                verify_ssl=coerce_verify_ssl(item.get("verify_ssl"), True),
            )
        )
    for site in sites:
        enrich_site_from_env(site)
    return sites


def build_staging_path(root_dir, when=None):
    from .data_summary_exporter import build_staging_path as _impl

    return _impl(root_dir, when)


def build_timestamped_path(root_dir, when=None):
    from .data_summary_exporter import build_timestamped_path as _impl

    return _impl(root_dir, when)


def latest_excel_path(root_dir):
    from .data_summary_exporter import latest_excel_path as _impl

    return _impl(root_dir)


def load_workbook_data(path):
    from .data_summary_exporter import load_workbook_data as _impl

    return _impl(path)


def save_workbook_data(path, **kwargs):
    from .data_summary_exporter import save_workbook_data as _impl

    return _impl(path, **kwargs)
