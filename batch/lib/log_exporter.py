"""使用日志抓取与导出（CSV / Excel）。

字段对齐 New API（ai-cloud-test）`model.Log` 及 `other` JSON 扩展
（见 service/log_info_generate.go），尽量全量导出可获取内容。
"""

from __future__ import annotations

import csv
import io
import json
import time
from dataclasses import dataclass
from datetime import datetime
from typing import Any, Callable

import requests

from .api_client import NewAPIClient


class LogExportCancelled(Exception):
    """用户点击终止，停止继续拉取使用日志。"""

LOG_TYPE_LABELS = {
    0: "未知",
    1: "充值",
    2: "消费",
    3: "管理",
    4: "系统",
    5: "错误",
    6: "退款",
}

DEFAULT_QUOTA_PER_UNIT = 500_000.0

# other JSON 扁平字段（与上游 Generate*OtherInfo 对齐）
_OTHER_FLAT_FIELDS: list[tuple[str, str]] = [
    ("模型倍率", "model_ratio"),
    ("分组倍率", "group_ratio"),
    ("用户专属倍率", "user_group_ratio"),
    ("补全倍率", "completion_ratio"),
    ("模型固定价", "model_price"),
    ("缓存Tokens", "cache_tokens"),
    ("缓存倍率", "cache_ratio"),
    ("缓存创建Tokens", "cache_creation_tokens"),
    ("缓存创建倍率", "cache_creation_ratio"),
    ("缓存创建Tokens_5m", "cache_creation_tokens_5m"),
    ("缓存创建倍率_5m", "cache_creation_ratio_5m"),
    ("缓存创建Tokens_1h", "cache_creation_tokens_1h"),
    ("缓存创建倍率_1h", "cache_creation_ratio_1h"),
    ("上游模型", "upstream_model_name"),
    ("模型已映射", "is_model_mapped"),
    ("请求路径", "request_path"),
    ("首字延迟ms", "frt"),
    ("ReasoningEffort", "reasoning_effort"),
    ("系统提示被覆盖", "is_system_prompt_overwritten"),
    ("Claude", "claude"),
    ("WebSocket", "ws"),
    ("Audio", "audio"),
    ("音频输入Tokens", "audio_input"),
    ("音频输出Tokens", "audio_output"),
    ("文本输入Tokens", "text_input"),
    ("文本输出Tokens", "text_output"),
    ("音频倍率", "audio_ratio"),
    ("音频补全倍率", "audio_completion_ratio"),
    ("音频输入单独计价", "audio_input_seperate_price"),
    ("音频输入Token计数", "audio_input_token_count"),
    ("图像", "image"),
    ("图像倍率", "image_ratio"),
    ("图像输出Tokens", "image_output"),
    ("Web搜索", "web_search"),
    ("Web搜索次数", "web_search_call_count"),
    ("Web搜索价格", "web_search_price"),
    ("文件搜索", "file_search"),
    ("文件搜索次数", "file_search_call_count"),
    ("文件搜索价格", "file_search_price"),
    ("图像生成调用", "image_generation_call"),
    ("图像生成价格", "image_generation_call_price"),
]

_ADMIN_OTHER_FIELDS: list[tuple[str, str]] = [
    ("重试渠道链", "admin_use_channel"),
    ("多Key渠道", "admin_is_multi_key"),
    ("多Key索引", "admin_multi_key_index"),
    ("本地计Token", "admin_local_count_tokens"),
]

# 顶层 Log 字段 + 衍生字段（Admin 全量）
EXPORT_COLUMNS_ADMIN: list[tuple[str, str]] = [
    ("日志ID", "id"),
    ("用户ID", "user_id"),
    ("时间戳", "created_at"),
    ("时间", "time"),
    ("类型码", "type_code"),
    ("类型", "type"),
    ("渠道ID", "channel"),
    ("渠道名称", "channel_name"),
    ("用户", "username"),
    ("令牌", "token_name"),
    ("令牌ID", "token_id"),
    ("分组", "group"),
    ("模型", "model_name"),
    ("模型展示", "model_name_display"),
    ("是否流式", "is_stream"),
    ("用时秒", "use_time_sec"),
    ("用时/首字", "use_time"),
    ("输入Tokens", "prompt_tokens"),
    ("输出Tokens", "completion_tokens"),
    ("额度quota", "quota"),
    ("花费(USD)", "cost_usd"),
    ("IP", "ip"),
    ("重试", "retry"),
    *_OTHER_FLAT_FIELDS,
    *_ADMIN_OTHER_FIELDS,
    ("计费摘要", "details"),
    ("内容", "content"),
    ("other原文", "other_raw"),
]

# 普通用户：上游会清空 channel_name、剔除 admin_info、混淆 id
EXPORT_COLUMNS_USER: list[tuple[str, str]] = [
    ("日志ID", "id"),
    ("时间戳", "created_at"),
    ("时间", "time"),
    ("类型码", "type_code"),
    ("类型", "type"),
    ("令牌", "token_name"),
    ("令牌ID", "token_id"),
    ("分组", "group"),
    ("模型", "model_name"),
    ("模型展示", "model_name_display"),
    ("是否流式", "is_stream"),
    ("用时秒", "use_time_sec"),
    ("用时/首字", "use_time"),
    ("输入Tokens", "prompt_tokens"),
    ("输出Tokens", "completion_tokens"),
    ("额度quota", "quota"),
    ("花费(USD)", "cost_usd"),
    ("IP", "ip"),
    *_OTHER_FLAT_FIELDS,
    ("计费摘要", "details"),
    ("内容", "content"),
    ("other原文", "other_raw"),
]


@dataclass
class LogExportResult:
    total: int
    rows: list[dict[str, Any]]
    filename: str
    content_type: str
    data: bytes
    is_admin: bool
    gateway_total: int = 0
    model_query: str = ""
    direct_export: bool = False


def _parse_other(raw: Any) -> dict[str, Any]:
    if not raw:
        return {}
    if isinstance(raw, dict):
        return raw
    text = str(raw).strip()
    if not text:
        return {}
    try:
        parsed = json.loads(text)
        return parsed if isinstance(parsed, dict) else {}
    except json.JSONDecodeError:
        return {}


def model_name_matches(log: dict[str, Any], query: str) -> bool:
    """模型名称模糊匹配：不区分大小写的子串包含。

    例：查询 ``glm-x`` 可命中 ``demo-alpha-test-glm-x``；
    填写完整模型名同样可命中。同时匹配 ``model_name``、``model`` 与上游映射名。
    """
    needle = str(query or "").strip().lower()
    if not needle:
        return True
    other = _parse_other(log.get("other"))
    candidates = (
        log.get("model_name"),
        log.get("model"),
        other.get("upstream_model_name"),
        other.get("model_name"),
        other.get("model"),
    )
    return any(needle in str(item or "").lower() for item in candidates)


def _format_timestamp(created_at: int | float | str | None) -> str:
    if created_at in (None, ""):
        return ""
    try:
        ts = int(float(created_at))
        return datetime.fromtimestamp(ts).strftime("%Y-%m-%d %H:%M:%S")
    except (TypeError, ValueError, OSError):
        return str(created_at)


def _format_cost_usd(quota: int | float | None, quota_per_unit: float) -> str:
    if quota in (None, ""):
        return ""
    unit = quota_per_unit or DEFAULT_QUOTA_PER_UNIT
    try:
        value = float(quota) / unit
        return f"{value:.6f}"
    except (TypeError, ValueError, ZeroDivisionError):
        return ""


def _cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)
    return value


def _effective_group_ratio(other: dict[str, Any]) -> tuple[float, str]:
    group_ratio = other.get("group_ratio")
    user_group_ratio = other.get("user_group_ratio")
    try:
        user_ratio = float(user_group_ratio)
        if user_ratio != -1:
            return user_ratio, "专属倍率"
    except (TypeError, ValueError):
        pass
    try:
        return float(group_ratio or 1), "分组倍率"
    except (TypeError, ValueError):
        return 1.0, "分组倍率"


def _format_details(log: dict[str, Any], other: dict[str, Any] | None = None) -> str:
    log_type = int(log.get("type") or 0)
    if log_type != 2:
        return str(log.get("content") or "")

    if other is None:
        other = _parse_other(log.get("other"))
    if not other:
        return str(log.get("content") or "")

    group_ratio, ratio_label = _effective_group_ratio(other)
    model_price = other.get("model_price", -1)
    try:
        model_price = float(model_price)
    except (TypeError, ValueError):
        model_price = -1

    if model_price != -1:
        return f"价格 ${model_price:.6f} * {ratio_label} {group_ratio:g}x"

    model_ratio = other.get("model_ratio")
    completion_ratio = other.get("completion_ratio", 0)
    try:
        model_ratio = float(model_ratio or 0)
        completion_ratio = float(completion_ratio or 0)
    except (TypeError, ValueError):
        model_ratio = 0.0
        completion_ratio = 0.0

    input_price = model_ratio * 2.0
    completion_price = model_ratio * 2.0 * completion_ratio
    parts = [
        f"{ratio_label} {group_ratio:g}x",
        f"输入 ${input_price:g} / 1M tokens",
        f"输出 ${completion_price:g} / 1M tokens",
    ]

    cache_tokens = int(other.get("cache_tokens") or 0)
    if cache_tokens > 0:
        cache_ratio = float(other.get("cache_ratio") or 1.0)
        parts.append(f"缓存 {cache_ratio:g}x")

    if other.get("request_path"):
        parts.append(f"路径 {other['request_path']}")

    return ", ".join(parts)


def _format_use_time(log: dict[str, Any], other: dict[str, Any] | None = None) -> str:
    use_time = log.get("use_time")
    if use_time in (None, ""):
        return ""
    text = f"{use_time}s"
    if log.get("is_stream"):
        if other is None:
            other = _parse_other(log.get("other"))
        frt = other.get("frt")
        if frt not in (None, ""):
            try:
                text += f" / {float(frt) / 1000.0:.1f}s"
            except (TypeError, ValueError):
                pass
        text += " 流"
    else:
        text += " 非流"
    return text


def _format_retry(log: dict[str, Any], *, is_admin: bool, other: dict[str, Any] | None = None) -> str:
    if not is_admin:
        return ""
    log_type = int(log.get("type") or 0)
    if log_type not in (2, 5):
        return ""

    if other is None:
        other = _parse_other(log.get("other"))
    admin_info = other.get("admin_info") if isinstance(other.get("admin_info"), dict) else {}
    use_channel = admin_info.get("use_channel")
    if isinstance(use_channel, list) and use_channel:
        return "渠道：" + "->".join(str(item) for item in use_channel)
    channel = log.get("channel")
    return f"渠道：{channel}" if channel not in (None, "") else ""


def _resolve_group(log: dict[str, Any], other: dict[str, Any] | None = None) -> str:
    group = str(log.get("group") or "").strip()
    if group:
        return group
    if other is None:
        other = _parse_other(log.get("other"))
    return str(other.get("group") or "")


def _other_raw_text(log: dict[str, Any]) -> str:
    raw = log.get("other")
    if raw in (None, ""):
        return ""
    if isinstance(raw, dict):
        return json.dumps(raw, ensure_ascii=False)
    return str(raw)


def format_log_row(
    log: dict[str, Any],
    *,
    is_admin: bool,
    quota_per_unit: float,
) -> dict[str, Any]:
    other = _parse_other(log.get("other"))
    admin_info = other.get("admin_info") if isinstance(other.get("admin_info"), dict) else {}

    model_name = str(log.get("model_name") or "")
    upstream = str(other.get("upstream_model_name") or "")
    model_display = model_name
    if other.get("is_model_mapped") and upstream and upstream != model_name:
        model_display = f"{model_name} -> {upstream}"

    type_code = int(log.get("type") or 0)
    use_channel = admin_info.get("use_channel")
    if isinstance(use_channel, list):
        use_channel_text = "->".join(str(x) for x in use_channel)
    else:
        use_channel_text = _cell(use_channel)

    row: dict[str, Any] = {
        "id": log.get("id", ""),
        "user_id": log.get("user_id", "") if is_admin else "",
        "created_at": log.get("created_at", ""),
        "time": _format_timestamp(log.get("created_at")),
        "type_code": type_code,
        "type": LOG_TYPE_LABELS.get(type_code, "未知"),
        "channel": log.get("channel", "") if is_admin else "",
        "channel_name": (log.get("channel_name") or "") if is_admin else "",
        "username": (log.get("username") or "") if is_admin else "",
        "token_name": log.get("token_name") or "",
        "token_id": log.get("token_id", ""),
        "group": _resolve_group(log, other),
        "model_name": model_name,
        "model_name_display": model_display,
        "is_stream": _cell(log.get("is_stream")),
        "use_time_sec": log.get("use_time", ""),
        "use_time": _format_use_time(log, other),
        "prompt_tokens": log.get("prompt_tokens", ""),
        "completion_tokens": log.get("completion_tokens", ""),
        "quota": log.get("quota", ""),
        "cost_usd": _format_cost_usd(log.get("quota"), quota_per_unit),
        "ip": log.get("ip") or "",
        "retry": _format_retry(log, is_admin=is_admin, other=other),
        "details": _format_details(log, other),
        "content": log.get("content") or "",
        "other_raw": _other_raw_text(log),
        "admin_use_channel": use_channel_text if is_admin else "",
        "admin_is_multi_key": _cell(admin_info.get("is_multi_key")) if is_admin else "",
        "admin_multi_key_index": _cell(admin_info.get("multi_key_index")) if is_admin else "",
        "admin_local_count_tokens": _cell(admin_info.get("local_count_tokens")) if is_admin else "",
    }

    for _label, key in _OTHER_FLAT_FIELDS:
        row[key] = _cell(other.get(key))

    return row


def _is_rate_limited(exc: BaseException) -> bool:
    if isinstance(exc, requests.HTTPError) and exc.response is not None:
        return int(exc.response.status_code) == 429
    text = str(exc)
    return "429" in text or "Too Many Requests" in text


def _retry_wait_seconds(exc: BaseException, attempt: int) -> float:
    """attempt 从 1 起；优先尊重 Retry-After。"""
    resp = getattr(exc, "response", None)
    if resp is not None:
        raw = resp.headers.get("Retry-After") or resp.headers.get("retry-after")
        if raw:
            try:
                return max(1.0, float(raw))
            except ValueError:
                pass
    # 2, 4, 8, 16, 32, 60…
    return min(60.0, float(2 ** attempt))


def _raise_if_cancelled(cancel: Callable[[], bool] | None) -> None:
    if cancel and cancel():
        raise LogExportCancelled("已终止导出")


def _sleep_interruptible(seconds: float, cancel: Callable[[], bool] | None) -> None:
    """可被终止打断的等待，避免页间隔/429 重试期间点了终止仍空等。"""
    if seconds <= 0:
        _raise_if_cancelled(cancel)
        return
    deadline = time.monotonic() + seconds
    while True:
        _raise_if_cancelled(cancel)
        remaining = deadline - time.monotonic()
        if remaining <= 0:
            return
        time.sleep(min(0.2, remaining))


def _gateway_export_filters(filters: dict[str, Any]) -> dict[str, Any]:
    """把本地筛选参数映射到新版 /api/log/export。"""
    params: dict[str, Any] = {"format": "placeholder"}
    for source_key, target_key in (
        ("log_type", "type"),
        ("start_timestamp", "start_timestamp"),
        ("end_timestamp", "end_timestamp"),
        ("model_name", "model_name"),
        ("username", "username"),
        ("token_name", "token_name"),
        ("channel", "channel"),
        ("group", "group"),
    ):
        value = filters.get(source_key)
        if value not in (None, ""):
            params[target_key] = value
    return params


def _looks_like_export_payload(data: bytes, export_format: str, content_type: str) -> bool:
    if not data:
        return False
    if export_format == "xlsx":
        return data[:2] == b"PK"
    content_type = content_type.lower()
    return data.startswith(b"\xef\xbb\xbf") or "text/csv" in content_type


def _gateway_direct_export(
    client: NewAPIClient,
    *,
    filters: dict[str, Any],
    export_format: str,
    progress: Callable[[int, int], None] | None,
    cancel: Callable[[], bool] | None,
    dest_dir: Any,
) -> LogExportResult | None:
    """新版上游直接导出；老版本 404/405 或返回非日志文件时返回 None。"""
    _raise_if_cancelled(cancel)
    if progress:
        try:
            progress(0, 0, "正在调用网关直接导出…")
        except TypeError:
            progress(0, 0)

    params = _gateway_export_filters(filters)
    params["format"] = export_format
    try:
        resp = client.session.get(
            f"{client.base_url}/api/log/export",
            params=params,
            headers=client._headers(),
            timeout=max(int(client.timeout or 30), 900),
        )
    except requests.RequestException:
        return None

    if resp.status_code in (404, 405):
        return None
    resp.raise_for_status()
    content_type = resp.headers.get("Content-Type") or resp.headers.get("content-type") or ""
    if not _looks_like_export_payload(resp.content, export_format, content_type):
        return None

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"usage_logs_{stamp}.{export_format}"
    content_type = (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        if export_format == "xlsx"
        else "text/csv; charset=utf-8"
    )
    data = resp.content
    if dest_dir is not None:
        from pathlib import Path

        path = Path(dest_dir) / filename
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(data)
        data = b""

    return LogExportResult(
        total=1,
        rows=[],
        filename=filename,
        content_type=content_type,
        data=data,
        is_admin=True,
        gateway_total=1,
        direct_export=True,
    )


def fetch_all_usage_logs(
    client: NewAPIClient,
    *,
    is_admin: bool,
    filters: dict[str, Any] | None = None,
    page_size: int = 100,
    progress: Callable[[int, int], None] | None = None,
    # 两次 list_logs 的启动间隔下限。NA 全局管理 API 约 180 次/180 秒（1 次/秒）。
    # 请求本身若已超过该间隔则不再额外空等，避免「接口 1.5s + 再睡 1.5s」变成 3 秒/页。
    page_delay: float = 1.05,
    max_retries: int = 10,
    cancel: Callable[[], bool] | None = None,
    stats: dict[str, Any] | None = None,
) -> list[dict[str, Any]]:
    filters = dict(filters or {})
    # NA 接口多为精确匹配；模型名改为本地模糊过滤（子串、忽略大小写）
    model_query = str(filters.pop("model_name", "") or "").strip()
    page = 1
    items: list[dict[str, Any]] = []
    total = 0

    def _emit(fetched: int, known_total: int, message: str) -> None:
        if not progress:
            return
        try:
            progress(fetched, known_total, message)
        except TypeError:
            progress(fetched, known_total)

    while True:
        _raise_if_cancelled(cancel)
        _emit(len(items), total, f"正在拉取第 {page} 页…")
        payload: dict[str, Any] | None = None
        last_exc: BaseException | None = None
        req_started = time.monotonic()
        for attempt in range(max_retries + 1):
            _raise_if_cancelled(cancel)
            try:
                req_started = time.monotonic()
                payload = client.list_logs(
                    page=page,
                    page_size=page_size,
                    admin=is_admin,
                    **filters,
                )
                break
            except LogExportCancelled:
                raise
            except Exception as exc:  # noqa: BLE001
                last_exc = exc
                if not _is_rate_limited(exc) or attempt >= max_retries:
                    raise
                wait = _retry_wait_seconds(exc, attempt + 1)
                _sleep_interruptible(wait, cancel)

        if payload is None:
            raise RuntimeError(str(last_exc) if last_exc else "获取使用日志失败")

        batch = payload.get("items") or []
        if not isinstance(batch, list):
            batch = []
        page_size_eff = int(payload.get("page_size") or page_size or 100)
        reported_total = payload.get("total")
        try:
            reported_total = int(reported_total) if reported_total is not None else None
        except (TypeError, ValueError):
            reported_total = None
        if reported_total is not None and reported_total > 0:
            total = reported_total
        if not batch:
            break
        items.extend(batch)
        shown = len(items) if total <= 0 else min(len(items), total)
        _emit(shown, total, f"第 {page} 页完成")
        if total > 0 and len(items) >= total:
            items = items[:total]
            break
        short_page = len(batch) < page_size_eff
        if short_page and not (total > 0 and len(items) < total):
            break
        page += 1
        leftover = float(page_delay) - (time.monotonic() - req_started)
        if leftover > 0:
            _sleep_interruptible(leftover, cancel)

    gateway_count = len(items)
    if stats is not None:
        stats["gateway_total"] = gateway_count
        stats["model_query"] = model_query
        stats["matched"] = gateway_count
    if model_query:
        items = [log for log in items if model_name_matches(log, model_query)]
        if stats is not None:
            stats["matched"] = len(items)
        _emit(
            len(items),
            gateway_count or total,
            f"模型名「{model_query}」筛选后保留 {len(items)} 条（网关共 {gateway_count} 条）",
        )
    return items


def _cell_has_content(value: Any) -> bool:
    """数据格是否有实质内容（空白字符串 / None 视为空）。"""
    if value is None:
        return False
    if isinstance(value, str):
        return bool(value.strip())
    return True


def _drop_empty_columns(
    headers: list[str],
    rows: list[list[Any]],
) -> tuple[list[str], list[list[Any]]]:
    """删除所有数据行均为空的列（仅表头有字的列不导出）。"""
    if not headers:
        return headers, rows
    if not rows:
        return headers, rows

    keep: list[int] = []
    for col_idx in range(len(headers)):
        if any(_cell_has_content(row[col_idx] if col_idx < len(row) else None) for row in rows):
            keep.append(col_idx)

    if not keep:
        # 极端情况：全部列空时仍保留原表头，避免导出无列文件
        return headers, rows

    new_headers = [headers[i] for i in keep]
    new_rows = [[row[i] if i < len(row) else "" for i in keep] for row in rows]
    return new_headers, new_rows


def _rows_to_table(
    logs: list[dict[str, Any]],
    *,
    is_admin: bool,
    quota_per_unit: float,
    progress: Callable[..., None] | None = None,
    cancel: Callable[[], bool] | None = None,
) -> tuple[list[str], list[list[Any]]]:
    columns = EXPORT_COLUMNS_ADMIN if is_admin else EXPORT_COLUMNS_USER
    headers = [label for label, _ in columns]
    keys = [key for _, key in columns]
    nonempty = [False] * len(keys)
    rows: list[list[Any]] = []
    total = len(logs)
    for idx, log in enumerate(logs):
        if cancel and idx % 4000 == 0:
            _raise_if_cancelled(cancel)
        if progress and idx > 0 and idx % 25000 == 0:
            try:
                progress(total, total, f"正在整理 {idx} / {total} 条…")
            except TypeError:
                progress(total, total)
        formatted = format_log_row(log, is_admin=is_admin, quota_per_unit=quota_per_unit)
        row = [formatted.get(key, "") for key in keys]
        for col_idx, value in enumerate(row):
            if not nonempty[col_idx] and _cell_has_content(value):
                nonempty[col_idx] = True
        rows.append(row)
    keep = [i for i, flag in enumerate(nonempty) if flag]
    if not keep:
        return headers, rows
    if len(keep) == len(headers):
        return headers, rows
    new_headers = [headers[i] for i in keep]
    new_rows = [[row[i] for i in keep] for row in rows]
    return new_headers, new_rows


def export_usage_logs_csv(
    logs: list[dict[str, Any]],
    *,
    is_admin: bool,
    quota_per_unit: float,
) -> bytes:
    headers, rows = _rows_to_table(logs, is_admin=is_admin, quota_per_unit=quota_per_unit)
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(headers)
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8-sig")


def _header_write_only_cells(ws, headers: list[str]) -> list:
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill

    fill = PatternFill("solid", fgColor="FF2F5496")
    font = Font(name="宋体", bold=True, color="FFFFFFFF", size=11)
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cells = []
    for header in headers:
        cell = WriteOnlyCell(ws, value=header)
        cell.fill = fill
        cell.font = font
        cell.alignment = align
        cells.append(cell)
    return cells


def _set_log_column_widths(ws, headers: list[str]) -> None:
    from openpyxl.utils import get_column_letter

    for idx, header in enumerate(headers, start=1):
        letter = get_column_letter(idx)
        if header in {"other原文", "内容", "计费摘要"}:
            ws.column_dimensions[letter].width = 36
        elif header in {"模型展示", "请求路径", "重试渠道链"}:
            ws.column_dimensions[letter].width = 28
        else:
            ws.column_dimensions[letter].width = min(18, max(10, len(header) + 2))


def export_usage_logs_xlsx(
    logs: list[dict[str, Any]],
    *,
    is_admin: bool,
    quota_per_unit: float,
    progress: Callable[..., None] | None = None,
    cancel: Callable[[], bool] | None = None,
    dest_path: Any = None,
) -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("导出 Excel 需要安装 openpyxl，请改用 CSV 或执行 pip install openpyxl") from exc

    total = len(logs)
    if progress:
        try:
            progress(total, total, f"正在整理 {total} 条日志…")
        except TypeError:
            progress(total, total)
    headers, rows = _rows_to_table(
        logs,
        is_admin=is_admin,
        quota_per_unit=quota_per_unit,
        progress=progress,
        cancel=cancel,
    )
    _raise_if_cancelled(cancel)
    if progress:
        try:
            progress(total, total, f"正在写入 Excel（{len(rows)} 行）…")
        except TypeError:
            progress(total, total)

    # write_only：大批量不再逐格套边框（30 万行 × 数十列会卡十几分钟）
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("使用日志")
    ws.append(_header_write_only_cells(ws, headers))
    for idx, row in enumerate(rows):
        if cancel and idx % 8000 == 0:
            _raise_if_cancelled(cancel)
        if progress and idx > 0 and idx % 50000 == 0:
            try:
                progress(total, total, f"正在写入 Excel {idx} / {len(rows)} 行…")
            except TypeError:
                progress(total, total)
        ws.append(row)

    _set_log_column_widths(ws, headers)
    ws.freeze_panes = "A2"
    if headers:
        last_col = get_column_letter(len(headers))
        ws.auto_filter.ref = f"A1:{last_col}{max(1, len(rows) + 1)}"

    if dest_path is not None:
        from pathlib import Path

        path = Path(dest_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
        return b""

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def run_logs_export(
    client: NewAPIClient,
    *,
    is_admin: bool,
    filters: dict[str, Any] | None = None,
    export_format: str = "xlsx",
    quota_per_unit: float | None = None,
    progress: Callable[[int, int], None] | None = None,
    cancel: Callable[[], bool] | None = None,
    dest_dir: Any = None,
) -> LogExportResult:
    fmt = (export_format or "xlsx").lower()
    if fmt not in {"xlsx", "csv"}:
        raise ValueError("导出格式仅支持 xlsx 或 csv")

    if is_admin:
        direct = _gateway_direct_export(
            client,
            filters=filters or {},
            export_format=fmt,
            progress=progress,
            cancel=cancel,
            dest_dir=dest_dir,
        )
        if direct is not None:
            return direct

    if quota_per_unit is None:
        try:
            _raise_if_cancelled(cancel)
            if progress:
                try:
                    progress(0, 0, "正在读取站点配置…")
                except TypeError:
                    progress(0, 0)
            status = client.get_status()
            quota_per_unit = float((status.get("data") or status).get("quota_per_unit") or DEFAULT_QUOTA_PER_UNIT)
        except LogExportCancelled:
            raise
        except Exception:
            quota_per_unit = DEFAULT_QUOTA_PER_UNIT

    fetch_stats: dict[str, Any] = {}
    logs = fetch_all_usage_logs(
        client,
        is_admin=is_admin,
        filters=filters,
        progress=progress,
        cancel=cancel,
        stats=fetch_stats,
    )
    _raise_if_cancelled(cancel)
    gateway_total = int(fetch_stats.get("gateway_total") or len(logs))
    model_query = str(fetch_stats.get("model_query") or "").strip()
    if not logs:
        return LogExportResult(
            total=0,
            rows=[],
            filename="",
            content_type="",
            data=b"",
            is_admin=is_admin,
            gateway_total=gateway_total,
            model_query=model_query,
        )
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    dest_path = None
    if dest_dir is not None:
        from pathlib import Path

        dest_path = Path(dest_dir)
    if fmt == "xlsx":
        filename = f"usage_logs_{stamp}.xlsx"
        content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        data = export_usage_logs_xlsx(
            logs,
            is_admin=is_admin,
            quota_per_unit=quota_per_unit,
            progress=progress,
            cancel=cancel,
            dest_path=(dest_path / filename) if dest_path is not None else None,
        )
    else:
        if progress:
            try:
                progress(len(logs), len(logs), f"正在写入 CSV（{len(logs)} 条）…")
            except TypeError:
                progress(len(logs), len(logs))
        data = export_usage_logs_csv(logs, is_admin=is_admin, quota_per_unit=quota_per_unit)
        filename = f"usage_logs_{stamp}.csv"
        content_type = "text/csv; charset=utf-8"
        if dest_path is not None:
            dest_path.mkdir(parents=True, exist_ok=True)
            (dest_path / filename).write_bytes(data)
            data = b""

    return LogExportResult(
        total=len(logs),
        rows=[],
        filename=filename,
        content_type=content_type,
        data=data,
        is_admin=is_admin,
        gateway_total=gateway_total,
        model_query=model_query,
    )
