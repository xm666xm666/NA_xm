# -*- coding: utf-8 -*-
"""上游账单 · 火山云 + Azure 汇总到同一工作簿。"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from .azure_billing import (
    AzureBillStats,
    DAY_HEADERS as AZURE_DAY_HEADERS,
    MONTH_HEADERS as AZURE_MONTH_HEADERS,
    SHEET_NOTE,
)
from .volc_billing import (
    BillCollectStats,
    COST_COL_HEADER,
    DAY_HEADERS as VOLC_DAY_HEADERS,
    MONTH_HEADERS as VOLC_MONTH_HEADERS,
    SAVING_PLAN_COL_HEADER,
    bill_month_label,
    excel_cost_headers,
    sheet_data_rows,
    write_bill_sheet,
)

COMBINED_STEM = "上游账单_汇总"
TIMESTAMP_FORMAT = "%Y%m%d%H%M%S"
SHEET_VOLC_DAY = "火山云-按天"
SHEET_VOLC_MONTH = "火山云-按月"
SHEET_AZURE_DAY = "Azure-按天"
SHEET_AZURE_MONTH = "Azure-按月"


def excel_dir(root_dir: Path) -> Path:
    return Path(root_dir) / "excel"


def build_combined_bill_filename(when: datetime | None = None, *, bill_period: Any = None) -> str:
    stamp = (when or datetime.now()).strftime(TIMESTAMP_FORMAT)
    return f"{COMBINED_STEM}_{bill_month_label(bill_period)}_{stamp}.xlsx"


def build_combined_bill_path(root_dir: Path, when: datetime | None = None, *, bill_period: Any = None) -> Path:
    return excel_dir(root_dir) / build_combined_bill_filename(when, bill_period=bill_period)


def list_combined_bill_excel_files(root_dir: Path) -> list[Path]:
    from .data_summary_exporter import sort_excel_files_latest_first

    directory = excel_dir(root_dir)
    if not directory.exists():
        return []
    return sort_excel_files_latest_first(list(directory.glob(f"{COMBINED_STEM}_*.xlsx")))


def latest_combined_bill_excel_path(root_dir: Path) -> Path | None:
    files = list_combined_bill_excel_files(root_dir)
    return files[0] if files else None


def save_combined_bill_workbook(
    path: Path,
    volc_stats: BillCollectStats | None,
    azure_stats: AzureBillStats | None,
) -> Path:
    from openpyxl import Workbook

    volc = volc_stats or BillCollectStats()
    azure = azure_stats or AzureBillStats()
    dest = Path(path)
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws_vd = wb.active
    ws_vd.title = SHEET_VOLC_DAY
    write_bill_sheet(ws_vd, VOLC_DAY_HEADERS, volc.rows_day)
    ws_vm = wb.create_sheet(SHEET_VOLC_MONTH)
    write_bill_sheet(ws_vm, VOLC_MONTH_HEADERS, volc.rows_month)
    ws_ad = wb.create_sheet(SHEET_AZURE_DAY)
    write_bill_sheet(ws_ad, AZURE_DAY_HEADERS, azure.rows_day)
    ws_am = wb.create_sheet(SHEET_AZURE_MONTH)
    write_bill_sheet(ws_am, AZURE_MONTH_HEADERS, azure.rows_month)
    period = volc.bill_period or azure.bill_period
    currency = "、".join(azure.currencies) if azure.currencies else "CNY"
    note = wb.create_sheet(SHEET_NOTE)
    write_bill_sheet(
        note,
        ["项", "说明"],
        [
            {"项": "账期", "说明": period},
            {"项": "火山云接口", "说明": volc.source_action or ""},
            {"项": "Azure 接口", "说明": azure.source_action or ""},
            {
                "项": "工作表",
                "说明": "同一工作簿四张费用表：火山云-按天/按月、Azure-按天/按月",
            },
            {
                "项": COST_COL_HEADER,
                "说明": f"人民币 ¥；末行 SUM 合计。Azure 币种：{currency}",
            },
            {
                "项": SAVING_PLAN_COL_HEADER,
                "说明": "火山云 SavingPlanDeductionDiscountTotalAmount（节省计划抵扣总额）；与应付金额分开汇总",
            },
        ],
        add_sum=False,
    )
    wb.save(dest)
    wb.close()
    return dest


def combined_bill_workbook_stats(path: Path) -> dict[str, Any]:
    if not path or not Path(path).exists():
        return {"rows": 0, "sheet": SHEET_VOLC_DAY, "sheets": {}, "total_rows": 0}
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = {name: sheet_data_rows(wb[name]) for name in wb.sheetnames}
        return {
            "rows": sheets.get(SHEET_VOLC_DAY, 0),
            "sheet": f"{SHEET_VOLC_DAY} + {SHEET_AZURE_DAY}",
            "sheets": sheets,
            "volc_day_rows": sheets.get(SHEET_VOLC_DAY, 0),
            "volc_month_rows": sheets.get(SHEET_VOLC_MONTH, 0),
            "azure_day_rows": sheets.get(SHEET_AZURE_DAY, 0),
            "azure_month_rows": sheets.get(SHEET_AZURE_MONTH, 0),
            "total_rows": (
                sheets.get(SHEET_VOLC_DAY, 0)
                + sheets.get(SHEET_VOLC_MONTH, 0)
                + sheets.get(SHEET_AZURE_DAY, 0)
                + sheets.get(SHEET_AZURE_MONTH, 0)
            ),
        }
    finally:
        wb.close()


def volc_excel_headers() -> list[str]:
    return excel_cost_headers(list(VOLC_DAY_HEADERS))


def azure_excel_headers() -> list[str]:
    return excel_cost_headers(list(AZURE_DAY_HEADERS))
