# -*- coding: utf-8 -*-
"""产物列表：含密钥须看单元格是否有值，不能只看列名或文件类型。"""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

BATCH_DIR = Path(__file__).resolve().parents[1]
if str(BATCH_DIR) not in sys.path:
    sys.path.insert(0, str(BATCH_DIR))

from lib.artifacts import list_artifacts, workbook_has_filled_secrets  # noqa: E402
from lib.data_summary import CHANNEL_HEADERS_WITH_KEY  # noqa: E402
from lib.data_summary_exporter import save_workbook_data  # noqa: E402
from lib.overview_analyzer import ANALYSIS_STEM  # noqa: E402
from lib.overview_exporter import OVERVIEW_STEM  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402


def _channel(name: str, key: str = "") -> dict:
    return {
        "名称": name,
        "分组": "g",
        "类型": "自定义渠道",
        "状态": "已启用",
        "模型列表": "m",
        "模型重定向": "",
        "base_url": "https://x",
        "密钥": key,
        "部署时间": "",
    }


def _write_summary(root: Path, filename: str, key: str = "") -> Path:
    dest = root / "excel" / filename
    save_workbook_data(
        dest,
        channel_env_rows={
            "test": [_channel("ch-a", key), _channel("ch-b")],
            "uat": [],
            "prod": [],
        },
        token_env_rows={"test": [], "uat": [], "prod": []},
        channel_headers=CHANNEL_HEADERS_WITH_KEY,
    )
    return dest


class ArtifactSecretFlagTests(unittest.TestCase):
    def test_empty_key_column_is_not_secret(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            dest = _write_summary(root, "ai网关令牌及渠道列表_20260908150402.xlsx")
            self.assertFalse(workbook_has_filled_secrets(dest))
            items = list_artifacts(root, "summary", limit=5)
            self.assertEqual(items[0]["filename"], dest.name)
            self.assertFalse(items[0]["has_secrets"])

    def test_placeholder_key_is_not_secret(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            dest = _write_summary(root, "ai网关令牌及渠道列表_20260908150402.xlsx", "N/A")
            self.assertFalse(workbook_has_filled_secrets(dest))
            items = list_artifacts(root, "summary", limit=5)
            self.assertFalse(items[0]["has_secrets"])

    def test_filled_key_cell_is_secret(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            dest = _write_summary(root, "ai网关令牌及渠道列表_20260908150402.xlsx", "sk-live")
            self.assertTrue(workbook_has_filled_secrets(dest))
            items = list_artifacts(root, "summary", limit=5)
            self.assertTrue(items[0]["has_secrets"])

    def test_overview_empty_keys_not_secret(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            dest = _write_summary(root, f"{OVERVIEW_STEM}_20260908150343.xlsx")
            items = list_artifacts(root, "overview", limit=5)
            self.assertEqual(items[0]["filename"], dest.name)
            self.assertFalse(items[0]["has_secrets"])

    def test_overview_upstream_api_key_counts(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            dest = _write_summary(root, f"{OVERVIEW_STEM}_20260908150343.xlsx")
            wb = load_workbook(dest)
            ws = wb.create_sheet("火山云")
            ws.append(["接入点名称", "API Key", "项目方"])
            ws.append(["ep-1", "ark-xxxxx", "demo"])
            wb.save(dest)
            wb.close()
            self.assertTrue(workbook_has_filled_secrets(dest))
            items = list_artifacts(root, "overview", limit=5)
            self.assertTrue(items[0]["has_secrets"])

    def test_analysis_key_digest_is_not_secret(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            dest = root / "excel" / f"{ANALYSIS_STEM}_20260908150000.xlsx"
            dest.parent.mkdir(parents=True, exist_ok=True)
            wb = Workbook()
            ws = wb.active
            ws.title = "建议处置"
            ws.append(["名称", "密钥摘要", "问题"])
            ws.append(["ch-a", "sk-abc…wxyz", "下游无上游"])
            wb.save(dest)
            wb.close()
            items = list_artifacts(root, "analysis", limit=5)
            self.assertFalse(items[0]["has_secrets"])
            self.assertFalse(workbook_has_filled_secrets(dest))


if __name__ == "__main__":
    unittest.main()
