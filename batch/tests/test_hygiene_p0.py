# -*- coding: utf-8 -*-
from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from lib.artifacts import list_artifacts, resolve_excel_path
from lib.hygiene_actions import (
    action_key,
    is_hold_suggestion,
    preview_disable_from_actions,
)
from lib.hygiene_policy import (
    uncommon_metric_label,
    unused_metric_label,
    normalize_policy,
    save_policy,
)
from lib.overview_analyzer import SHEET_SUMMARY, classify_usage, relax_for_new_deploy, save_analysis_workbook


class HygienePolicyTests(unittest.TestCase):
    def test_normalize_rejects_bad_thresholds(self):
        with self.assertRaises(ValueError):
            normalize_policy({"unused_used_usd": 2, "uncommon_used_usd": 1})

    def test_save_and_load_roundtrip(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            saved = save_policy(root, {"new_deploy_days": 7, "uncommon_used_usd": 2})
            self.assertEqual(saved["new_deploy_days"], 7)
            self.assertEqual(saved["uncommon_used_usd"], 2.0)
            loaded = json.loads((root / "env" / "hygiene_policy.json").read_text(encoding="utf-8"))
            self.assertEqual(loaded["new_deploy_days"], 7)


class ClassifyUsagePolicyTests(unittest.TestCase):
    def test_custom_thresholds(self):
        self.assertEqual(classify_usage(0.5, unused=0.4, uncommon=3), "不常用")
        self.assertEqual(classify_usage(0.4, unused=0.4, uncommon=3), "不用")
        self.assertEqual(classify_usage(3, unused=0.4, uncommon=3), "常用")
        self.assertEqual(classify_usage(4.9, unused=0, uncommon=5), "不常用")
        self.assertEqual(classify_usage(5, unused=0, uncommon=5), "常用")

    def test_metric_labels_follow_policy(self):
        policy = {"unused_used_usd": 0, "uncommon_used_usd": 5}
        self.assertEqual(unused_metric_label(policy), "不用（已用≤$0）")
        self.assertEqual(uncommon_metric_label(policy), "不常用（已用<$5）")

    def test_summary_sheet_writes_policy_thresholds(self):
        with tempfile.TemporaryDirectory() as tmp:
            dest = Path(tmp) / "智能分析表格_policy.xlsx"
            save_analysis_workbook(
                dest,
                {
                    "stats": {"balance_unused": 1, "balance_uncommon": 3},
                    "policy": {
                        "new_deploy_days": 10,
                        "unused_used_usd": 0,
                        "uncommon_used_usd": 5,
                    },
                },
            )
            from openpyxl import load_workbook

            wb = load_workbook(dest, read_only=True, data_only=True)
            try:
                labels = [
                    str(row[0].value or "")
                    for row in wb[SHEET_SUMMARY].iter_rows(min_row=2, max_col=1)
                ]
                notes = {
                    str(row[0].value or ""): str(row[1].value or "")
                    for row in wb["说明"].iter_rows(min_row=2, max_col=2)
                }
            finally:
                wb.close()
            self.assertIn("不用（已用≤$0）", labels)
            self.assertIn("不常用（已用<$5）", labels)
            self.assertNotIn("不常用（已用<$1）", labels)
            self.assertIn("已用 < $5 视为不常用", notes.get("渠道余额", ""))
            self.assertNotIn("已用 < $1 视为不常用", notes.get("渠道余额", ""))

    def test_relax_uses_days(self):
        from datetime import datetime, timedelta

        now = datetime(2026, 8, 28)
        fresh = (now - timedelta(days=3)).strftime("%Y-%m-%d %H:%M:%S")
        row = relax_for_new_deploy({"级别": "高", "建议": "禁用"}, fresh, now=now, days=5)
        self.assertEqual(row["级别"], "中")
        self.assertIn("暂缓禁用", row["建议"])
        old = relax_for_new_deploy({"级别": "高", "建议": "禁用"}, fresh, now=now, days=2)
        self.assertEqual(old["级别"], "高")


class ArtifactsTests(unittest.TestCase):
    def test_list_and_resolve(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            excel = root / "excel"
            excel.mkdir()
            target = excel / "智能分析表格_20260828101010.xlsx"
            target.write_bytes(b"PK\x03\x04")
            items = list_artifacts(root, "analysis", limit=5)
            self.assertEqual(len(items), 1)
            self.assertEqual(items[0]["id"], target.name)
            resolved = resolve_excel_path(root, items[0]["id"])
            self.assertEqual(resolved, target.resolve())
            with self.assertRaises(ValueError):
                resolve_excel_path(root, "../secrets.xlsx")


class HygieneActionsTests(unittest.TestCase):
    def test_preview_skips_hold_unless_forced(self):
        rows = [
            {
                "key": "1|uat|ch| |g|下游无上游",
                "级别": "高",
                "类别": "下游无上游",
                "站点": "uat",
                "渠道名称": "ch",
                "分组": "g",
                "状态": "已启用",
                "建议": "建议禁用；部署未满 10 天，暂缓禁用/停用，观察后再处置",
                "hold": True,
                "status": "held",
            },
            {
                "key": "2|uat|old| |g|下游无上游",
                "级别": "高",
                "类别": "下游无上游",
                "站点": "uat",
                "渠道名称": "old",
                "分组": "g",
                "状态": "已启用",
                "建议": "建议禁用",
                "hold": False,
                "status": "open",
            },
        ]
        preview = preview_disable_from_actions(rows, force_hold=False)
        self.assertEqual(preview["channel_count"]["uat"], 1)
        self.assertIn("old\tg", preview["channels"]["uat"])
        forced = preview_disable_from_actions(rows, force_hold=True)
        self.assertEqual(forced["channel_count"]["uat"], 2)

    def test_action_key_stable(self):
        row = {"站点": "test", "渠道名称": "a", "分组": "g", "类别": "下游无上游"}
        self.assertEqual(action_key(row, index=3), "3|test|a||g|下游无上游")
        self.assertTrue(is_hold_suggestion("部署未满 10 天，暂缓禁用/停用"))


if __name__ == "__main__":
    unittest.main()
