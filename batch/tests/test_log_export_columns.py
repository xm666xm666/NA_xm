"""使用日志导出会量字段。"""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from lib.log_exporter import (  # noqa: E402
    EXPORT_COLUMNS_ADMIN,
    EXPORT_COLUMNS_USER,
    LogExportCancelled,
    _drop_empty_columns,
    _rows_to_table,
    fetch_all_usage_logs,
    format_log_row,
    export_usage_logs_xlsx,
)
from lib.api_client import _parse_log_list_payload  # noqa: E402


SAMPLE_LOG = {
    "id": 1025,
    "user_id": 3,
    "created_at": 1700000000,
    "type": 2,
    "content": "consume detail",
    "username": "alice",
    "token_name": "tok-a",
    "token_id": 9,
    "model_name": "alias-m",
    "quota": 500000,
    "prompt_tokens": 10,
    "completion_tokens": 20,
    "use_time": 3,
    "is_stream": True,
    "channel": 7,
    "channel_name": "ch-demo",
    "group": "demo-uat",
    "ip": "1.2.3.4",
    "other": {
        "model_ratio": 1.5,
        "group_ratio": 1.0,
        "completion_ratio": 2.0,
        "cache_tokens": 4,
        "cache_ratio": 0.5,
        "model_price": -1,
        "user_group_ratio": -1,
        "frt": 1200,
        "is_model_mapped": True,
        "upstream_model_name": "upstream-m",
        "request_path": "/v1/chat/completions",
        "claude": False,
        "admin_info": {
            "use_channel": ["7", "8"],
            "is_multi_key": True,
            "multi_key_index": 1,
            "local_count_tokens": True,
        },
    },
}


def test_admin_row_has_core_and_other_fields():
    row = format_log_row(SAMPLE_LOG, is_admin=True, quota_per_unit=500_000)
    assert row["id"] == 1025
    assert row["user_id"] == 3
    assert row["username"] == "alice"
    assert row["channel_name"] == "ch-demo"
    assert row["token_id"] == 9
    assert row["quota"] == 500000
    assert row["cost_usd"] == "1.000000"
    assert row["model_name"] == "alias-m"
    assert "upstream-m" in row["model_name_display"]
    assert row["model_ratio"] == 1.5
    assert row["cache_tokens"] == 4
    assert row["frt"] == 1200
    assert row["request_path"] == "/v1/chat/completions"
    assert row["admin_use_channel"] == "7->8"
    assert row["admin_is_multi_key"] == "是"
    assert row["other_raw"]
    for _label, key in EXPORT_COLUMNS_ADMIN:
        assert key in row, f"missing key {key}"


def test_user_row_hides_admin_fields():
    row = format_log_row(SAMPLE_LOG, is_admin=False, quota_per_unit=500_000)
    assert row["username"] == ""
    assert row["channel_name"] == ""
    assert row["admin_use_channel"] == ""
    assert row["model_ratio"] == 1.5
    for _label, key in EXPORT_COLUMNS_USER:
        assert key in row, f"missing key {key}"


def test_xlsx_export_headers():
    data = export_usage_logs_xlsx([SAMPLE_LOG], is_admin=True, quota_per_unit=500_000)
    assert data[:2] == b"PK"
    from openpyxl import load_workbook
    import io

    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    headers = [c.value for c in ws[1]]
    assert "日志ID" in headers
    assert "额度quota" in headers
    assert "other原文" in headers
    assert "重试渠道链" in headers
    assert "上游模型" in headers
    a1 = ws.cell(row=1, column=1)
    assert a1.font.bold
    assert ws.max_row == 2
    id_idx = headers.index("日志ID") + 1
    assert ws.cell(row=2, column=id_idx).value == 1025


def test_xlsx_write_only_many_rows_is_fast():
    """大批量不再逐格套边框；两千行应在数秒内写完。"""
    import time

    logs = []
    for i in range(2000):
        item = dict(SAMPLE_LOG)
        item["id"] = i
        item["other"] = dict(SAMPLE_LOG["other"])
        logs.append(item)
    started = time.perf_counter()
    data = export_usage_logs_xlsx(logs, is_admin=True, quota_per_unit=500_000)
    elapsed = time.perf_counter() - started
    assert data[:2] == b"PK"
    assert elapsed < 8.0


def test_xlsx_saves_to_dest_path(tmp_path=None):
    import tempfile
    from pathlib import Path
    from openpyxl import load_workbook

    with tempfile.TemporaryDirectory() as tmp:
        dest = Path(tmp) / "usage_logs_fast.xlsx"
        payload = export_usage_logs_xlsx(
            [SAMPLE_LOG],
            is_admin=True,
            quota_per_unit=500_000,
            dest_path=dest,
        )
        assert payload == b""
        assert dest.is_file()
        wb = load_workbook(dest)
        assert wb.active.max_row == 2


def test_drop_empty_columns():
    headers = ["有值", "全空", "也有"]
    rows = [["a", "", "1"], ["b", "  ", "2"]]
    h2, r2 = _drop_empty_columns(headers, rows)
    assert h2 == ["有值", "也有"]
    assert r2 == [["a", "1"], ["b", "2"]]


def test_rows_to_table_omits_blank_other_fields():
    # 无 Claude / 音频等扩展时，对应列应被剔除
    sparse = {
        "id": 1,
        "user_id": 1,
        "created_at": 1700000000,
        "type": 2,
        "content": "x",
        "username": "u",
        "token_name": "t",
        "token_id": 1,
        "model_name": "m",
        "quota": 100,
        "prompt_tokens": 1,
        "completion_tokens": 2,
        "use_time": 1,
        "is_stream": False,
        "channel": 1,
        "channel_name": "c",
        "group": "g",
        "ip": "",
        "other": {"request_path": "/v1/chat/completions"},
    }
    headers, rows = _rows_to_table([sparse], is_admin=True, quota_per_unit=500_000)
    assert "请求路径" in headers
    assert "Claude" not in headers
    assert "缓存创建Tokens" not in headers
    assert "音频输入Tokens" not in headers
    path_idx = headers.index("请求路径")
    assert rows[0][path_idx] == "/v1/chat/completions"


def test_fetch_stops_when_cancelled():
    calls = {"n": 0}

    class FakeClient:
        def list_logs(self, **kwargs):
            calls["n"] += 1
            return {"items": [{"id": calls["n"]}], "total": 1000}

    def cancel():
        return calls["n"] >= 2

    try:
        fetch_all_usage_logs(FakeClient(), is_admin=True, page_delay=0, cancel=cancel)
        raise AssertionError("应在终止后抛出 LogExportCancelled")
    except LogExportCancelled:
        pass
    assert calls["n"] == 2


def test_fetch_continues_when_total_missing():
    """total 缺失/为 0 时不能只拉一页。"""
    pages = {
        1: {"items": [{"id": i} for i in range(100)], "total": 0, "page_size": 100},
        2: {"items": [{"id": i} for i in range(100, 200)], "total": None, "page_size": 100},
        3: {"items": [{"id": i} for i in range(200, 230)], "page_size": 100},
    }

    class FakeClient:
        def list_logs(self, **kwargs):
            return pages[kwargs["page"]]

    items = fetch_all_usage_logs(FakeClient(), is_admin=True, page_delay=0)
    assert len(items) == 230


def test_fetch_stops_at_reported_total():
    class FakeClient:
        def list_logs(self, **kwargs):
            start = (kwargs["page"] - 1) * 100
            return {
                "items": [{"id": i} for i in range(start, start + 100)],
                "total": 250,
                "page_size": 100,
            }

    items = fetch_all_usage_logs(FakeClient(), is_admin=True, page_delay=0)
    assert len(items) == 250


def test_fetch_filters_model_name_locally():
    class FakeClient:
        def list_logs(self, **kwargs):
            assert "model_name" not in kwargs
            return {
                "items": [
                    {"id": 1, "model_name": "deepseek-v3.2"},
                    {"id": 2, "model_name": "glm-4.7"},
                ],
                "total": 2,
                "page_size": 100,
            }

    stats = {}
    items = fetch_all_usage_logs(
        FakeClient(),
        is_admin=True,
        page_delay=0,
        filters={"model_name": "deepseek-v3.2", "group": "demo"},
        stats=stats,
    )
    assert [row["id"] for row in items] == [1]
    assert stats["gateway_total"] == 2
    assert stats["matched"] == 1
    assert stats["model_query"] == "deepseek-v3.2"


def test_parse_log_list_keeps_missing_total_as_none():
    parsed = _parse_log_list_payload(
        {"success": True, "data": {"items": [{"id": 1}], "page_size": 100}},
        page=1,
        page_size=100,
    )
    assert parsed["total"] is None
    assert parsed["items"] == [{"id": 1}]


if __name__ == "__main__":
    test_admin_row_has_core_and_other_fields()
    test_user_row_hides_admin_fields()
    test_xlsx_export_headers()
    test_xlsx_write_only_many_rows_is_fast()
    test_xlsx_saves_to_dest_path()
    test_drop_empty_columns()
    test_rows_to_table_omits_blank_other_fields()
    test_fetch_stops_when_cancelled()
    test_fetch_continues_when_total_missing()
    test_fetch_stops_at_reported_total()
    test_fetch_filters_model_name_locally()
    test_parse_log_list_keeps_missing_total_as_none()
    print("ok")
