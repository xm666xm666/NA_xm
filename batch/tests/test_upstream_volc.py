# -*- coding: utf-8 -*-
"""上游火山云导出格式单元测试（无网络）。"""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from lib.upstream_exporter import SHEET_VOLC, save_upstream_workbook, upstream_workbook_stats
from lib.upstream_volc import (
    COL_API_KEY,
    COL_API_KEY_CATALOG,
    UPSTREAM_HEADERS,
    apply_real_token_overlay,
    detect_env_tag,
    endpoint_to_row,
    extract_project_owner,
    format_api_key_catalog,
    format_model_service,
    format_project_party,
    load_real_tokens_by_deploy_id,
    match_api_key_for_endpoint,
    normalize_token_uuid,
    resolve_endpoint_token,
    resolve_rpm_tpm,
)


class UpstreamVolcFormatTests(unittest.TestCase):
    def test_format_model_service(self):
        text = format_model_service(
            {
                "ModelReference": {
                    "FoundationModel": {"Name": "glm-4-7", "ModelVersion": "251222"}
                }
            }
        )
        self.assertEqual(text, "glm-4-7｜251222")

    def test_detect_env_tag(self):
        self.assertEqual(detect_env_tag("Doubao-Seed-1.6-pvcopilot-uat"), "uat")
        self.assertEqual(detect_env_tag("xxx-prd"), "prd")
        self.assertEqual(detect_env_tag("finance-prod"), "prd")
        self.assertEqual(detect_env_tag("token-dev"), "uat")
        self.assertEqual(detect_env_tag("svc-test"), "test")
        self.assertEqual(detect_env_tag("model-demo"), "test")
        self.assertEqual(detect_env_tag("phivb-glm47"), "")

    def test_match_api_key_by_env(self):
        keys = [
            {"Id": 1, "Name": "uat-key-20260528142600"},
            {"Id": 2, "Name": "token-prd"},
            {"Id": 3, "Name": "token-dev"},
            {"Id": 4, "Name": "test-key"},
            {"Id": 5, "Name": "demo-token"},
        ]
        hit = match_api_key_for_endpoint("Doubao-Seed-1.6-pvcopilot-uat", keys)
        self.assertEqual(hit["Id"], 1)
        hit = match_api_key_for_endpoint("foo-prd", keys)
        self.assertEqual(hit["Id"], 2)
        hit = match_api_key_for_endpoint("kimi-dev", [{"Id": 2, "Name": "token-prd"}, {"Id": 3, "Name": "token-dev"}])
        self.assertEqual(hit["Id"], 3)
        hit = match_api_key_for_endpoint("glm-test", keys)
        self.assertEqual(hit["Id"], 4)
        hit = match_api_key_for_endpoint(
            "bar-demo",
            [{"Id": 2, "Name": "token-prd"}, {"Id": 5, "Name": "demo-token"}],
        )
        self.assertEqual(hit["Id"], 5)

    def test_match_single_key_fallback(self):
        keys = [{"Id": 9, "Name": "phivb-only"}]
        hit = match_api_key_for_endpoint("Doubao-Seed-2.0-pro-ivb", keys)
        self.assertEqual(hit["Id"], 9)
        hit = match_api_key_for_endpoint("anything-prd", keys)
        self.assertEqual(hit["Id"], 9)

    def test_match_unmatched_falls_back_to_default_key(self):
        keys = [{"Id": 1, "Name": "alpha"}, {"Id": 2, "Name": "beta"}]
        self.assertEqual(match_api_key_for_endpoint("phivb-glm47", keys)["Id"], 1)
        self.assertEqual(match_api_key_for_endpoint("model-uat", keys)["Id"], 1)
        self.assertEqual(
            match_api_key_for_endpoint(
                "phivb-glm47",
                keys,
                raw_key_by_id={1: "", 2: "ark-plain"},
            )["Id"],
            2,
        )

    def test_resolve_token_uses_default_plain_key(self):
        keys = [{"Id": 1, "Name": "alpha"}, {"Id": 2, "Name": "beta"}]
        token, name = resolve_endpoint_token("no-env-ep", keys, {1: "", 2: "ark-plain"})
        self.assertEqual(token, "ark-plain")
        self.assertEqual(name, "beta")

    def test_match_test_falls_back_to_uat_not_prd(self):
        keys = [
            {"Id": 1, "Name": "token-prd"},
            {"Id": 2, "Name": "token-dev"},
        ]
        self.assertEqual(match_api_key_for_endpoint("glm-5.2-fanyi-test", keys)["Id"], 2)
        self.assertEqual(match_api_key_for_endpoint("glm-5.2-fanyi-uat", keys)["Id"], 2)

    def test_dev_does_not_steal_test_key(self):
        keys = [
            {"Id": 1, "Name": "clinai"},
            {"Id": 9, "Name": "test-key"},
        ]
        self.assertEqual(match_api_key_for_endpoint("glm-4.7-clinai-dev", keys)["Id"], 1)
        self.assertEqual(match_api_key_for_endpoint("glm-4.7-clinai-prd", keys)["Id"], 1)

    def test_partial_split_old_test_keeps_generic_newest_uses_test_key(self):
        keys = [
            {"Id": 1, "Name": "clinai"},
            {"Id": 9, "Name": "test-key"},
        ]
        siblings = [
            ("doubao-seed-1.6-vision-clinai-test", "ep-20260430111318-76cxh"),
            ("Doubao-Seed-2.0-pro-clinai-test", "ep-20260430111431-2kkbm"),
            ("glm-5.2-clinai-test", "ep-20260820141814-ksmwh"),
        ]
        self.assertEqual(
            match_api_key_for_endpoint(
                "doubao-seed-1.6-vision-clinai-test",
                keys,
                endpoint_id="ep-20260430111318-76cxh",
                sibling_endpoints=siblings,
            )["Id"],
            1,
        )
        self.assertEqual(
            match_api_key_for_endpoint(
                "glm-5.2-clinai-test",
                keys,
                endpoint_id="ep-20260820141814-ksmwh",
                sibling_endpoints=siblings,
            )["Id"],
            9,
        )

    def test_pick_default_prefers_oldest_plain_key(self):
        keys = [{"Id": 20, "Name": "newer"}, {"Id": 3, "Name": "older"}]
        raw = {20: "ark-new", 3: "ark-old"}
        self.assertEqual(match_api_key_for_endpoint("no-env", keys, raw)["Id"], 3)

    def test_normalize_token_uuid(self):
        expected_uuid = "46f87179-1b5d-48e2-a7d4-9456514abe42"
        self.assertEqual(
            normalize_token_uuid("ark-" + expected_uuid),
            expected_uuid,
        )
        self.assertEqual(
            normalize_token_uuid("9ef24e14-8543-4ad8-bfff-a806f6017d85"),
            "9ef24e14-8543-4ad8-bfff-a806f6017d85",
        )

    def test_apply_real_token_overlay_by_deploy_id(self):
        rows = [
            {"部署名": "ep-20260302095255-vxrqq", COL_API_KEY: "ark-wrong"},
            {"部署名": "ep-only-upstream", COL_API_KEY: "ark-keep"},
        ]
        changed = apply_real_token_overlay(
            rows,
            {"ep-20260302095255-vxrqq": "9ef24e14-8543-4ad8-bfff-a806f6017d85"},
        )
        self.assertEqual(changed, 1)
        self.assertEqual(rows[0][COL_API_KEY], "9ef24e14-8543-4ad8-bfff-a806f6017d85")
        self.assertEqual(rows[1][COL_API_KEY], "ark-keep")

    def test_format_api_key_catalog(self):
        text = format_api_key_catalog(
            [
                {"Id": 2, "Name": "token-prd"},
                {"Id": 1, "Name": "uat-key-20260528142600"},
                {"Id": 3, "Name": "token-dev"},
            ],
            {
                1: "ark-test-key",
                2: "52988b4a-0c45-459e-9e19-81773379597e",
                3: "1524dacf-11c6-49d2-9b1f-a3426b608eaa",
            },
        )
        self.assertEqual(
            text,
            "\n".join(
                [
                    "uat-key-20260528142600 | ark-test-key",
                    "token-prd | 52988b4a-0c45-459e-9e19-81773379597e",
                    "token-dev | 1524dacf-11c6-49d2-9b1f-a3426b608eaa",
                ]
            ),
        )

    def test_load_real_tokens_by_deploy_id(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "火山真实数据.xlsx"
            save_upstream_workbook(
                path,
                [
                    endpoint_to_row(
                        {"Name": "glm-4.7-clinai-dev", "Id": "ep-20260302095255-vxrqq"},
                        project_name="clinai",
                        token="9ef24e14-8543-4ad8-bfff-a806f6017d85",
                    )
                ],
            )
            mapping = load_real_tokens_by_deploy_id(path)
            self.assertEqual(
                mapping["ep-20260302095255-vxrqq"],
                "9ef24e14-8543-4ad8-bfff-a806f6017d85",
            )

    def test_match_merged_uat_when_only_dev_key(self):
        keys = [{"Id": 2, "Name": "token-dev"}, {"Id": 3, "Name": "token-prd"}]
        self.assertEqual(match_api_key_for_endpoint("model-uat", keys)["Id"], 2)
        self.assertEqual(match_api_key_for_endpoint("model-dev", keys)["Id"], 2)

    def test_match_split_uat_dev_when_both_keys_exist(self):
        keys = [
            {"Id": 1, "Name": "token-uat"},
            {"Id": 2, "Name": "token-dev"},
            {"Id": 3, "Name": "token-prd"},
        ]
        self.assertEqual(match_api_key_for_endpoint("model-uat", keys)["Id"], 1)
        self.assertEqual(match_api_key_for_endpoint("model-dev", keys)["Id"], 2)
        self.assertEqual(match_api_key_for_endpoint("model-prd", keys)["Id"], 3)

    def test_resolve_rpm_tpm_from_model_map(self):
        endpoint = {
            "Name": "phivb-glm47",
            "ModelReference": {"FoundationModel": {"Name": "glm-4-7"}},
        }
        rpm, tpm = resolve_rpm_tpm(
            endpoint,
            model_rate_map={"glm-4-7": {"Rpm": 15000, "Tpm": 1500000}},
        )
        self.assertEqual(rpm, 15000)
        self.assertEqual(tpm, 1500000)

    def test_resolve_rpm_tpm_prefers_positive_endpoint_limit(self):
        endpoint = {
            "Name": "custom",
            "RateLimit": {"Rpm": 500, "Tpm": 1000000},
            "ModelReference": {"FoundationModel": {"Name": "glm-4-7"}},
        }
        rpm, tpm = resolve_rpm_tpm(
            endpoint,
            model_rate_map={"glm-4-7": {"Rpm": 15000, "Tpm": 1500000}},
        )
        self.assertEqual(rpm, 500)
        self.assertEqual(tpm, 1000000)

    def test_resolve_rpm_tpm_ignores_minus_one(self):
        endpoint = {
            "RateLimit": {"Rpm": -1, "Tpm": -1},
            "ModelReference": {"FoundationModel": {"Name": "glm-4-7"}},
        }
        rpm, tpm = resolve_rpm_tpm(
            endpoint,
            model_rate_map={"glm-4-7": {"Rpm": 15000, "Tpm": 1500000}},
        )
        self.assertEqual(rpm, 15000)
        self.assertEqual(tpm, 1500000)

    def test_endpoint_to_row_headers(self):
        row = endpoint_to_row(
            {
                "Name": "phivb-glm47",
                "Id": "ep-20260424142638-zxct2",
                "ProjectName": "phivb",
                "ModelReference": {
                    "FoundationModel": {"Name": "glm-4-7", "ModelVersion": "251222"}
                },
            },
            project_name="phivb",
            token="apikey-demo",
            model_rate_map={"glm-4-7": {"Rpm": 15000, "Tpm": 1500000}},
        )
        self.assertEqual(list(row.keys()), UPSTREAM_HEADERS)
        self.assertLess(UPSTREAM_HEADERS.index("TPM"), UPSTREAM_HEADERS.index("RPM"))
        self.assertEqual(row["接入点名称"], "phivb-glm47")
        self.assertEqual(row["部署名"], "ep-20260424142638-zxct2")
        self.assertEqual(row[COL_API_KEY], "apikey-demo")
        self.assertEqual(row[COL_API_KEY_CATALOG], "")
        self.assertEqual(row["RPM"], 15000)
        self.assertEqual(row["TPM"], 1500000)
        self.assertEqual(row["项目方"], "phivb-N/A")
        self.assertEqual(row["联系人"], "N/A")
        self.assertEqual(row["部署时间"], "2026-04-24 14:26:38")
        self.assertEqual(UPSTREAM_HEADERS[-1], "部署时间")

    def test_endpoint_to_row_uses_create_time(self):
        row = endpoint_to_row(
            {
                "Name": "a",
                "Id": "ep-other",
                "CreateTime": "2026-01-15T09:30:00+08:00",
            },
            project_name="p",
        )
        self.assertEqual(row["部署时间"], "2026-01-15 09:30:00")

    def test_format_project_party_with_owner(self):
        self.assertEqual(format_project_party("phivb", "phivb-范亚浩"), "phivb-范亚浩")
        self.assertEqual(extract_project_owner("phivb", "phivb-范亚浩"), "范亚浩")
        self.assertEqual(format_project_party("smo-rpm", "smo-rpm-金亥其"), "smo-rpm-金亥其")
        self.assertEqual(format_project_party("pv", "孙刘淘"), "pv-孙刘淘")

    def test_format_project_party_missing_owner_is_na(self):
        self.assertEqual(format_project_party("AI-Drug-Research", ""), "AI-Drug-Research-N/A")
        self.assertEqual(
            format_project_party("AI-Drug-Research", "AI-Drug-Research"),
            "AI-Drug-Research-N/A",
        )
        self.assertEqual(format_project_party("default", "默认项目"), "default-N/A")
        self.assertEqual(extract_project_owner("AI-Drug-Research", ""), "N/A")
        self.assertEqual(
            format_project_party("AI-Drug-Research", "AI-Drug-Research-N/A"),
            "AI-Drug-Research-N/A",
        )

    def test_endpoint_to_row_uses_display_owner(self):
        row = endpoint_to_row(
            {"Name": "a", "Id": "ep-1", "ProjectName": "phivb"},
            project_name="phivb",
            display_name="phivb-范亚浩",
        )
        self.assertEqual(row["项目方"], "phivb-范亚浩")
        self.assertEqual(row["联系人"], "范亚浩")

    def test_save_workbook_independent_sheet(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "上游信息_火山云_test.xlsx"
            rows = [
                endpoint_to_row(
                    {
                        "Name": "a",
                        "Id": "ep-1",
                        "ProjectName": "p1",
                        "ModelReference": {
                            "FoundationModel": {"Name": "m", "ModelVersion": "v"}
                        },
                    },
                    project_name="p1",
                )
            ]
            save_upstream_workbook(path, rows)
            stats = upstream_workbook_stats(path)
            self.assertEqual(stats["sheet"], SHEET_VOLC)
            self.assertEqual(stats["rows"], 1)
            from openpyxl import load_workbook
            from openpyxl.utils import get_column_letter

            wb = load_workbook(path)
            try:
                ws = wb[SHEET_VOLC]
                last_col = get_column_letter(ws.max_column)
                self.assertEqual(ws.auto_filter.ref, f"A1:{last_col}{ws.max_row}")
            finally:
                wb.close()


if __name__ == "__main__":
    unittest.main()
