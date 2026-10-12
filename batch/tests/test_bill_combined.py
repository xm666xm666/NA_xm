# -*- coding: utf-8 -*-
"""火山云 + Azure 账单汇总工作簿（无网络）。"""

from __future__ import annotations

import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from lib.azure_billing import AzureBillStats
from lib.bill_combined import (
    COMBINED_STEM,
    SHEET_AZURE_DAY,
    SHEET_VOLC_DAY,
    build_combined_bill_filename,
    save_combined_bill_workbook,
)
from lib.volc_billing import BillCollectStats, SUM_LABEL


class CombinedBillTests(unittest.TestCase):
    def test_combined_workbook_sheets_sum_and_yen(self):
        with tempfile.TemporaryDirectory() as tmp:
            dest = Path(tmp) / f"{COMBINED_STEM}_20260901120000.xlsx"
            volc = BillCollectStats(
                bill_period="2026-08",
                source_action="ListSplitBillDetail",
                rows_day=[
                    {
                        "日期": "2026-08-01",
                        "项目名": "p",
                        "产品": "方舟",
                        "配置名称": "c",
                        "分拆项ID": "x",
                        "费用": 1.5,
                    }
                ],
                rows_month=[
                    {
                        "日期": "2026-08",
                        "项目名": "p",
                        "产品": "方舟",
                        "配置名称": "c",
                        "分拆项ID": "x",
                        "费用": 1.5,
                    }
                ],
            )
            azure = AzureBillStats(
                bill_period="2026-08",
                rows_day=[
                    {
                        "日期": "2026-08-02",
                        "项目名": "临床AI",
                        "项目负责人": "张三",
                        "模型名": "gpt-4o",
                        "配置名称": "acct",
                        "地区": "eastus",
                        "订阅": "订阅1",
                        "费用": 12.5,
                    }
                ],
                rows_month=[
                    {
                        "日期": "2026-08",
                        "项目名": "临床AI",
                        "项目负责人": "张三",
                        "模型名": "gpt-4o",
                        "配置名称": "acct",
                        "地区": "eastus",
                        "订阅": "订阅1",
                        "费用": 12.5,
                    }
                ],
            )
            save_combined_bill_workbook(dest, volc, azure)
            from openpyxl import load_workbook

            wb = load_workbook(dest)
            try:
                self.assertIn(SHEET_VOLC_DAY, wb.sheetnames)
                self.assertIn(SHEET_AZURE_DAY, wb.sheetnames)
                volc_headers = [c.value for c in wb[SHEET_VOLC_DAY][1]]
                self.assertEqual(volc_headers[-1], "费用（¥）")
                self.assertEqual(wb[SHEET_VOLC_DAY]["A3"].value, SUM_LABEL)
                self.assertEqual(wb[SHEET_VOLC_DAY]["F3"].value, "=SUM(F2:F2)")
                azure_headers = [c.value for c in wb[SHEET_AZURE_DAY][1]]
                self.assertEqual(azure_headers[-1], "费用（¥）")
                self.assertEqual(wb[SHEET_AZURE_DAY]["A3"].value, SUM_LABEL)
                self.assertEqual(wb[SHEET_AZURE_DAY]["H3"].value, "=SUM(H2:H2)")
                self.assertEqual(wb[SHEET_VOLC_DAY].auto_filter.ref, "A1:F2")
                self.assertEqual(wb[SHEET_AZURE_DAY].auto_filter.ref, "A1:H2")
            finally:
                wb.close()

    def test_combined_filename_includes_month(self):
        name = build_combined_bill_filename(
            when=datetime(2026, 9, 16, 10, 28, 22),
            bill_period="2026-09",
        )
        self.assertEqual(name, f"{COMBINED_STEM}_9月_20260916102822.xlsx")


if __name__ == "__main__":
    unittest.main()
