"""渠道余额查询单元测试。"""

from __future__ import annotations

import io
import sys
import unittest
from pathlib import Path

BATCH_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BATCH_DIR))

from lib.channel_balance_query import (  # noqa: E402
    BALANCE_EXCEL_STEM,
    TIME_PRESET_DAYS,
    attach_activity_to_rows,
    balance_sheet_title,
    channel_balance_row,
    export_channel_balance_xlsx,
    latest_balance_excel_path,
    filter_balance_rows,
    format_usd,
    in_range,
    parse_query_timestamp,
    resolve_created_time_range,
    resolve_range,
    used_quota_to_usd,
    USED_RANGE_PRESETS,
)


class ChannelBalanceQueryTests(unittest.TestCase):
    def test_used_quota_to_usd(self):
        self.assertAlmostEqual(used_quota_to_usd(500_000, 500_000), 1.0)
        self.assertAlmostEqual(used_quota_to_usd(1_380_000, 500_000), 2.76)

    def test_format_usd(self):
        self.assertEqual(format_usd(0), "$0")
        self.assertEqual(format_usd(2.76), "$2.76")

    def test_in_range_eq0(self):
        self.assertTrue(in_range(0.0, 0.0, 0.0))
        self.assertFalse(in_range(0.01, 0.0, 0.0))

    def test_resolve_custom_overrides_preset(self):
        mn, mx = resolve_range("ge100", USED_RANGE_PRESETS, 5.0, 20.0)
        self.assertEqual(mn, 5.0)
        self.assertEqual(mx, 20.0)

    def test_le1_le5_presets(self):
        from lib.channel_balance_query import REMAIN_RANGE_PRESETS

        self.assertEqual(USED_RANGE_PRESETS["le1"], (None, 1.0))
        self.assertEqual(USED_RANGE_PRESETS["le5"], (None, 5.0))
        self.assertEqual(REMAIN_RANGE_PRESETS["le1"], (None, 1.0))
        self.assertEqual(REMAIN_RANGE_PRESETS["le5"], (None, 5.0))
        self.assertTrue(in_range(1.0, *USED_RANGE_PRESETS["le1"]))
        self.assertTrue(in_range(5.0, *USED_RANGE_PRESETS["le5"]))
        self.assertFalse(in_range(1.01, *USED_RANGE_PRESETS["le1"]))
        self.assertFalse(in_range(5.01, *USED_RANGE_PRESETS["le5"]))

    def test_created_time_label(self):
        row = channel_balance_row(
            {
                "id": 1,
                "name": "a",
                "group": "g",
                "type": 8,
                "status": 1,
                "used_quota": 0,
                "balance": 0,
                "created_time": 1700000000,
            },
            quota_per_unit=500_000,
        )
        self.assertEqual(row["created_time"], 1700000000)
        self.assertNotEqual(row["created_time_label"], "—")
        self.assertIn("-", row["created_time_label"])

    def test_export_xlsx(self):
        items = [
            channel_balance_row(
                {
                    "id": 9,
                    "name": "demo",
                    "group": "g1",
                    "type": 8,
                    "status": 1,
                    "used_quota": 500_000,
                    "balance": 1.5,
                    "created_time": 1700000000,
                },
                quota_per_unit=500_000,
                type_label="自定义渠道",
            )
        ]
        items[0]["activity_label"] = "gpt-4: 有活跃\nglm-5: 近一月无活跃"
        data, filename = export_channel_balance_xlsx(items, include_activity=True)
        self.assertTrue(filename.startswith(f"{BALANCE_EXCEL_STEM}_"))
        self.assertTrue(filename.endswith(".xlsx"))
        self.assertGreater(len(data), 100)
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data))
        try:
            ws = wb.active
            headers = [cell.value for cell in ws[1]]
            self.assertIn("近一月活跃", headers)
            idx = headers.index("近一月活跃")
            self.assertIn("gpt-4: 有活跃", str(ws.cell(2, idx + 1).value or ""))
            self.assertIn("glm-5: 近一月无活跃", str(ws.cell(2, idx + 1).value or ""))
        finally:
            wb.close()

    def test_attach_activity_multi_model_and_fetch_fail(self):
        class _OkClient:
            def fetch_channel_activity_stats(self, items, batch_size=500):
                return [
                    {
                        "channel_id": 9,
                        "model_name": "gpt-4",
                        "total_requests": 2,
                        "active_minutes": 1,
                        "avg_rpm": 0.1,
                        "avg_tpm": 10,
                    }
                ]

        class _FailClient:
            def fetch_channel_activity_stats(self, items, batch_size=500):
                raise RuntimeError("boom")

        rows = [
            {"id": 9, "models": "gpt-4, glm-5", "test_model": "gpt-4"},
            {"id": 8, "models": "", "test_model": ""},
        ]
        attach_activity_to_rows(_OkClient(), rows)
        self.assertIn("gpt-4: 有活跃", rows[0]["activity_label"])
        self.assertIn("glm-5: 近一月无活跃", rows[0]["activity_label"])
        self.assertEqual(rows[1]["activity_label"], "N/A")

        failed = [{"id": 1, "models": "a,b"}]
        attach_activity_to_rows(_FailClient(), failed)
        self.assertIn("a: N/A", failed[0]["activity_label"])
        self.assertIn("b: N/A", failed[0]["activity_label"])

    def test_multi_sheet_titles(self):
        self.assertEqual(balance_sheet_title("uat"), "uat-渠道余额")
        self.assertEqual(balance_sheet_title("test"), "test-渠道余额")
        items = [
            channel_balance_row(
                {"id": 1, "name": "a", "group": "g", "type": 8, "status": 1, "used_quota": 0, "balance": 0, "created_time": 1700000000},
                quota_per_unit=500_000,
            ),
            channel_balance_row(
                {"id": 2, "name": "b", "group": "g", "type": 8, "status": 1, "used_quota": 0, "balance": 1, "created_time": 1700000000},
                quota_per_unit=500_000,
            ),
        ]
        items[0]["env"] = "test"
        items[0]["site_label"] = "na-test"
        items[1]["env"] = "uat"
        items[1]["site_label"] = "na-uat"
        data, filename = export_channel_balance_xlsx(items, by_env=True)
        self.assertTrue(filename.endswith(".xlsx"))
        self.assertGreater(len(data), 200)
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data))
        self.assertIn("test-渠道余额", wb.sheetnames)
        self.assertIn("uat-渠道余额", wb.sheetnames)

    def test_filter_requires_selected_metrics(self):
        rows = [
            channel_balance_row(
                {"id": 1, "name": "a", "group": "g", "type": 8, "status": 1, "used_quota": 500_000, "balance": 0},
                quota_per_unit=500_000,
            ),
            channel_balance_row(
                {"id": 2, "name": "b", "group": "g", "type": 8, "status": 1, "used_quota": 0, "balance": 10},
                quota_per_unit=500_000,
            ),
        ]
        matched = filter_balance_rows(
            rows,
            query_used=True,
            query_remain=False,
            used_min=1.0,
            used_max=None,
            remain_min=None,
            remain_max=None,
        )
        self.assertEqual([r["id"] for r in matched], [1])

        matched2 = filter_balance_rows(
            rows,
            query_used=False,
            query_remain=True,
            used_min=None,
            used_max=None,
            remain_min=0.0,
            remain_max=0.0,
        )
        self.assertEqual([r["id"] for r in matched2], [1])

    def test_time_preset_days(self):
        self.assertEqual(list(TIME_PRESET_DAYS), ["all", "1d", "3d", "5d", "15d", "30d", "60d", "180d", "365d"])
        self.assertIsNone(TIME_PRESET_DAYS["all"])
        self.assertEqual(TIME_PRESET_DAYS["1d"], 1)
        self.assertEqual(TIME_PRESET_DAYS["365d"], 365)

    def test_resolve_created_time_range_all(self):
        self.assertEqual(resolve_created_time_range("all"), (None, None))

    def test_resolve_created_time_range_preset(self):
        from datetime import datetime

        now = datetime(2026, 8, 26, 15, 30, 0)
        start_ts, end_ts = resolve_created_time_range("3d", now=now)
        self.assertEqual(datetime.fromtimestamp(start_ts), datetime(2026, 8, 24, 0, 0, 0))
        self.assertEqual(datetime.fromtimestamp(end_ts), datetime(2026, 8, 26, 23, 59, 59))
        start1, end1 = resolve_created_time_range("1d", now=now)
        self.assertEqual(datetime.fromtimestamp(start1), datetime(2026, 8, 26, 0, 0, 0))
        self.assertEqual(end1, end_ts)

    def test_resolve_created_time_range_custom_overrides(self):
        start_ts, end_ts = resolve_created_time_range("30d", 1700000000, 1700003600)
        self.assertEqual((start_ts, end_ts), (1700000000, 1700003600))
        with self.assertRaises(ValueError):
            resolve_created_time_range("all", 200, 100)

    def test_parse_query_timestamp(self):
        self.assertIsNone(parse_query_timestamp(""))
        self.assertIsNone(parse_query_timestamp(0))
        self.assertEqual(parse_query_timestamp(1700000000), 1700000000)
        self.assertEqual(parse_query_timestamp("1700000000"), 1700000000)

    def test_filter_by_created_time(self):
        rows = [
            channel_balance_row(
                {"id": 1, "name": "old", "group": "g", "type": 8, "status": 1, "used_quota": 0, "balance": 1, "created_time": 1_700_000_000},
                quota_per_unit=500_000,
            ),
            channel_balance_row(
                {"id": 2, "name": "new", "group": "g", "type": 8, "status": 1, "used_quota": 0, "balance": 1, "created_time": 1_800_000_000},
                quota_per_unit=500_000,
            ),
            channel_balance_row(
                {"id": 3, "name": "none", "group": "g", "type": 8, "status": 1, "used_quota": 0, "balance": 1, "created_time": 0},
                quota_per_unit=500_000,
            ),
        ]
        matched = filter_balance_rows(
            rows,
            query_used=False,
            query_remain=False,
            used_min=None,
            used_max=None,
            remain_min=None,
            remain_max=None,
            start_ts=1_750_000_000,
            end_ts=1_900_000_000,
        )
        self.assertEqual([r["id"] for r in matched], [2])

    def test_latest_balance_excel_path(self):
        import os
        import tempfile

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            excel = root / "excel"
            excel.mkdir()
            older = excel / "渠道余额_20200101000000.xlsx"
            newer = excel / "渠道余额_20260101000000.xlsx"
            legacy = excel / "channel_balance_20200101_000000.xlsx"
            older.write_bytes(b"old")
            newer.write_bytes(b"new")
            legacy.write_bytes(b"legacy")
            os.utime(older, (1, 1))
            os.utime(legacy, (2, 2))
            os.utime(newer, (3, 3))
            self.assertEqual(latest_balance_excel_path(root).resolve(), newer.resolve())
            self.assertIsNone(latest_balance_excel_path(root / "missing"))


if __name__ == "__main__":
    unittest.main()
