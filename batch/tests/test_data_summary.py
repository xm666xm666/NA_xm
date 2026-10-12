"""数据汇总：字段映射与重复检测。"""

from __future__ import annotations

import sys
from pathlib import Path
from unittest.mock import patch

BATCH_DIR = Path(__file__).resolve().parents[1]
if str(BATCH_DIR) not in sys.path:
    sys.path.insert(0, str(BATCH_DIR))

from lib.data_summary import (  # noqa: E402
    ENVS,
    TOKEN_HEADERS,
    SiteConfig,
    _cross_env_duplicate_names,
    _exclude_duplicate_names,
    build_duplicate_channels,
    build_duplicate_tokens,
    channel_to_row,
    enrich_auth_from_env,
    enrich_site_from_env,
    parse_sites,
    token_to_row,
)


def test_token_to_row_has_all_headers():
    row = token_to_row(
        {
            "name": "demo-token",
            "group": "default",
            "status": 1,
            "unlimited_quota": True,
            "allow_ips": "",
            "created_time": 0,
            "expired_time": -1,
        },
        [],
    )
    assert list(row.keys()) == TOKEN_HEADERS
    assert row["项目环境"] == "default"
    assert row["分组"] == "default"


def test_token_to_row_project_env_matches_group():
    row = token_to_row(
        {
            "name": "财务小助手-费仕忆",
            "group": "doubao",
            "status": 1,
            "unlimited_quota": True,
            "allow_ips": "",
            "created_time": 0,
            "expired_time": -1,
        },
        [],
    )
    assert row["分组"] == "doubao"
    assert row["项目环境"] == row["分组"]


def test_channel_to_row_deploy_time_from_created_time():
    row = channel_to_row(
        {
            "name": "demo",
            "type": 1,
            "status": 1,
            "models": "m1,m2",
            "created_time": 1700000000,
            "remark": "管理员备注",
        },
        {1: "OpenAI"},
    )
    assert row["部署时间"]
    assert "-" in row["部署时间"]
    assert row["模型列表"] == "m1, m2"
    assert row["备注"] == "管理员备注"


def test_channel_to_row_deploy_time_empty_when_missing():
    row = channel_to_row({"name": "demo", "type": 1, "status": 1, "models": "m1,m2"}, {1: "OpenAI"})
    assert row["部署时间"] == ""
    assert row["模型列表"] == "m1, m2"
    assert row["响应时间"] == "未测试"


def test_channel_to_row_response_time_after_status():
    from lib.data_summary import CHANNEL_HEADERS, CHANNEL_HEADERS_WITH_KEY, format_channel_response_time

    assert CHANNEL_HEADERS.index("响应时间") == CHANNEL_HEADERS.index("状态") + 1
    assert CHANNEL_HEADERS_WITH_KEY.index("响应时间") == CHANNEL_HEADERS_WITH_KEY.index("状态") + 1
    assert CHANNEL_HEADERS.index("备注") == CHANNEL_HEADERS.index("base_url") + 1
    assert CHANNEL_HEADERS_WITH_KEY.index("备注") == CHANNEL_HEADERS_WITH_KEY.index("base_url") + 1
    assert format_channel_response_time(0) == "未测试"
    assert format_channel_response_time(None) == "未测试"
    assert format_channel_response_time(850) == "0.85 秒"
    row = channel_to_row(
        {"name": "demo", "type": 1, "status": 1, "models": "m1", "response_time": 1234},
        {1: "OpenAI"},
    )
    assert row["响应时间"] == "1.23 秒"


def test_duplicate_channels_across_envs():
    rows = {env: [] for env in ENVS}
    base = {
        "名称": "same-model",
        "分组": "g",
        "类型": "OpenAI",
        "状态": "已启用",
        "模型列表": "m",
        "模型重定向": "",
        "base_url": "https://x",
        "部署时间": "",
    }
    rows["test"].append(dict(base))
    rows["uat"].append(dict(base))
    dup = build_duplicate_channels(rows)
    assert len(dup) == 2
    assert {item["来源"] for item in dup} == {"test-渠道管理", "uat-渠道管理"}


def test_duplicate_tokens_across_envs():
    rows = {env: [] for env in ENVS}
    base = {
        "名称": "token-a",
        "项目环境": "deepseek",
        "状态": "已启用",
        "剩余额度": "无限额度",
        "分组": "deepseek",
        "可用模型": "无限制",
        "包含渠道": "无限制",
        "IP限制": "无限制",
        "创建时间": "",
        "过期时间": "永不过期",
    }
    rows["test"].append(dict(base))
    rows["prod"].append(dict(base))
    dup = build_duplicate_tokens(rows)
    assert len(dup) == 2
    assert dup[0]["重复序号"] == 1


def test_exclude_duplicates_from_main_sheets():
    rows = {env: [] for env in ENVS}
    base = {
        "名称": "shared",
        "分组": "g",
        "类型": "OpenAI",
        "状态": "已启用",
        "模型列表": "m",
        "模型重定向": "",
        "base_url": "https://x",
        "部署时间": "",
    }
    rows["test"].append(dict(base))
    rows["uat"].append(dict(base))
    rows["prod"].append({**base, "名称": "unique-prod"})

    dup_names = _cross_env_duplicate_names(rows)
    assert dup_names == {"shared"}

    filtered, skipped = _exclude_duplicate_names(rows, dup_names)
    assert skipped == 2
    assert all(row["名称"] != "shared" for row in filtered["test"])
    assert all(row["名称"] != "shared" for row in filtered["uat"])
    assert filtered["prod"][0]["名称"] == "unique-prod"


def test_merge_rows_keeps_same_name_different_ids():
    from lib.data_summary import _merge_rows

    base = {
        "名称": "dup-name",
        "分组": "g",
        "类型": "OpenAI",
        "状态": "已启用",
        "模型列表": "m",
        "模型重定向": "",
        "base_url": "https://x",
        "部署时间": "",
    }
    rows = [{**base, "_sort_id": 10}, {**base, "_sort_id": 9, "base_url": "https://y"}]
    merged, added = _merge_rows([], rows)
    assert len(merged) == 2
    assert [row["_sort_id"] for row in merged] == [9, 10]


def test_sort_rows_by_id_asc():
    from lib.data_summary import _sort_rows_by_id

    rows = [{"_sort_id": 3, "名称": "c"}, {"_sort_id": 99, "名称": "a"}, {"_sort_id": 50, "名称": "b"}]
    assert [row["_sort_id"] for row in _sort_rows_by_id(rows)] == [3, 50, 99]


def test_collect_keeps_cross_env_duplicates_in_main_sheet(tmp_path: Path):
    """主表保留各环境全量，跨环境同名仍写入重复模型表。"""
    from lib.data_summary import SiteConfig, run_data_summary_collect
    from lib.data_summary_exporter import load_workbook_data, save_workbook_data, SHEET_DUP_CHANNELS

    root = tmp_path
    excel_dir = root / "excel"
    excel_dir.mkdir()

    shared = {
        "名称": "shared-channel",
        "分组": "g",
        "类型": "OpenAI",
        "状态": "已启用",
        "模型列表": "m",
        "模型重定向": "",
        "base_url": "https://x",
        "部署时间": "",
    }
    unique_uat = {**shared, "名称": "uat-only"}

    def fake_fetch(site, *, type_labels, include_keys=False):
        env = site.env
        meta = {"keys_fetched": 0, "keys_failed": 0, "key_warnings": []}
        if env == "test":
            return env, [dict(shared)], [], meta
        if env == "uat":
            return env, [dict(shared), unique_uat], [], meta
        return env, [], [], meta

    with patch("lib.data_summary.fetch_site_data", fake_fetch), patch(
        "lib.data_summary.latest_excel_path", lambda _root: None
    ), patch("lib.data_summary.build_timestamped_path", lambda _root: excel_dir / "new.xlsx"):
        stats, _saved = run_data_summary_collect(
            sites=[
                SiteConfig(base_url="https://test.example.com", username="a", password="x", env="test"),
                SiteConfig(base_url="https://uat.example.com", username="a", password="x", env="uat"),
            ],
            excel_path=excel_dir / "new.xlsx",
        )

    channel_rows, _ = load_workbook_data(excel_dir / "new.xlsx")
    assert [row["名称"] for row in channel_rows["test"]] == ["shared-channel"]
    assert {row["名称"] for row in channel_rows["uat"]} == {"shared-channel", "uat-only"}
    assert stats.channels_in_sheet["uat"] == 2
    assert stats.channels_skipped_duplicate == 0
    assert stats.duplicate_channels == 1

    from openpyxl import load_workbook

    wb = load_workbook(excel_dir / "new.xlsx")
    ws = wb[SHEET_DUP_CHANNELS]
    dup_names = {ws.cell(row=i, column=1).value for i in range(2, ws.max_row + 1)}
    assert "shared-channel" in dup_names
    wb.close()


def test_collect_creates_new_file_not_latest(tmp_path: Path):
    """一键采集新建 Excel，不覆盖 excel/ 下已有最新文件。"""
    from lib.data_summary import EXCEL_STEM, SiteConfig, run_data_summary_collect
    from lib.data_summary_exporter import load_workbook_data, save_workbook_data

    root = tmp_path
    excel_dir = root / "excel"
    excel_dir.mkdir()
    old_path = excel_dir / f"{EXCEL_STEM}.xlsx"
    new_path = excel_dir / f"{EXCEL_STEM}_new.xlsx"

    stale_uat = {
        "名称": "stale-uat",
        "分组": "g",
        "类型": "OpenAI",
        "状态": "已启用",
        "模型列表": "m",
        "模型重定向": "",
        "base_url": "https://stale",
        "部署时间": "",
    }
    keep_test = {**stale_uat, "名称": "keep-test", "base_url": "https://test"}
    save_workbook_data(
        old_path,
        channel_env_rows={"test": [keep_test], "uat": [stale_uat], "prod": []},
        token_env_rows={env: [] for env in ENVS},
    )
    before_old = old_path.read_bytes()

    def fake_fetch(site, *, type_labels, include_keys=False):
        env = site.env or "uat"
        row = {
            "名称": "fresh-uat",
            "分组": "g",
            "类型": "OpenAI",
            "状态": "已启用",
            "模型列表": "m",
            "模型重定向": "",
            "base_url": site.base_url,
            "部署时间": "",
        }
        return env, [row], [], {"keys_fetched": 0, "keys_failed": 0, "key_warnings": []}

    with patch("lib.data_summary.fetch_site_data", fake_fetch), patch(
        "lib.data_summary.build_timestamped_path", lambda _root: new_path
    ):
        stats, saved = run_data_summary_collect(
            sites=[SiteConfig(base_url="https://uat.example.com", username="admin", password="x", env="uat")],
            excel_path=old_path,
        )

    assert saved == new_path
    assert new_path.exists()
    assert old_path.read_bytes() == before_old
    channel_rows, _token_rows = load_workbook_data(new_path)
    assert [row["名称"] for row in channel_rows["uat"]] == ["fresh-uat"]
    assert channel_rows["test"] == []
    assert channel_rows["prod"] == []
    assert stats.channels_in_sheet["uat"] == 1


def test_collect_session_output_path_merges_envs(tmp_path: Path):
    """同一次一键采集多站：第二站写入同一新文件并保留已采环境。"""
    from lib.data_summary import SiteConfig, run_data_summary_collect
    from lib.data_summary_exporter import load_workbook_data

    root = tmp_path
    excel_dir = root / "excel"
    excel_dir.mkdir()
    session_path = excel_dir / "session.xlsx"

    def fake_fetch(site, *, type_labels, include_keys=False):
        env = site.env
        row = {
            "名称": f"{env}-only",
            "分组": "g",
            "类型": "OpenAI",
            "状态": "已启用",
            "模型列表": "m",
            "模型重定向": "",
            "base_url": site.base_url,
            "部署时间": "",
        }
        meta = {"keys_fetched": 0, "keys_failed": 0, "key_warnings": []}
        return env, [row], [], meta

    with patch("lib.data_summary.fetch_site_data", fake_fetch), patch(
        "lib.data_summary.build_timestamped_path", lambda _root: session_path
    ):
        _stats1, saved1 = run_data_summary_collect(
            sites=[SiteConfig(base_url="https://test.example.com", username="a", password="x", env="test")],
            excel_path=session_path,
        )
        stats2, saved2 = run_data_summary_collect(
            sites=[SiteConfig(base_url="https://uat.example.com", username="a", password="x", env="uat")],
            excel_path=session_path,
            output_path=session_path,
        )

    assert saved1 == session_path
    assert saved2 == session_path
    channel_rows, _ = load_workbook_data(session_path)
    assert [row["名称"] for row in channel_rows["test"]] == ["test-only"]
    assert [row["名称"] for row in channel_rows["uat"]] == ["uat-only"]
    assert stats2.channels_in_sheet["test"] == 1
    assert stats2.channels_in_sheet["uat"] == 1


def test_collect_all_fail_does_not_write(tmp_path: Path):
    """全部站点失败时不新建/不覆盖 Excel。"""
    from lib.data_summary import EXCEL_STEM, SiteConfig, run_data_summary_collect
    from lib.data_summary_exporter import save_workbook_data

    root = tmp_path
    excel_dir = root / "excel"
    excel_dir.mkdir()
    old_path = excel_dir / f"{EXCEL_STEM}.xlsx"
    save_workbook_data(
        old_path,
        channel_env_rows={
            "test": [{
                "名称": "keep",
                "分组": "g",
                "类型": "OpenAI",
                "状态": "已启用",
                "模型列表": "m",
                "模型重定向": "",
                "base_url": "https://x",
                "部署时间": "",
            }],
            "uat": [],
            "prod": [],
        },
        token_env_rows={env: [] for env in ENVS},
    )
    before = old_path.read_bytes()

    def boom(*_a, **_k):
        raise RuntimeError("429 Too Many Requests")

    with patch("lib.data_summary.fetch_site_data", boom), patch(
        "lib.data_summary.build_timestamped_path", lambda _root: excel_dir / "should-not-create.xlsx"
    ):
        stats, saved = run_data_summary_collect(
            sites=[SiteConfig(base_url="https://test.example.com", username="a", password="x", env="test")],
            excel_path=old_path,
        )

    assert stats.sites_ok == 0
    assert saved == old_path
    assert old_path.read_bytes() == before
    assert not (excel_dir / "should-not-create.xlsx").exists()


SAMPLE_TOTP = "ZDNVM5SXU3MYGNGUKIVITQ4ZA5BL2RGO"
_ENV_SITE = {
    "test": {
        "base_url": "https://na-test.example.com",
        "username": "env-admin",
        "password": "from-env",
        "totp_secret": SAMPLE_TOTP,
        "default_key": "sk-env",
    }
}


def test_enrich_site_from_env_fills_password_and_totp():
    site = SiteConfig(base_url="https://na-test.example.com", username="", password="", env="test")
    with patch("lib.data_summary.load_raw_env_sites", return_value=_ENV_SITE):
        enrich_site_from_env(site)
    assert site.username == "env-admin"
    assert site.password == "from-env"
    assert site.totp_secret == SAMPLE_TOTP


def test_enrich_site_from_env_keeps_explicit_password():
    site = SiteConfig(
        base_url="https://na-test.example.com",
        username="page-user",
        password="page-pass",
        env="test",
    )
    with patch("lib.data_summary.load_raw_env_sites", return_value=_ENV_SITE):
        enrich_site_from_env(site)
    assert site.username == "page-user"
    assert site.password == "page-pass"
    assert site.totp_secret == SAMPLE_TOTP


def test_enrich_auth_from_env_fills_password_and_totp():
    with patch("lib.data_summary.load_raw_env_sites", return_value=_ENV_SITE):
        auth = enrich_auth_from_env({"username": "", "password": ""}, base_url="https://na-test.example.com")
    assert auth["username"] == "env-admin"
    assert auth["password"] == "from-env"
    assert auth["totp_secret"] == SAMPLE_TOTP
    assert auth["default_key"] == "sk-env"


def test_parse_sites_fills_password_and_totp_from_env():
    with patch("lib.data_summary.load_raw_env_sites", return_value=_ENV_SITE):
        sites = parse_sites(
            [{"base_url": "https://na-test.example.com", "username": "", "password": "", "env": "test"}]
        )
    assert len(sites) == 1
    assert sites[0].username == "env-admin"
    assert sites[0].password == "from-env"
    assert sites[0].totp_secret == SAMPLE_TOTP


if __name__ == "__main__":
    import tempfile

    test_token_to_row_has_all_headers()
    test_token_to_row_project_env_matches_group()
    test_channel_to_row_deploy_time_from_created_time()
    test_channel_to_row_deploy_time_empty_when_missing()
    test_channel_to_row_response_time_after_status()
    test_duplicate_channels_across_envs()
    test_duplicate_tokens_across_envs()
    test_exclude_duplicates_from_main_sheets()
    test_merge_rows_keeps_same_name_different_ids()
    test_sort_rows_by_id_asc()
    with tempfile.TemporaryDirectory() as tmp:
        test_collect_keeps_cross_env_duplicates_in_main_sheet(Path(tmp))
    with tempfile.TemporaryDirectory() as tmp:
        test_collect_creates_new_file_not_latest(Path(tmp))
    with tempfile.TemporaryDirectory() as tmp:
        test_collect_session_output_path_merges_envs(Path(tmp))
    with tempfile.TemporaryDirectory() as tmp:
        test_collect_all_fail_does_not_write(Path(tmp))
    test_enrich_site_from_env_fills_password_and_totp()
    test_enrich_site_from_env_keeps_explicit_password()
    test_enrich_auth_from_env_fills_password_and_totp()
    test_parse_sites_fills_password_and_totp_from_env()
    print("all tests passed")
