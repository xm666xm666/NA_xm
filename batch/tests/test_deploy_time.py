# -*- coding: utf-8 -*-
"""上游表按部署时间降序。"""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from lib.data_summary import CHANNEL_HEADERS  # noqa: F401
from lib.deploy_time import parse_deploy_datetime, sort_rows_by_deploy_time_desc
from lib.upstream_azure import AZURE_HEADERS
from lib.upstream_exporter import SHEET_AZURE, SHEET_VOLC, save_combined_upstream_workbook, save_upstream_workbook
from lib.upstream_volc import UPSTREAM_HEADERS


class DeployTimeSortTests(unittest.TestCase):
    def test_sort_newest_first_empty_last(self):
        rows = [
            {"部署名": "old", "部署时间": "2026-01-01 00:00:00"},
            {"部署名": "empty", "部署时间": ""},
            {"部署名": "new", "部署时间": "2026-08-25 10:00:00"},
            {"部署名": "bad", "部署时间": "not-a-date"},
        ]
        ordered = sort_rows_by_deploy_time_desc(rows)
        self.assertEqual([r["部署名"] for r in ordered], ["new", "old", "empty", "bad"])

    def test_parse_common_formats(self):
        self.assertEqual(
            parse_deploy_datetime("2026-08-25 10:00:00").strftime("%Y-%m-%d %H:%M:%S"),
            "2026-08-25 10:00:00",
        )
        self.assertIsNone(parse_deploy_datetime(""))
        self.assertIsNone(parse_deploy_datetime(None))

    def test_volc_workbook_newest_on_top(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "上游信息_火山云_test.xlsx"
            save_upstream_workbook(
                path,
                [
                    {"接入点名称": "old", "部署时间": "2025-01-01 00:00:00"},
                    {"接入点名称": "new", "部署时间": "2026-08-01 12:00:00"},
                    {"接入点名称": "empty", "部署时间": ""},
                ],
                headers=list(UPSTREAM_HEADERS),
            )
            from openpyxl import load_workbook

            wb = load_workbook(path, read_only=True)
            try:
                names = [row[0] for row in wb[SHEET_VOLC].iter_rows(min_row=2, values_only=True)]
            finally:
                wb.close()
            self.assertEqual(names, ["new", "old", "empty"])

    def test_azure_workbook_newest_on_top(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "上游信息_Azure_test.xlsx"
            save_upstream_workbook(
                path,
                [
                    {"部署名": "old", "部署时间": "2025-03-01 00:00:00"},
                    {"部署名": "new", "部署时间": "2026-08-20 08:00:00"},
                ],
                headers=list(AZURE_HEADERS),
                sheet_name=SHEET_AZURE,
            )
            from openpyxl import load_workbook

            wb = load_workbook(path, read_only=True)
            try:
                headers = [cell.value for cell in next(wb[SHEET_AZURE].iter_rows(min_row=1, max_row=1))]
                idx = headers.index("部署名")
                names = [row[idx] for row in wb[SHEET_AZURE].iter_rows(min_row=2, values_only=True)]
            finally:
                wb.close()
            self.assertEqual(names, ["new", "old"])

    def test_combined_workbook_both_sheets_sorted(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "上游信息_汇总_test.xlsx"
            save_combined_upstream_workbook(
                path,
                [
                    {"接入点名称": "v-old", "部署时间": "2025-01-01 00:00:00"},
                    {"接入点名称": "v-new", "部署时间": "2026-08-01 00:00:00"},
                ],
                [
                    {"部署名": "a-old", "部署时间": "2024-12-01 00:00:00"},
                    {"部署名": "a-new", "部署时间": "2026-07-01 00:00:00"},
                ],
            )
            from openpyxl import load_workbook

            wb = load_workbook(path, read_only=True)
            try:
                volc_names = [row[0] for row in wb[SHEET_VOLC].iter_rows(min_row=2, values_only=True)]
                az_headers = [cell.value for cell in next(wb[SHEET_AZURE].iter_rows(min_row=1, max_row=1))]
                az_idx = az_headers.index("部署名")
                azure_names = [row[az_idx] for row in wb[SHEET_AZURE].iter_rows(min_row=2, values_only=True)]
            finally:
                wb.close()
            self.assertEqual(volc_names, ["v-new", "v-old"])
            self.assertEqual(azure_names, ["a-new", "a-old"])

    def test_upstream_workbook_header_autofilter(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "上游信息_汇总_filter.xlsx"
            save_combined_upstream_workbook(
                path,
                [{"接入点名称": "v1", "部署时间": "2026-08-01 00:00:00"}],
                [{"部署名": "a1", "部署时间": "2026-07-01 00:00:00"}],
            )
            from openpyxl import load_workbook
            from openpyxl.utils import get_column_letter

            wb = load_workbook(path)
            try:
                for sheet in (wb[SHEET_VOLC], wb[SHEET_AZURE]):
                    last_col = get_column_letter(sheet.max_column)
                    self.assertEqual(sheet.auto_filter.ref, f"A1:{last_col}{sheet.max_row}")
            finally:
                wb.close()


if __name__ == "__main__":
    unittest.main()
