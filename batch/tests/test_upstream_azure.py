# -*- coding: utf-8 -*-
"""Azure OpenAI 上游采集格式单元测试（无网络）。"""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from lib.azure_arm import parse_resource_group
from lib.upstream_azure import (
    AZURE_HEADERS,
    build_endpoint_url,
    deployment_to_row,
    format_location,
    format_sku_label,
    is_openai_account,
    parse_tag_project_owner,
    resolve_rpm,
    resolve_tpm,
)
from lib.upstream_exporter import SHEET_AZURE, save_upstream_workbook, upstream_workbook_stats


class UpstreamAzureFormatTests(unittest.TestCase):
    def test_parse_resource_group(self):
        account_id = (
            "/subscriptions/76c8315b-d084-4f21-8003-cf6365e711bb"
            "/resourceGroups/phc_openai_order1"
            "/providers/Microsoft.CognitiveServices/accounts/phc-openai-clinai-dev1"
        )
        self.assertEqual(parse_resource_group(account_id), "phc_openai_order1")

    def test_format_location_and_sku(self):
        self.assertEqual(format_location("eastus2"), "East US2")
        self.assertEqual(format_sku_label("GlobalStandard"), "全局标准")
        self.assertEqual(format_sku_label("Standard"), "标准")

    def test_tpm_from_capacity(self):
        self.assertEqual(resolve_tpm({"sku": {"capacity": 10000}}), 10000000)
        self.assertEqual(resolve_tpm({"sku": {"capacity": 0}}), "")

    def test_rpm_from_rate_limits(self):
        dep = {
            "properties": {
                "rateLimits": [
                    {"key": "token", "count": 10000000},
                    {"key": "request", "count": 10000},
                ]
            }
        }
        self.assertEqual(resolve_rpm(dep), 10000)

    def test_endpoint_url(self):
        self.assertIn(
            "/openai/responses?api-version=2025-04-01-preview",
            build_endpoint_url(
                "https://phc-openai-clinai-dev1.openai.azure.com/",
                "phc-openai-clinai-dev1",
                "clinai-dev-gpt-5.1",
                "gpt-5.1",
            ),
        )
        self.assertIn(
            "/embeddings?api-version=2023-05-15",
            build_endpoint_url(
                "https://phc-openai-clinai-dev.openai.azure.com/",
                "phc-openai-clinai-dev",
                "clinai-dev-text-embedding-ada-002",
                "text-embedding-ada-002",
            ),
        )

    def test_is_openai_account(self):
        self.assertTrue(is_openai_account({"kind": "OpenAI", "name": "x"}))
        self.assertTrue(is_openai_account({"kind": "AIServices", "name": "hub"}))
        self.assertFalse(is_openai_account({"kind": "TextAnalytics", "name": "lang"}))

    def test_deployment_to_row_headers(self):
        row = deployment_to_row(
            subscription_id="76c8315b-d084-4f21-8003-cf6365e711bb",
            subscription_label="订阅1",
            location="eastus2",
            account_name="phc-openai-clinai-dev1",
            account_endpoint="https://phc-openai-clinai-dev1.openai.azure.com/",
            deployment={
                "name": "clinai-dev-gpt-5.1",
                "sku": {"name": "GlobalStandard", "capacity": 10000},
                "properties": {
                    "model": {"name": "gpt-5.1", "version": "2025-11-13"},
                    "rateLimits": [{"key": "request", "count": 100000}],
                },
            },
            api_key="key-demo",
        )
        self.assertEqual(list(row.keys()), AZURE_HEADERS)
        self.assertEqual(row["区域"], "East US2")
        self.assertEqual(row["Azure OpenAI"], "phc-openai-clinai-dev1")
        self.assertEqual(row["部署名"], "clinai-dev-gpt-5.1")
        self.assertEqual(row["模型名"], "gpt-5.1")
        self.assertEqual(row["TPM"], 10000000)
        self.assertEqual(row["RPM"], 100000)
        self.assertEqual(row["密钥一"], "key-demo")
        self.assertEqual(row["订阅"], "订阅1")
        self.assertEqual(row["项目"], "N/A")
        self.assertEqual(row["负责人"], "N/A")
        self.assertEqual(row["部署时间"], "")
        self.assertEqual(AZURE_HEADERS[-1], "部署时间")
        self.assertEqual(AZURE_HEADERS[-3:-1], ["项目", "负责人"])
        self.assertEqual(AZURE_HEADERS[-4], "密钥一")

    def test_deployment_to_row_uses_system_data_created_at(self):
        row = deployment_to_row(
            subscription_id="sub",
            location="eastus2",
            account_name="acc",
            account_endpoint="https://acc.openai.azure.com/",
            deployment={
                "name": "d1",
                "sku": {"name": "Standard", "capacity": 1},
                "properties": {},
                "systemData": {"createdAt": "2025-11-13T06:30:00Z"},
            },
        )
        self.assertEqual(row["部署时间"], "2025-11-13 14:30:00")

    def test_parse_tag_project_owner(self):
        self.assertEqual(parse_tag_project_owner(None), ("N/A", "N/A"))
        self.assertEqual(parse_tag_project_owner({}), ("N/A", "N/A"))
        self.assertEqual(parse_tag_project_owner({"env": "prod"}), ("N/A", "N/A"))
        self.assertEqual(parse_tag_project_owner({"1": ""}), ("N/A", "N/A"))
        self.assertEqual(
            parse_tag_project_owner({"1": "产研AI团队-乔加奇"}),
            ("产研AI团队", "乔加奇"),
        )
        self.assertEqual(
            parse_tag_project_owner({"1": "产研AI团队-乔加奇，李四"}),
            ("产研AI团队", "乔加奇，李四"),
        )
        self.assertEqual(
            parse_tag_project_owner({"1": "产研AI团队-乔加奇,李四"}),
            ("产研AI团队", "乔加奇，李四"),
        )
        self.assertEqual(
            parse_tag_project_owner({"1": "产研AI团队-乔加奇、李四"}),
            ("产研AI团队", "乔加奇，李四"),
        )
        self.assertEqual(
            parse_tag_project_owner({"1": "产研AI团队-乔加奇, 李四、王五，赵六"}),
            ("产研AI团队", "乔加奇，李四，王五，赵六"),
        )
        self.assertEqual(
            parse_tag_project_owner([{"name": "1", "value": "临床-张三，王五"}]),
            ("临床", "张三，王五"),
        )
        self.assertEqual(parse_tag_project_owner({"1": "仅项目名"}), ("仅项目名", "N/A"))
        self.assertEqual(
            parse_tag_project_owner({"1": "AI-Product-金鑫"}),
            ("AI-Product", "金鑫"),
        )
        self.assertEqual(
            parse_tag_project_owner({"1": "AI-Product-金鑫，李四"}),
            ("AI-Product", "金鑫，李四"),
        )

    def test_deployment_to_row_from_tags(self):
        row = deployment_to_row(
            subscription_id="sub",
            location="eastus2",
            account_name="phc-openai-aiic-llmg",
            account_endpoint="https://phc-openai-aiic-llmg.openai.azure.com/",
            deployment={"name": "d1", "sku": {"name": "Standard", "capacity": 1}, "properties": {}},
            tags={"1": "产研AI团队-乔加奇"},
        )
        self.assertEqual(row["项目"], "产研AI团队")
        self.assertEqual(row["负责人"], "乔加奇")

    def test_save_azure_workbook(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "上游信息_Azure_test.xlsx"
            rows = [
                deployment_to_row(
                    subscription_id="sub",
                    location="eastus",
                    account_name="acc",
                    account_endpoint="https://acc.openai.azure.com/",
                    deployment={
                        "name": "d1",
                        "sku": {"name": "Standard", "capacity": 1},
                        "properties": {"model": {"name": "gpt-4o", "version": "2024-08-06"}},
                    },
                )
            ]
            save_upstream_workbook(path, rows, headers=AZURE_HEADERS, sheet_name=SHEET_AZURE)
            stats = upstream_workbook_stats(path)
            self.assertEqual(stats["sheet"], SHEET_AZURE)
            self.assertEqual(stats["rows"], 1)
            from openpyxl import load_workbook

            wb = load_workbook(path, read_only=True)
            try:
                headers = [cell.value for cell in next(wb[SHEET_AZURE].iter_rows(min_row=1, max_row=1))]
                self.assertEqual(headers, AZURE_HEADERS)
                values = [cell.value for cell in next(wb[SHEET_AZURE].iter_rows(min_row=2, max_row=2))]
                self.assertEqual(values[headers.index("项目")], "N/A")
                self.assertEqual(values[headers.index("负责人")], "N/A")
            finally:
                wb.close()

    def test_acquire_token_error_message(self):
        from lib.azure_arm import AzureARMError, acquire_arm_token

        with self.assertRaises(AzureARMError):
            acquire_arm_token("", "id", "secret")
        with patch("lib.azure_arm.requests.post") as mocked:
            mocked.return_value.status_code = 401
            mocked.return_value.json.return_value = {"error_description": "invalid client"}
            mocked.return_value.text = "invalid client"
            with self.assertRaises(AzureARMError) as ctx:
                acquire_arm_token("tenant", "id", "secret")
            self.assertIn("401", str(ctx.exception))

    def test_parse_subscription_entries_labels(self):
        from lib.azure_arm import parse_subscription_entries

        entries = parse_subscription_entries(
            [
                "76c8315b-d084-4f21-8003-cf6365e711bb",
                "22222222-2222-2222-2222-222222222222",
            ],
            default="",
        )
        self.assertEqual(entries[0]["label"], "订阅1")
        self.assertEqual(entries[1]["label"], "订阅2")
        named = parse_subscription_entries(
            [{"id": "76c8315b-d084-4f21-8003-cf6365e711bb", "name": "订阅1"}],
            default="",
        )
        self.assertEqual(named[0]["label"], "订阅1")
        from lib.azure_arm import normalize_subscription_ids

        ids = normalize_subscription_ids(
            ["aaa", "bbb"],
            "aaa\nccc",
            default="",
        )
        self.assertEqual(ids, ["aaa", "bbb", "ccc"])
        self.assertEqual(normalize_subscription_ids(default="76c8315b-d084-4f21-8003-cf6365e711bb"), [
            "76c8315b-d084-4f21-8003-cf6365e711bb"
        ])

    def test_collect_multiple_subscriptions(self):
        from lib.upstream_azure import collect_azure_upstream

        def make_account(sub, name):
            return {
                "name": name,
                "kind": "OpenAI",
                "location": "eastus2",
                "id": f"/subscriptions/{sub}/resourceGroups/rg/providers/Microsoft.CognitiveServices/accounts/{name}",
                "properties": {"endpoint": f"https://{name}.openai.azure.com/"},
            }

        dep = {
            "name": "d1",
            "sku": {"name": "Standard", "capacity": 1},
            "properties": {"model": {"name": "gpt-4o", "version": "1"}},
        }
        sub_a = "11111111-1111-1111-1111-111111111111"
        sub_b = "22222222-2222-2222-2222-222222222222"
        client_a = Mock()
        client_a.subscription_id = sub_a
        client_a.list_accounts.return_value = [make_account(sub_a, "acc-a")]
        client_a.list_deployments.return_value = [dep]
        client_a.list_keys.return_value = {"key1": "ka"}
        client_b = Mock()
        client_b.subscription_id = sub_b
        client_b.list_accounts.return_value = [make_account(sub_b, "acc-b")]
        client_b.list_deployments.return_value = [dep]
        client_b.list_keys.return_value = {"key1": "kb"}
        with patch("lib.upstream_azure.acquire_arm_token", return_value="tok"):
            with patch("lib.upstream_azure.AzureARMClient", side_effect=[client_a, client_b]):
                stats = collect_azure_upstream(
                    "tenant",
                    "client",
                    "secret",
                    subscription_ids=[sub_a, sub_b],
                )
        self.assertEqual(stats.subscriptions_ok, 2)
        self.assertEqual(stats.deployments, 2)
        self.assertEqual({row["订阅"] for row in stats.rows}, {"订阅1", "订阅2"})
        self.assertEqual({row["Azure OpenAI"] for row in stats.rows}, {"acc-a", "acc-b"})
        from lib.azure_arm import AzureARMError, acquire_arm_token

        with self.assertRaises(AzureARMError):
            acquire_arm_token("", "id", "secret")
        with patch("lib.azure_arm.requests.post") as mocked:
            mocked.return_value.status_code = 401
            mocked.return_value.json.return_value = {"error_description": "invalid client"}
            mocked.return_value.text = "invalid client"
            with self.assertRaises(AzureARMError) as ctx:
                acquire_arm_token("tenant", "id", "secret")
            self.assertIn("401", str(ctx.exception))

    def test_collect_reads_tag_project_owner(self):
        from lib.upstream_azure import collect_azure_upstream

        account = {
            "name": "phc-openai-clinai-dev1",
            "kind": "OpenAI",
            "location": "eastus2",
            "id": (
                "/subscriptions/76c8315b-d084-4f21-8003-cf6365e711bb"
                "/resourceGroups/phc_openai_order1"
                "/providers/Microsoft.CognitiveServices/accounts/phc-openai-clinai-dev1"
            ),
            "tags": {"1": "产研AI团队-乔加奇"},
            "properties": {"endpoint": "https://phc-openai-clinai-dev1.openai.azure.com/"},
        }
        deployment = {
            "name": "clinai-dev-gpt-5.1",
            "sku": {"name": "GlobalStandard", "capacity": 10},
            "properties": {
                "model": {"name": "gpt-5.1", "version": "2025-11-13"},
                "rateLimits": [{"key": "request", "count": 60}],
            },
        }
        with patch("lib.upstream_azure.acquire_arm_token", return_value="tok"):
            with patch("lib.upstream_azure.AzureARMClient") as mocked:
                inst = mocked.return_value
                inst.subscription_id = "76c8315b-d084-4f21-8003-cf6365e711bb"
                inst.list_accounts.return_value = [account]
                inst.list_deployments.return_value = [deployment]
                inst.list_keys.return_value = {"key1": "abc", "key2": "def"}
                stats = collect_azure_upstream("tenant", "client", "secret")
        self.assertEqual(stats.deployments, 1)
        self.assertEqual(list(stats.rows[0].keys()), AZURE_HEADERS)
        self.assertEqual(stats.rows[0]["项目"], "产研AI团队")
        self.assertEqual(stats.rows[0]["负责人"], "乔加奇")
        self.assertEqual(stats.rows[0]["密钥一"], "abc")
        self.assertEqual(list(stats.rows[0].keys())[-3:-1], ["项目", "负责人"])
        self.assertEqual(list(stats.rows[0].keys())[-1], "部署时间")
        self.assertEqual(stats.rows[0]["Azure OpenAI"], "phc-openai-clinai-dev1")
        self.assertEqual(stats.rows[0]["订阅"], "订阅1")

    def test_collect_get_account_when_list_omits_tags(self):
        from lib.upstream_azure import collect_azure_upstream

        account = {
            "name": "phc-openai-aiic-llmg",
            "kind": "OpenAI",
            "location": "eastus2",
            "id": (
                "/subscriptions/76c8315b-d084-4f21-8003-cf6365e711bb"
                "/resourceGroups/phc_openai_order1"
                "/providers/Microsoft.CognitiveServices/accounts/phc-openai-aiic-llmg"
            ),
            "properties": {"endpoint": "https://phc-openai-aiic-llmg.openai.azure.com/"},
        }
        deployment = {
            "name": "d1",
            "sku": {"name": "Standard", "capacity": 1},
            "properties": {"model": {"name": "gpt-4o", "version": "1"}},
        }
        with patch("lib.upstream_azure.acquire_arm_token", return_value="tok"):
            with patch("lib.upstream_azure.AzureARMClient") as mocked:
                inst = mocked.return_value
                inst.subscription_id = "76c8315b-d084-4f21-8003-cf6365e711bb"
                inst.list_accounts.return_value = [account]
                inst.list_deployments.return_value = [deployment]
                inst.list_keys.return_value = {"key1": "k1"}
                inst.get_account.return_value = {**account, "tags": {"1": "产研AI团队-乔加奇，李四"}}
                stats = collect_azure_upstream("tenant", "client", "secret")
                inst.get_account.assert_called_once()
        self.assertEqual(stats.rows[0]["项目"], "产研AI团队")
        self.assertEqual(stats.rows[0]["负责人"], "乔加奇，李四")
        self.assertEqual(list(stats.rows[0].keys())[-3:-1], ["项目", "负责人"])
        self.assertEqual(list(stats.rows[0].keys())[-1], "部署时间")

    def test_collect_missing_tag_is_na(self):
        from lib.upstream_azure import collect_azure_upstream

        account = {
            "name": "phc-openai-clinai-dev1",
            "kind": "OpenAI",
            "location": "eastus2",
            "id": (
                "/subscriptions/76c8315b-d084-4f21-8003-cf6365e711bb"
                "/resourceGroups/phc_openai_order1"
                "/providers/Microsoft.CognitiveServices/accounts/phc-openai-clinai-dev1"
            ),
            "properties": {"endpoint": "https://phc-openai-clinai-dev1.openai.azure.com/"},
        }
        deployment = {
            "name": "clinai-dev-gpt-5.1",
            "sku": {"name": "GlobalStandard", "capacity": 10},
            "properties": {
                "model": {"name": "gpt-5.1", "version": "2025-11-13"},
                "rateLimits": [{"key": "request", "count": 60}],
            },
        }
        with patch("lib.upstream_azure.acquire_arm_token", return_value="tok"):
            with patch("lib.upstream_azure.AzureARMClient") as mocked:
                inst = mocked.return_value
                inst.subscription_id = "76c8315b-d084-4f21-8003-cf6365e711bb"
                inst.list_accounts.return_value = [account]
                inst.list_deployments.return_value = [deployment]
                inst.list_keys.return_value = {"key1": "abc", "key2": "def"}
                stats = collect_azure_upstream("tenant", "client", "secret")
        self.assertEqual(list(stats.rows[0].keys()), AZURE_HEADERS)
        self.assertEqual(stats.rows[0]["项目"], "N/A")
        self.assertEqual(stats.rows[0]["负责人"], "N/A")
        self.assertNotIn("备注", stats.rows[0])
        self.assertEqual(stats.rows[0]["密钥一"], "abc")
        self.assertEqual(stats.rows[0]["Azure OpenAI"], "phc-openai-clinai-dev1")
        self.assertEqual(stats.rows[0]["订阅"], "订阅1")


class CombinedUpstreamWorkbookTests(unittest.TestCase):
    def test_save_combined_has_both_sheets(self):
        from lib.upstream_exporter import (
            SHEET_AZURE,
            SHEET_VOLC,
            list_combined_upstream_excel_files,
            save_combined_upstream_workbook,
            upstream_workbook_stats,
        )

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            path = root / "excel" / "上游信息_汇总_test.xlsx"
            save_combined_upstream_workbook(
                path,
                [{"接入点名称": "ep-a", "部署名": "ep-1"}],
                [{"订阅": "订阅1", "部署名": "d1", "Azure OpenAI": "acc"}],
            )
            stats = upstream_workbook_stats(path)
            self.assertEqual(stats["sheets"][SHEET_VOLC], 1)
            self.assertEqual(stats["sheets"][SHEET_AZURE], 1)
            self.assertEqual(stats["total_rows"], 2)
            self.assertIn(SHEET_VOLC, stats["sheet"])
            self.assertIn(SHEET_AZURE, stats["sheet"])
            from openpyxl import load_workbook

            wb = load_workbook(path, read_only=True)
            try:
                headers = [cell.value for cell in next(wb[SHEET_AZURE].iter_rows(min_row=1, max_row=1))]
                self.assertEqual(headers, AZURE_HEADERS)
                values = [cell.value for cell in next(wb[SHEET_AZURE].iter_rows(min_row=2, max_row=2))]
                self.assertEqual(values[headers.index("项目")], "N/A")
                self.assertEqual(values[headers.index("负责人")], "N/A")
            finally:
                wb.close()

            (root / "excel" / "上游信息_火山云_20260101000000.xlsx").write_bytes(b"x")
            (root / "excel" / "上游信息_Azure_20260101000000.xlsx").write_bytes(b"x")
            files = list_combined_upstream_excel_files(root)
            self.assertEqual([item.name for item in files], ["上游信息_汇总_test.xlsx"])

    def test_combined_keeps_empty_sheet(self):
        from lib.upstream_exporter import SHEET_AZURE, SHEET_VOLC, save_combined_upstream_workbook
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "combined.xlsx"
            save_combined_upstream_workbook(path, [{"接入点名称": "a"}], [])
            wb = load_workbook(path, read_only=True)
            try:
                self.assertEqual(wb.sheetnames, [SHEET_VOLC, SHEET_AZURE])
                self.assertEqual(wb[SHEET_AZURE].max_row, 1)
            finally:
                wb.close()


if __name__ == "__main__":
    unittest.main()
