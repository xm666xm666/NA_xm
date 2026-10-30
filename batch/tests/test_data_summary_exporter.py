"""Excel 导出样式：边框与同分组着色。"""

from __future__ import annotations

import sys
from pathlib import Path

BATCH_DIR = Path(__file__).resolve().parents[1]
if str(BATCH_DIR) not in sys.path:
    sys.path.insert(0, str(BATCH_DIR))

from openpyxl import load_workbook  # noqa: E402

from lib.data_summary import (  # noqa: E402
    ACTIVITY_HEADER,
    CHANNEL_HEADERS,
    DUP_CHANNEL_HEADERS,
    ENVS,
    EXCEL_STEM,
    SHEET_CHANNELS,
    SHEET_DUP_CHANNELS,
    SHEET_TOKENS,
    TOKEN_HEADERS,
)
from lib.data_summary_exporter import (  # noqa: E402
    build_timestamped_filename,
    load_workbook_data,
    is_staging_excel,
    latest_excel_path,
    list_summary_excel_files,
    publish_staging_excel,
    save_workbook_data,
)
from datetime import datetime  # noqa: E402


def test_build_timestamped_filename():
    name = build_timestamped_filename(datetime(2026, 7, 16, 16, 5, 30))
    assert name == f"{EXCEL_STEM}_20260716160530.xlsx"


def test_staging_excel_not_listed_as_latest(tmp_path: Path):
    excel = tmp_path / "excel"
    excel.mkdir()
    part = excel / f"{EXCEL_STEM}_20260820180000.part.xlsx"
    final = excel / f"{EXCEL_STEM}_20260820170000.xlsx"
    final.write_bytes(b"complete")
    part.write_bytes(b"partial")
    assert is_staging_excel(part)
    assert not is_staging_excel(final)
    listed = list_summary_excel_files(tmp_path)
    assert listed == [final]
    assert latest_excel_path(tmp_path) == final
    published = publish_staging_excel(part)
    assert published.name == f"{EXCEL_STEM}_20260820180000.xlsx"
    assert published.exists()
    assert not part.exists()
    assert latest_excel_path(tmp_path) == published


def test_overview_excel_not_listed_as_gateway_latest(tmp_path: Path):
    excel = tmp_path / "excel"
    excel.mkdir()
    gateway = excel / f"{EXCEL_STEM}_20260820170000.xlsx"
    overview = excel / f"{EXCEL_STEM}_一表总览_20260820180000.xlsx"
    gateway.write_bytes(b"gateway")
    overview.write_bytes(b"overview")
    listed = list_summary_excel_files(tmp_path)
    assert listed == [gateway]
    assert latest_excel_path(tmp_path) == gateway


def test_gateway_latest_uses_filename_timestamp(tmp_path: Path):
    import os

    excel = tmp_path / "excel"
    excel.mkdir()
    newer_stamp = excel / f"{EXCEL_STEM}_20260820190000.xlsx"
    older_stamp = excel / f"{EXCEL_STEM}_20260820170000.xlsx"
    newer_stamp.write_bytes(b"new")
    older_stamp.write_bytes(b"old")
    os.utime(older_stamp, None)
    assert latest_excel_path(tmp_path) == newer_stamp


def test_export_applies_borders_and_group_colors(tmp_path: Path):
    excel_path = tmp_path / "summary.xlsx"
    rows = {
        "test": [
            {
                "名称": "ch-a",
                "分组": "group-a",
                "类型": "OpenAI",
                "状态": "已启用",
                "模型列表": "m1",
                "模型重定向": "",
                "base_url": "https://a",
                "部署时间": "",
            },
            {
                "名称": "ch-b",
                "分组": "group-a",
                "类型": "OpenAI",
                "状态": "已启用",
                "模型列表": "m2",
                "模型重定向": "",
                "base_url": "https://b",
                "部署时间": "",
            },
            {
                "名称": "ch-c",
                "分组": "group-b",
                "类型": "OpenAI",
                "状态": "已启用",
                "模型列表": "m3",
                "模型重定向": "",
                "base_url": "https://c",
                "部署时间": "",
            },
        ]
    }
    save_workbook_data(
        excel_path,
        channel_env_rows={env: rows.get("test", []) if env == "test" else [] for env in ENVS},
        token_env_rows={env: [] for env in ENVS},
    )

    wb = load_workbook(excel_path)
    ws = wb[SHEET_CHANNELS["test"]]
    assert ws["A1"].font.bold is True
    assert ws["A1"].font.name == "宋体"
    assert ws["A1"].font.size == 11
    assert ws["A2"].font.name == "宋体"
    assert ws["A2"].font.size == 11
    assert ws["A1"].border.left.style == "thin"
    assert ws["A1"].border.left.color.rgb == "FF000000"
    assert ws["A2"].border.left.style == "thin"
    assert ws["A2"].border.right.style == "thin"
    assert ws["A1"].border.bottom.style == "thin"
    assert ws.cell(row=4, column=len(CHANNEL_HEADERS)).border.right.style == "thin"
    assert ws.cell(row=4, column=1).border.bottom.style == "thin"

    group_col = CHANNEL_HEADERS.index("分组") + 1
    fill_a = ws.cell(row=2, column=group_col).fill.fgColor.rgb
    fill_b = ws.cell(row=3, column=group_col).fill.fgColor.rgb
    fill_c = ws.cell(row=4, column=group_col).fill.fgColor.rgb
    assert fill_a == fill_b
    assert fill_a != fill_c
    wb.close()


def test_disabled_row_uses_red_font_not_red_fill(tmp_path: Path):
    excel_path = tmp_path / "disabled.xlsx"
    rows = {
        "test": [
            {
                "名称": "ch-on",
                "分组": "group-a",
                "类型": "OpenAI",
                "状态": "已启用",
                "模型列表": "m1",
                "模型重定向": "",
                "base_url": "https://a",
                "部署时间": "",
            },
            {
                "名称": "ch-off",
                "分组": "group-a",
                "类型": "OpenAI",
                "状态": "已禁用",
                "模型列表": "m2",
                "模型重定向": "",
                "base_url": "https://b",
                "部署时间": "",
            },
        ]
    }
    save_workbook_data(
        excel_path,
        channel_env_rows={env: rows.get("test", []) if env == "test" else [] for env in ENVS},
        token_env_rows={env: [] for env in ENVS},
    )
    wb = load_workbook(excel_path)
    ws = wb[SHEET_CHANNELS["test"]]
    assert ws["A2"].font.color is None or (
        getattr(ws["A2"].font.color, "rgb", None) in (None, "FF000000", "00000000")
    )
    assert ws["A3"].font.color.rgb == "FFFF0000"
    assert ws["D3"].font.color.rgb == "FFFF0000"
    fill_rgb = str(ws["A3"].fill.fgColor.rgb or "")
    assert "FF0000" not in fill_rgb.upper()
    assert fill_rgb.upper() not in {"FFFF0000", "FFCC4125", "FFEA9999", "FFF4CCCC"}
    wb.close()


def test_channel_activity_column_is_after_batch_config(tmp_path: Path):
    excel_path = tmp_path / "activity.xlsx"
    save_workbook_data(
        excel_path,
        channel_env_rows={
            "test": [{"名称": "ch-a", "分组": "g", "activity_label": "m: 近一月无活跃"}],
            "uat": [],
            "prod": [],
        },
        token_env_rows={env: [] for env in ENVS},
    )
    wb = load_workbook(excel_path)
    ws = wb[SHEET_CHANNELS["test"]]
    headers = [cell.value for cell in ws[1]]
    assert headers[-2] == "批量配置"
    assert headers[-1] == ACTIVITY_HEADER
    assert ws.cell(2, headers.index(ACTIVITY_HEADER) + 1).value == "m: 近一月无活跃"
    wb.close()


def test_channel_remark_mismatch_only_uses_distinct_remark_font_color(tmp_path: Path):
    excel_path = tmp_path / "remark-mismatch.xlsx"
    rows = {
        "test": [
            {
                "名称": "matched",
                "分组": "group-a",
                "类型": "OpenAI",
                "状态": "已启用",
                "模型列表": "m1",
                "模型重定向": "",
                "base_url": "https://matched.example.com",
                "备注": "管理员说明\nhttps://matched.example.com",
                "部署时间": "",
            },
            {
                "名称": "missing-url",
                "分组": "group-a",
                "类型": "OpenAI",
                "状态": "已启用",
                "模型列表": "m2",
                "模型重定向": "",
                "base_url": "https://missing.example.com",
                "备注": "管理员说明",
                "部署时间": "",
            },
            {
                "名称": "orphan-remark",
                "分组": "group-a",
                "类型": "OpenAI",
                "状态": "已启用",
                "模型列表": "m3",
                "模型重定向": "",
                "base_url": "",
                "备注": "没有对应地址",
                "部署时间": "",
            },
        ]
    }
    save_workbook_data(
        excel_path,
        channel_env_rows={env: rows.get("test", []) if env == "test" else [] for env in ENVS},
        token_env_rows={env: [] for env in ENVS},
    )
    wb = load_workbook(excel_path)
    ws = wb[SHEET_CHANNELS["test"]]
    base_col = CHANNEL_HEADERS.index("base_url") + 1
    remark_col = CHANNEL_HEADERS.index("备注") + 1
    assert ws.cell(row=2, column=base_col).fill.fgColor.rgb == ws.cell(row=3, column=base_col).fill.fgColor.rgb
    for row_idx in (3, 4):
        assert ws.cell(row=row_idx, column=remark_col).fill.fgColor.rgb == ws.cell(row=row_idx, column=base_col).fill.fgColor.rgb
        assert ws.cell(row=row_idx, column=remark_col).font.color.rgb == "FF7030A0"
        assert getattr(ws.cell(row=row_idx, column=base_col).font.color, "rgb", None) != "FF7030A0"
    wb.close()


def test_token_sheet_has_full_headers(tmp_path: Path):
    excel_path = tmp_path / "tokens.xlsx"
    token_row = {
        "名称": "token-a",
        "项目环境": "default",
        "状态": "已启用",
        "剩余额度": "无限额度",
        "分组": "default",
        "可用模型": "无限制",
        "包含渠道": "无限制",
        "IP限制": "无限制",
        "创建时间": "",
        "过期时间": "永不过期",
    }
    save_workbook_data(
        excel_path,
        channel_env_rows={env: [] for env in ENVS},
        token_env_rows={"test": [token_row], "uat": [], "prod": []},
    )
    wb = load_workbook(excel_path)
    ws = wb[SHEET_TOKENS["test"]]
    headers = [ws.cell(row=1, column=i).value for i in range(1, len(TOKEN_HEADERS) + 1)]
    assert headers == TOKEN_HEADERS
    assert ws.cell(row=2, column=2).value == "default"
    assert ws.cell(row=2, column=5).value == "default"
    assert ws.cell(row=2, column=2).border.left.style == "thin"
    wb.close()


def test_token_sheet_colors_by_group(tmp_path: Path):
    excel_path = tmp_path / "tokens-color.xlsx"
    rows = [
        {
            "名称": "token-a",
            "项目环境": "group-a",
            "状态": "已启用",
            "剩余额度": "无限额度",
            "分组": "group-a",
            "可用模型": "无限制",
            "包含渠道": "无限制",
            "IP限制": "无限制",
            "创建时间": "",
            "过期时间": "永不过期",
        },
        {
            "名称": "token-b",
            "项目环境": "group-a",
            "状态": "已启用",
            "剩余额度": "无限额度",
            "分组": "group-a",
            "可用模型": "无限制",
            "包含渠道": "无限制",
            "IP限制": "无限制",
            "创建时间": "",
            "过期时间": "永不过期",
        },
        {
            "名称": "token-c",
            "项目环境": "group-b",
            "状态": "已启用",
            "剩余额度": "无限额度",
            "分组": "group-b",
            "可用模型": "无限制",
            "包含渠道": "无限制",
            "IP限制": "无限制",
            "创建时间": "",
            "过期时间": "永不过期",
        },
    ]
    save_workbook_data(
        excel_path,
        channel_env_rows={env: [] for env in ENVS},
        token_env_rows={"test": rows, "uat": [], "prod": []},
    )
    wb = load_workbook(excel_path)
    ws = wb[SHEET_TOKENS["test"]]
    group_col = TOKEN_HEADERS.index("分组") + 1
    fill_a = ws.cell(row=2, column=group_col).fill.fgColor.rgb
    fill_b = ws.cell(row=3, column=group_col).fill.fgColor.rgb
    fill_c = ws.cell(row=4, column=group_col).fill.fgColor.rgb
    assert fill_a == fill_b
    assert fill_a != fill_c
    wb.close()


def test_many_groups_avoid_sequential_color_collision(tmp_path: Path):
    """第 2 与第 12 个分组不应因调色板取模而必然同色（旧逻辑缺陷）。"""
    excel_path = tmp_path / "many-groups.xlsx"
    group_names = [
        "ai-product-test",
        "demo-alpha-test",
        "g02",
        "g03",
        "g04",
        "g05",
        "g06",
        "g07",
        "g08",
        "g09",
        "g10",
        "demo-beta-test",
    ]
    rows = [
        {
            "名称": f"token-{i}",
            "项目环境": group,
            "状态": "已启用",
            "剩余额度": "无限额度",
            "分组": group,
            "可用模型": "无限制",
            "包含渠道": "无限制",
            "IP限制": "无限制",
            "创建时间": "",
            "过期时间": "永不过期",
        }
        for i, group in enumerate(group_names)
    ]
    save_workbook_data(
        excel_path,
        channel_env_rows={env: [] for env in ENVS},
        token_env_rows={"test": rows, "uat": [], "prod": []},
    )
    wb = load_workbook(excel_path)
    ws = wb[SHEET_TOKENS["test"]]
    group_col = TOKEN_HEADERS.index("分组") + 1
    by_group = {}
    for row_idx in range(2, ws.max_row + 1):
        name = ws.cell(row=row_idx, column=group_col).value
        fill = ws.cell(row=row_idx, column=group_col).fill.fgColor.rgb
        by_group[name] = fill
    assert by_group["demo-alpha-test"] == by_group["demo-alpha-test"]
    assert by_group["demo-alpha-test"] != by_group["demo-beta-test"]
    wb.close()


def test_dup_channel_sheet_colors_by_model_list(tmp_path: Path):
    from openpyxl import Workbook

    from lib.data_summary_exporter import _ensure_sheet, _write_sheet  # noqa: E402

    wb = Workbook()
    wb.remove(wb.active)
    dup_rows = [
        {
            "名称": "same-model",
            "类型": "OpenAI",
            "状态": "已启用",
            "模型列表": "gpt-4o",
            "模型重定向": "",
            "base_url": "https://a",
            "部署时间": "",
            "来源": "test-渠道管理",
            "重复序号": 1,
        },
        {
            "名称": "same-model",
            "类型": "OpenAI",
            "状态": "已启用",
            "模型列表": "gpt-4o",
            "模型重定向": "",
            "base_url": "https://b",
            "部署时间": "",
            "来源": "uat-渠道管理",
            "重复序号": 1,
        },
        {
            "名称": "other-model",
            "类型": "OpenAI",
            "状态": "已启用",
            "模型列表": "glm-5",
            "模型重定向": "",
            "base_url": "https://c",
            "部署时间": "",
            "来源": "test-渠道管理",
            "重复序号": 2,
        },
    ]
    _write_sheet(
        _ensure_sheet(wb, SHEET_DUP_CHANNELS, DUP_CHANNEL_HEADERS),
        DUP_CHANNEL_HEADERS,
        dup_rows,
        color_by="模型列表",
    )
    ws = wb[SHEET_DUP_CHANNELS]
    model_col = DUP_CHANNEL_HEADERS.index("模型列表") + 1
    fill_same = ws.cell(row=2, column=model_col).fill.fgColor.rgb
    fill_same2 = ws.cell(row=3, column=model_col).fill.fgColor.rgb
    fill_other = ws.cell(row=4, column=model_col).fill.fgColor.rgb
    assert fill_same == fill_same2
    assert fill_same != fill_other
    wb.close()


def test_custom_env_sheets_are_saved_and_loaded(tmp_path: Path):
    excel_path = tmp_path / "custom-env.xlsx"
    save_workbook_data(
        excel_path,
        channel_env_rows={"k8s-test": [{"名称": "channel-a"}]},
        token_env_rows={"k8s-test": [{"名称": "token-a"}]},
    )
    assert "k8s-test-渠道管理" in load_workbook(excel_path).sheetnames
    assert "k8s-test-令牌管理" in load_workbook(excel_path).sheetnames
    channel_rows, token_rows = load_workbook_data(excel_path)
    assert channel_rows["k8s-test"][0]["名称"] == "channel-a"
    assert token_rows["k8s-test"][0]["名称"] == "token-a"
