# -*- coding: utf-8 -*-
"""火山云账单汇总（无网络）。"""

from __future__ import annotations

import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import MagicMock

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from lib.volc_billing import (
    SHEET_DAY,
    SHEET_MONTH,
    SHEET_NOTE,
    aggregate_rows,
    bill_month_label,
    bill_periods_from_payload,
    build_bill_filename,
    collect_volc_bills,
    default_bill_period,
    expand_bill_period_range,
    is_blank_or_default_project,
    map_bill_item,
    parse_bill_period,
    save_bill_workbook,
)
from lib.volc_openapi import VolcOpenAPIError


class VolcBillingTests(unittest.TestCase):
    def test_parse_bill_period(self):
        self.assertEqual(parse_bill_period("2026-08"), "2026-08")
        self.assertEqual(parse_bill_period("2026/8"), "2026-08")
        with self.assertRaises(ValueError):
            parse_bill_period("2026-13")

    def test_default_bill_period_is_previous_month(self):
        from datetime import datetime, timezone, timedelta

        now = datetime(2026, 8, 31, tzinfo=timezone(timedelta(hours=8)))
        self.assertEqual(default_bill_period(now=now), "2026-07")
        jan = datetime(2026, 1, 2, tzinfo=timezone(timedelta(hours=8)))
        self.assertEqual(default_bill_period(now=jan), "2025-12")

    def test_map_and_aggregate(self):
        items = [
            {
                "ExpenseDate": "2026-08-01",
                "ProjectDisplayName": "demo-alpha",
                "ProductZh": "方舟",
                "ConfigName": "input",
                "SplitItemID": "ep-1",
                "PayableAmount": "1.20",
                "SavingPlanDeductionDiscountTotalAmount": "0.60",
            },
            {
                "ExpenseDate": "2026-08-01",
                "Project": "demo-alpha",
                "ProductZh": "方舟",
                "ConfigName": "input",
                "SplitItemID": "ep-1",
                "PayableAmount": "0.80",
                "SavingPlanDeductionDiscountTotalAmount": "0.40",
            },
            {
                "ExpenseDate": "2026-08-02",
                "ProjectDisplayName": "demo-alpha",
                "ProductZh": "方舟",
                "ConfigName": "input",
                "SplitItemID": "ep-1",
                "PayableAmount": "3",
                "SavingPlanDeductionDiscountTotalAmount": "1.00",
            },
        ]
        mapped = [map_bill_item(item, bill_period="2026-08") for item in items]
        day = aggregate_rows(mapped, month=False)
        month = aggregate_rows(mapped, month=True)
        self.assertEqual(len(day), 2)
        day_by_date = {row["日期"]: row["费用"] for row in day}
        saving_by_date = {row["日期"]: row["节省计划抵扣总额"] for row in day}
        self.assertEqual(day_by_date["2026-08-01"], 2.0)
        self.assertEqual(day_by_date["2026-08-02"], 3.0)
        self.assertEqual(saving_by_date["2026-08-01"], 1.0)
        self.assertEqual(saving_by_date["2026-08-02"], 1.0)
        self.assertEqual(len(month), 1)
        self.assertEqual(month[0]["日期"], "2026-08")
        self.assertEqual(month[0]["费用"], 5.0)
        self.assertEqual(month[0]["节省计划抵扣总额"], 2.0)
        self.assertEqual(month[0]["分拆项ID"], "ep-1")

    def test_blank_or_default_project_skipped(self):
        self.assertTrue(is_blank_or_default_project(""))
        self.assertTrue(is_blank_or_default_project("default"))
        self.assertTrue(is_blank_or_default_project("默认项目"))
        self.assertTrue(is_blank_or_default_project("default(默认项目)"))
        self.assertFalse(is_blank_or_default_project("demo-alpha"))
        client = MagicMock()
        client.request.return_value = {
            "List": [
                {
                    "ExpenseDate": "2026-08-01",
                    "Project": "",
                    "ProductZh": "方舟",
                    "ConfigName": "in",
                    "PayableAmount": "1",
                },
                {
                    "ExpenseDate": "2026-08-01",
                    "Project": "default",
                    "ProjectDisplayName": "默认项目",
                    "ProductZh": "方舟",
                    "ConfigName": "in",
                    "PayableAmount": "2",
                },
                {
                    "ExpenseDate": "2026-08-01",
                    "Project": "demo-alpha",
                    "ProjectDisplayName": "demo-alpha",
                    "ProductZh": "方舟",
                    "ConfigName": "in",
                    "PayableAmount": "3",
                },
            ],
            "Total": 3,
        }
        stats = collect_volc_bills("ak", "sk", bill_period="2026-08", client=client)
        self.assertEqual(stats.skipped_empty_or_default, 2)
        self.assertEqual(stats.day_rows, 1)
        self.assertEqual(stats.rows_day[0]["项目名"], "demo-alpha")
        self.assertEqual(stats.rows_day[0]["费用"], 3.0)

    def test_save_workbook_two_sheets(self):
        from lib.volc_billing import BillCollectStats

        with tempfile.TemporaryDirectory() as tmp:
            dest = Path(tmp) / "上游账单_火山云_20260831120000.xlsx"
            stats = BillCollectStats(
                bill_period="2026-08",
                source_action="ListSplitBillDetail",
                split_used=True,
                rows_day=[
                    {
                        "日期": "2026-08-01",
                        "项目名": "p",
                        "产品": "方舟",
                        "配置名称": "c",
                        "分拆项ID": "x",
                        "费用": 1.5,
                        "节省计划抵扣总额": 0.75,
                    }
                ],
                rows_month=[
                    {
                        "日期": "2026-08",
                        "项目名": "p",
                        "产品": "方舟",
                        "配置名称": "c",
                        "分拆项ID": "x",
                        "费用": 1.5,
                        "节省计划抵扣总额": 0.75,
                    }
                ],
            )
            save_bill_workbook(dest, stats)
            from openpyxl import load_workbook

            wb = load_workbook(dest)
            try:
                self.assertEqual(wb.sheetnames[:2], [SHEET_DAY, SHEET_MONTH])
                self.assertIn(SHEET_NOTE, wb.sheetnames)
                headers = [cell.value for cell in next(wb[SHEET_DAY].iter_rows(min_row=1, max_row=1))]
                self.assertEqual(headers, ["日期", "项目名", "产品", "配置名称", "分拆项ID", "费用（¥）", "节省计划抵扣总额（¥）"])
                self.assertEqual(wb[SHEET_DAY]["A2"].value, "2026-08-01")
                self.assertEqual(wb[SHEET_MONTH]["A2"].value, "2026-08")
                self.assertEqual(wb[SHEET_DAY]["A3"].value, "合计")
                self.assertEqual(wb[SHEET_DAY]["F3"].value, "=SUM(F2:F2)")
                self.assertIn("¥", str(wb[SHEET_DAY]["F2"].number_format))
                self.assertEqual(wb[SHEET_MONTH]["A3"].value, "合计")
                self.assertEqual(wb[SHEET_MONTH]["F3"].value, "=SUM(F2:F2)")
                self.assertEqual(wb[SHEET_DAY]["G3"].value, "=SUM(G2:G2)")
                self.assertEqual(wb[SHEET_MONTH]["G3"].value, "=SUM(G2:G2)")
                self.assertEqual(wb[SHEET_DAY].auto_filter.ref, "A1:G2")
                self.assertEqual(wb[SHEET_MONTH].auto_filter.ref, "A1:G2")
            finally:
                wb.close()

    def test_collect_falls_back_to_bill_detail(self):
        client = MagicMock()

        def fake_request(**kwargs):
            action = kwargs.get("action")
            if action == "ListSplitBillDetail":
                raise VolcOpenAPIError("split not open", code="OperationDenied")
            return {
                "List": [
                    {
                        "ExpenseDate": "2026-08-10",
                        "Project": "demo",
                        "ProductZh": "方舟",
                        "ConfigName": "out",
                        "PayableAmount": "4.5",
                        "SavingPlanDeductionDiscountTotalAmount": "2.5",
                    }
                ],
                "Total": 1,
            }

        client.request.side_effect = fake_request
        stats = collect_volc_bills("ak", "sk", bill_period="2026-08", client=client)
        self.assertEqual(stats.source_action, "ListBillDetail")
        self.assertFalse(stats.split_used)
        self.assertEqual(stats.day_rows, 1)
        self.assertEqual(stats.rows_day[0]["费用"], 4.5)
        self.assertEqual(stats.rows_day[0]["节省计划抵扣总额"], 2.5)
        self.assertEqual(stats.rows_month[0]["日期"], "2026-08")

    def test_both_apis_denied_explains_iam(self):
        client = MagicMock()
        client.request.side_effect = VolcOpenAPIError("no auth", code="AccessDenied")
        with self.assertRaises(VolcOpenAPIError) as ctx:
            collect_volc_bills("ak", "sk", bill_period="2026-08", client=client)
        self.assertIn("账单查询权限", str(ctx.exception))

    def test_filename_includes_month_label(self):
        self.assertEqual(bill_month_label("2026-09"), "9月")
        self.assertEqual(bill_month_label("2026-08"), "8月")
        name = build_bill_filename(when=datetime(2026, 9, 16, 10, 28, 22), bill_period="2026-09")
        self.assertEqual(name, "上游账单_火山云_9月_20260916102822.xlsx")

    def test_expand_bill_period_range(self):
        self.assertEqual(
            expand_bill_period_range("2026-07", "2026-09"),
            ["2026-07", "2026-08", "2026-09"],
        )
        self.assertEqual(
            expand_bill_period_range("2026-09", "2026-07"),
            ["2026-07", "2026-08", "2026-09"],
        )
        self.assertEqual(expand_bill_period_range("2025-12", "2026-01"), ["2025-12", "2026-01"])
        self.assertEqual(len(expand_bill_period_range("2024-01", "2025-12")), 24)
        with self.assertRaises(ValueError):
            expand_bill_period_range("2024-01", "2026-01")

    def test_bill_periods_from_payload(self):
        self.assertEqual(bill_periods_from_payload({"bill_period": "2026-08"}), ["2026-08"])
        self.assertEqual(
            bill_periods_from_payload({"bill_period_from": "2026-07", "bill_period_to": "2026-08"}),
            ["2026-07", "2026-08"],
        )
        self.assertEqual(
            bill_periods_from_payload({"bill_periods": ["2026-05", "2026-07"]}),
            ["2026-05", "2026-07"],
        )
        self.assertEqual(
            bill_periods_from_payload({"bill_periods": "2026-05, 2026-06"}),
            ["2026-05", "2026-06"],
        )


if __name__ == "__main__":
    unittest.main()
