# -*- coding: utf-8 -*-
"""⑥ 智能分析：上下游密钥匹配。"""

from __future__ import annotations

import sys
import tempfile
import unittest
from datetime import datetime, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from lib.channel_balance_query import export_channel_balance_xlsx
from lib.data_summary import CHANNEL_HEADERS_WITH_KEY
from lib.data_summary_exporter import save_workbook_data
from lib.overview_analyzer import (
    ACTION_HEADERS,
    ANALYSIS_STEM,
    BATCH_DISABLE_HEADERS,
    SHEET_ACTIONS,
    SHEET_BALANCE,
    SHEET_BATCH_DISABLE,
    SHEET_DISABLED_ACTIONS,
    USAGE_COMMON,
    USAGE_UNCOMMON,
    USAGE_UNUSED,
    analyze_overview_workbook,
    build_batch_disable_rows,
    classify_usage,
    classify_vendor,
    format_name_group_line,
    is_action_already_disabled,
    is_new_deploy,
    parse_usd,
    partition_disabled_actions,
    run_overview_analysis,
    secrets_match,
    usage_action_level,
)
from lib.overview_exporter import OVERVIEW_STEM, save_overview_workbook


class OverviewAnalyzerTests(unittest.TestCase):
    def test_secrets_match_uuid_fixture(self):
        self.assertTrue(
            secrets_match(
                "00000000-0000-0000-0000-000000000000",
                "00000000-0000-0000-0000-000000000000",
            )
        )
        self.assertFalse(secrets_match("aaa", "bbb"))

    def test_classify_vendor(self):
        self.assertEqual(classify_vendor({"类型": "字节火山方舟、豆包通用", "base_url": ""}), "火山云")
        self.assertEqual(
            classify_vendor({"类型": "自定义渠道", "base_url": "https://ark.cn-beijing.volces.com"}),
            "火山云",
        )
        self.assertEqual(classify_vendor({"类型": "Azure OpenAI", "base_url": ""}), "Azure OpenAI")
        self.assertEqual(classify_vendor({"类型": "自定义渠道", "base_url": "https://chat.example.com"}), "外部/其它")

    def test_down_no_up_and_up_no_down(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            dest = root / "excel" / f"{OVERVIEW_STEM}_src.xlsx"
            save_workbook_data(
                dest,
                channel_env_rows={
                    "test": [
                        {
                            "名称": "stale-ep",
                            "分组": "g",
                            "类型": "字节火山方舟、豆包通用",
                            "状态": "已启用",
                            "模型列表": "ep-20990101000000-zzzzz",
                            "模型重定向": "",
                            "base_url": "https://ark.cn-beijing.volces.com",
                            "密钥": "aaaaaaaa-bbbb-cccc-dddd-eeeeeeeeeeee",
                            "部署时间": "",
                        },
                        {
                            "名称": "live-ep",
                            "分组": "g",
                            "类型": "字节火山方舟、豆包通用",
                            "状态": "已启用",
                            "模型列表": "ep-20260101000000-abcde",
                            "模型重定向": "",
                            "base_url": "https://ark.cn-beijing.volces.com",
                            "密钥": "11111111-2222-3333-4444-555555555555",
                            "部署时间": "",
                        },
                    ],
                    "uat": [],
                    "prod": [],
                },
                token_env_rows={
                    "test": [
                        {
                            "名称": "stale-token-张三",
                            "项目环境": "g",
                            "状态": "已启用",
                            "剩余额度": "",
                            "分组": "g",
                            "可用模型": "stale-ep",
                            "包含渠道": "stale-ep",
                            "IP限制": "",
                            "创建时间": "",
                            "过期时间": "永不过期",
                        }
                    ],
                    "uat": [],
                    "prod": [],
                },
                channel_headers=CHANNEL_HEADERS_WITH_KEY,
            )
            save_overview_workbook(
                dest,
                gateway_path=dest,
                volc_rows=[
                    {
                        "接入点名称": "live-ep",
                        "部署名": "ep-20260101000000-abcde",
                        "API Key": "ark-test-key",
                        "项目方": "demo",
                        "联系人": "n",
                    },
                    {
                        "接入点名称": "orphan-up",
                        "部署名": "ep-20260101000000-orphan",
                        "API Key": "ark-test-key",
                        "项目方": "demo",
                        "联系人": "n",
                    },
                ],
                azure_rows=[],
            )
            payload = analyze_overview_workbook(dest)
            down_names = {row["渠道名称"] for row in payload["down_rows"]}
            up_names = {row["上游名称"] for row in payload["up_rows"]}
            self.assertIn("stale-ep", down_names)
            self.assertNotIn("live-ep", down_names)
            self.assertIn("orphan-up", up_names)
            self.assertNotIn("live-ep", up_names)
            stale_actions = [
                row for row in payload["actions"] if row.get("渠道名称") == "stale-ep" and row.get("类别") == "下游无上游"
            ]
            self.assertTrue(stale_actions)
            self.assertEqual(stale_actions[0].get("令牌名称"), "stale-token-张三")
            self.assertTrue(stale_actions[0].get("密钥摘要"))
            self.assertEqual(ACTION_HEADERS[3:6], ["渠道名称", "令牌名称", "密钥摘要"])
            saved, _ = run_overview_analysis(root, dest)
            self.assertTrue(saved.name.startswith(ANALYSIS_STEM))
            self.assertTrue(saved.exists())
            self.assertEqual(payload["stats"].get("balance_file"), "")
            self.assertEqual(payload["stats"].get("balance_unused"), 0)
            disable = payload.get("disable_rows") or []
            self.assertTrue(disable)
            self.assertEqual(disable[0].get("test渠道"), "stale-ep\tg")
            self.assertEqual(disable[0].get("test令牌"), "stale-token-张三\tg")
            self.assertFalse(any(row.get("test渠道") == "live-ep\tg" for row in disable))
            from openpyxl import load_workbook

            wb = load_workbook(saved)
            try:
                action_idx = wb.sheetnames.index(SHEET_ACTIONS)
                self.assertEqual(wb.sheetnames[action_idx + 1], SHEET_DISABLED_ACTIONS)
                self.assertEqual(wb.sheetnames[action_idx + 2], SHEET_BATCH_DISABLE)
                headers = [cell.value for cell in next(wb[SHEET_BATCH_DISABLE].iter_rows(min_row=1, max_row=1))]
                self.assertEqual(headers, BATCH_DISABLE_HEADERS)
                done_headers = [cell.value for cell in next(wb[SHEET_DISABLED_ACTIONS].iter_rows(min_row=1, max_row=1))]
                self.assertEqual(done_headers, ACTION_HEADERS)
            finally:
                wb.close()

    def test_parse_usd_and_classify_usage(self):
        self.assertEqual(parse_usd("$0"), 0.0)
        self.assertEqual(parse_usd("$0.50"), 0.5)
        self.assertEqual(parse_usd("$1,234.56"), 1234.56)
        self.assertIsNone(parse_usd(""))
        self.assertEqual(classify_usage(0), USAGE_UNUSED)
        self.assertEqual(classify_usage(0.2), USAGE_UNCOMMON)
        self.assertEqual(classify_usage(0.999), USAGE_UNCOMMON)
        self.assertEqual(classify_usage(1.0), USAGE_COMMON)
        self.assertEqual(classify_usage(12), USAGE_COMMON)
        self.assertEqual(usage_action_level(USAGE_UNUSED, month_idle=False), "中")
        self.assertEqual(usage_action_level(USAGE_UNUSED, month_idle=True), "中")
        self.assertEqual(usage_action_level(USAGE_UNCOMMON, month_idle=False), "低")
        self.assertEqual(usage_action_level(USAGE_UNCOMMON, month_idle=True), "中")
        self.assertEqual(usage_action_level(USAGE_COMMON, month_idle=True), "低")

    def test_balance_unused_and_uncommon(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            dest = root / "excel" / f"{OVERVIEW_STEM}_src.xlsx"
            save_workbook_data(
                dest,
                channel_env_rows={
                    "test": [
                        {
                            "名称": "idle-ch",
                            "分组": "g",
                            "类型": "自定义渠道",
                            "状态": "已启用",
                            "模型列表": "m",
                            "模型重定向": "",
                            "base_url": "https://chat.example.com",
                            "密钥": "k1",
                            "部署时间": "",
                        },
                        {
                            "名称": "rare-ch",
                            "分组": "g",
                            "类型": "自定义渠道",
                            "状态": "已启用",
                            "模型列表": "m",
                            "模型重定向": "",
                            "base_url": "https://chat.example.com",
                            "密钥": "k2",
                            "部署时间": "",
                        },
                        {
                            "名称": "hot-ch",
                            "分组": "g",
                            "类型": "自定义渠道",
                            "状态": "已启用",
                            "模型列表": "m",
                            "模型重定向": "",
                            "base_url": "https://chat.example.com",
                            "密钥": "k3",
                            "部署时间": "",
                        },
                        {
                            "名称": "off-ch",
                            "分组": "g",
                            "类型": "自定义渠道",
                            "状态": "已禁用",
                            "模型列表": "m",
                            "模型重定向": "",
                            "base_url": "https://chat.example.com",
                            "密钥": "k4",
                            "部署时间": "",
                        },
                    ],
                    "uat": [],
                    "prod": [],
                },
                token_env_rows={"test": [], "uat": [], "prod": []},
                channel_headers=CHANNEL_HEADERS_WITH_KEY,
            )
            save_overview_workbook(dest, gateway_path=dest, volc_rows=[], azure_rows=[])
            items = [
                _balance_item("test", "idle-ch", "$0", "已启用"),
                _balance_item("test", "rare-ch", "$0.50", "已启用"),
                _balance_item("test", "hot-ch", "$12.00", "已启用"),
                _balance_item("test", "off-ch", "$0", "已禁用"),
            ]
            data, _ = export_channel_balance_xlsx(items, include_activity=True)
            (root / "excel" / "渠道余额_20260825120000.xlsx").write_bytes(data)
            payload = analyze_overview_workbook(
                dest,
                root_dir=root,
                policy={"na_channel_activity_enabled": True},
            )
            by_name = {row["渠道名称"]: row for row in payload["balance_rows"]}
            self.assertEqual(by_name["idle-ch"]["判定"], USAGE_UNUSED)
            self.assertEqual(by_name["rare-ch"]["判定"], USAGE_UNCOMMON)
            self.assertEqual(by_name["hot-ch"]["判定"], USAGE_COMMON)
            self.assertEqual(payload["stats"]["balance_unused"], 2)
            self.assertEqual(payload["stats"]["balance_uncommon"], 1)
            action_names = {
                row["渠道名称"]
                for row in payload["actions"]
                if row.get("类别") == "渠道余额"
            }
            self.assertIn("idle-ch", action_names)
            self.assertIn("rare-ch", action_names)
            self.assertNotIn("hot-ch", action_names)
            self.assertNotIn("off-ch", action_names)
            action_levels = {
                row["渠道名称"]: row["级别"]
                for row in payload["actions"]
                if row.get("类别") == "渠道余额"
            }
            self.assertEqual(action_levels["idle-ch"], "中")
            self.assertEqual(action_levels["rare-ch"], "低")
            done_names = {
                row["渠道名称"]
                for row in payload.get("disabled_actions") or []
                if row.get("类别") == "渠道余额"
            }
            self.assertIn("off-ch", done_names)
            self.assertNotIn("idle-ch", done_names)
            saved, _ = run_overview_analysis(root, dest)
            from openpyxl import load_workbook

            wb = load_workbook(saved)
            self.assertIn(SHEET_BALANCE, wb.sheetnames)
            wb.close()

    def test_activity_blocks_usage_disable_when_any_model_active_or_na(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            dest = root / "excel" / f"{OVERVIEW_STEM}_act.xlsx"
            save_workbook_data(
                dest,
                channel_env_rows={
                    "test": [
                        {
                            "名称": "mixed-act",
                            "分组": "g",
                            "类型": "自定义渠道",
                            "状态": "已启用",
                            "模型列表": "a,b",
                            "模型重定向": "",
                            "base_url": "https://chat.example.com",
                            "密钥": "k1",
                            "部署时间": "",
                        },
                        {
                            "名称": "all-idle",
                            "分组": "g",
                            "类型": "自定义渠道",
                            "状态": "已启用",
                            "模型列表": "only",
                            "模型重定向": "",
                            "base_url": "https://chat.example.com",
                            "密钥": "k2",
                            "部署时间": "",
                        },
                        {
                            "名称": "na-act",
                            "分组": "g",
                            "类型": "自定义渠道",
                            "状态": "已启用",
                            "模型列表": "x",
                            "模型重定向": "",
                            "base_url": "https://chat.example.com",
                            "密钥": "k3",
                            "部署时间": "",
                        },
                    ],
                    "uat": [],
                    "prod": [],
                },
                token_env_rows={"test": [], "uat": [], "prod": []},
                channel_headers=CHANNEL_HEADERS_WITH_KEY,
            )
            save_overview_workbook(dest, gateway_path=dest, volc_rows=[], azure_rows=[])
            items = [
                _balance_item(
                    "test",
                    "mixed-act",
                    "$0",
                    "已启用",
                    activity_label="a: 有活跃\nb: 近一月无活跃",
                ),
                _balance_item(
                    "test",
                    "all-idle",
                    "$0",
                    "已启用",
                    activity_label="only: 近一月无活跃",
                ),
                _balance_item(
                    "test",
                    "na-act",
                    "$0",
                    "已启用",
                    activity_label="x: N/A",
                ),
            ]
            data, _ = export_channel_balance_xlsx(items, include_activity=True)
            (root / "excel" / "渠道余额_20260825120001.xlsx").write_bytes(data)
            payload = analyze_overview_workbook(
                dest,
                root_dir=root,
                policy={"na_channel_activity_enabled": True},
            )
            by_name = {row["渠道名称"]: row for row in payload["balance_rows"]}
            self.assertIn("不禁用", by_name["mixed-act"]["建议"])
            self.assertIn("不据此禁用", by_name["na-act"]["建议"])
            self.assertNotIn("不禁用", by_name["all-idle"]["建议"])
            self.assertNotIn("不据此禁用", by_name["all-idle"]["建议"])
            action_names = {
                row["渠道名称"]
                for row in payload["actions"]
                if row.get("类别") == "渠道余额"
            }
            self.assertNotIn("mixed-act", action_names)
            self.assertNotIn("na-act", action_names)
            self.assertIn("all-idle", action_names)

    def test_uncommon_and_month_idle_is_medium(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            dest = root / "excel" / f"{OVERVIEW_STEM}_uncommon.xlsx"
            save_workbook_data(
                dest,
                channel_env_rows={
                    "test": [
                        {
                            "名称": "rare-idle",
                            "分组": "g",
                            "类型": "自定义渠道",
                            "状态": "已启用",
                            "模型列表": "only",
                            "模型重定向": "",
                            "base_url": "https://chat.example.com",
                            "密钥": "k1",
                            "部署时间": "",
                        },
                        {
                            "名称": "rare-legacy",
                            "分组": "g",
                            "类型": "自定义渠道",
                            "状态": "已启用",
                            "模型列表": "only",
                            "模型重定向": "",
                            "base_url": "https://chat.example.com",
                            "密钥": "k2",
                            "部署时间": "",
                        },
                    ],
                    "uat": [],
                    "prod": [],
                },
                token_env_rows={"test": [], "uat": [], "prod": []},
                channel_headers=CHANNEL_HEADERS_WITH_KEY,
            )
            save_overview_workbook(dest, gateway_path=dest, volc_rows=[], azure_rows=[])
            items = [
                _balance_item(
                    "test",
                    "rare-idle",
                    "$0.50",
                    "已启用",
                    activity_label="only: 近一月无活跃",
                ),
                _balance_item("test", "rare-legacy", "$0.50", "已启用"),
            ]
            data, _ = export_channel_balance_xlsx(items, include_activity=True)
            (root / "excel" / "渠道余额_20260825120002.xlsx").write_bytes(data)
            payload = analyze_overview_workbook(
                dest,
                root_dir=root,
                policy={"na_channel_activity_enabled": True},
            )
            action_by_name = {
                row["渠道名称"]: row
                for row in payload["actions"]
                if row.get("类别") == "渠道余额"
            }
            self.assertEqual(action_by_name["rare-idle"]["级别"], "中")
            self.assertIn("近一月无活跃", action_by_name["rare-idle"]["问题"])
            self.assertEqual(action_by_name["rare-legacy"]["级别"], "低")
            disable_text = " ".join(
                str(row.get("test渠道") or "") for row in payload.get("disable_rows") or []
            )
            self.assertIn("rare-idle", disable_text)
            self.assertNotIn("rare-legacy", disable_text)

    def test_activity_disabled_keeps_legacy_usage_rules(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            dest = root / "excel" / f"{OVERVIEW_STEM}_legacy.xlsx"
            save_workbook_data(
                dest,
                channel_env_rows={
                    "test": [
                        {
                            "名称": "rare-legacy",
                            "分组": "g",
                            "类型": "自定义渠道",
                            "状态": "已启用",
                            "模型列表": "only",
                            "模型重定向": "",
                            "base_url": "https://chat.example.com",
                            "密钥": "k1",
                            "部署时间": "",
                        }
                    ],
                    "uat": [],
                    "prod": [],
                },
                token_env_rows={"test": [], "uat": [], "prod": []},
                channel_headers=CHANNEL_HEADERS_WITH_KEY,
            )
            save_overview_workbook(dest, gateway_path=dest, volc_rows=[], azure_rows=[])
            data, _ = export_channel_balance_xlsx(
                [_balance_item("test", "rare-legacy", "$0.50", "已启用")],
                include_activity=True,
            )
            (root / "excel" / "渠道余额_20260825120003.xlsx").write_bytes(data)
            payload = analyze_overview_workbook(dest, root_dir=root)
            action = next(
                row
                for row in payload["actions"]
                if row.get("类别") == "渠道余额" and row.get("渠道名称") == "rare-legacy"
            )
            self.assertEqual(action["级别"], "低")
            self.assertNotIn("近一月无活跃", action["问题"])

    def test_channel_sheet_activity_feeds_analysis(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            dest = root / "excel" / f"{OVERVIEW_STEM}_channel-activity.xlsx"
            save_workbook_data(
                dest,
                channel_env_rows={
                    "test": [
                        {
                            "名称": "rare-channel-activity",
                            "分组": "g",
                            "类型": "自定义渠道",
                            "状态": "已启用",
                            "模型列表": "only",
                            "模型重定向": "",
                            "base_url": "https://chat.example.com",
                            "密钥": "k1",
                            "部署时间": "",
                            "近一月活跃": "only: 近一月无活跃",
                        }
                    ],
                    "uat": [],
                    "prod": [],
                },
                token_env_rows={"test": [], "uat": [], "prod": []},
                channel_headers=CHANNEL_HEADERS_WITH_KEY,
            )
            save_overview_workbook(dest, gateway_path=dest, volc_rows=[], azure_rows=[])
            data, _ = export_channel_balance_xlsx(
                [_balance_item("test", "rare-channel-activity", "$0.50", "已启用")],
                include_activity=False,
            )
            (root / "excel" / "渠道余额_20260825120004.xlsx").write_bytes(data)
            payload = analyze_overview_workbook(
                dest,
                root_dir=root,
                policy={"na_channel_activity_enabled": True},
            )
            action = next(
                row
                for row in payload["actions"]
                if row.get("类别") == "渠道余额"
                and row.get("渠道名称") == "rare-channel-activity"
            )
            self.assertEqual(action["级别"], "中")
            self.assertEqual(action["近一月活跃"], "only: 近一月无活跃")

    def test_no_used_column_skips_usage(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            dest = root / "excel" / f"{OVERVIEW_STEM}_src.xlsx"
            save_workbook_data(
                dest,
                channel_env_rows={
                    "test": [
                        {
                            "名称": "idle-ch",
                            "分组": "g",
                            "类型": "自定义渠道",
                            "状态": "已启用",
                            "模型列表": "m",
                            "模型重定向": "",
                            "base_url": "https://chat.example.com",
                            "密钥": "k1",
                            "部署时间": "",
                        },
                    ],
                    "uat": [],
                    "prod": [],
                },
                token_env_rows={"test": [], "uat": [], "prod": []},
                channel_headers=CHANNEL_HEADERS_WITH_KEY,
            )
            save_overview_workbook(dest, gateway_path=dest, volc_rows=[], azure_rows=[])
            data, _ = export_channel_balance_xlsx(
                [_balance_item("test", "idle-ch", "$0", "已启用")],
                include_used=False,
            )
            (root / "excel" / "渠道余额_20260825120000.xlsx").write_bytes(data)
            payload = analyze_overview_workbook(dest, root_dir=root)
            self.assertEqual(payload["balance_rows"], [])
            self.assertEqual(payload["stats"]["balance_unused"], 0)

    def test_idle_activity_creates_action_without_used_column(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            dest = root / "excel" / f"{OVERVIEW_STEM}_idle.xlsx"
            save_workbook_data(
                dest,
                channel_env_rows={
                    "test": [
                        {
                            "名称": "idle-month",
                            "分组": "g",
                            "类型": "自定义渠道",
                            "状态": "已启用",
                            "模型列表": "m",
                            "模型重定向": "",
                            "base_url": "https://chat.example.com",
                            "密钥": "k1",
                            "部署时间": "",
                            "近一月活跃": "m: 近一月无活跃",
                        }
                    ],
                    "uat": [],
                    "prod": [],
                },
                token_env_rows={"test": [], "uat": [], "prod": []},
                channel_headers=CHANNEL_HEADERS_WITH_KEY,
            )
            save_overview_workbook(dest, gateway_path=dest, volc_rows=[], azure_rows=[])
            data, _ = export_channel_balance_xlsx(
                [_balance_item("test", "idle-month", "$0", "已启用")],
                include_used=False,
                include_activity=True,
            )
            (root / "excel" / "渠道余额_20260825120005.xlsx").write_bytes(data)
            payload = analyze_overview_workbook(
                dest,
                root_dir=root,
                policy={"na_channel_activity_enabled": True},
            )
            action = next(
                row
                for row in payload["actions"]
                if row.get("类别") == "渠道活跃"
                and row.get("渠道名称") == "idle-month"
            )
            self.assertEqual(action["级别"], "中")
            self.assertEqual(action["问题"], "近一月全部模型无活跃")
            self.assertEqual(payload["stats"]["activity_idle"], 1)

    def test_is_new_deploy_window(self):
        now = datetime(2026, 8, 25, 16, 0, 0)
        self.assertTrue(is_new_deploy("2026-08-20 10:00:00", now=now))
        self.assertFalse(is_new_deploy("2026-08-01 10:00:00", now=now))
        self.assertFalse(is_new_deploy("", now=now))
        self.assertFalse(is_new_deploy("2026-08-15 16:00:00", now=now))

    def test_new_deploy_relaxes_down_up_and_unused(self):
        now = datetime(2026, 8, 25, 16, 0, 0)
        fresh = (now - timedelta(days=3)).strftime("%Y-%m-%d %H:%M:%S")
        old = (now - timedelta(days=40)).strftime("%Y-%m-%d %H:%M:%S")
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            dest = root / "excel" / f"{OVERVIEW_STEM}_src.xlsx"
            save_workbook_data(
                dest,
                channel_env_rows={
                    "test": [
                        {
                            "名称": "fresh-stale",
                            "分组": "g",
                            "类型": "字节火山方舟、豆包通用",
                            "状态": "已启用",
                            "模型列表": "ep-20990101000000-zzzzz",
                            "模型重定向": "",
                            "base_url": "https://ark.cn-beijing.volces.com",
                            "密钥": "aaaaaaaa-bbbb-cccc-dddd-eeeeeeeeeeee",
                            "部署时间": fresh,
                        },
                        {
                            "名称": "old-stale",
                            "分组": "g",
                            "类型": "字节火山方舟、豆包通用",
                            "状态": "已启用",
                            "模型列表": "ep-20990101000000-yyyyy",
                            "模型重定向": "",
                            "base_url": "https://ark.cn-beijing.volces.com",
                            "密钥": "bbbbbbbb-bbbb-cccc-dddd-eeeeeeeeeeee",
                            "部署时间": old,
                        },
                    ],
                    "uat": [],
                    "prod": [],
                },
                token_env_rows={"test": [], "uat": [], "prod": []},
                channel_headers=CHANNEL_HEADERS_WITH_KEY,
            )
            save_overview_workbook(
                dest,
                gateway_path=dest,
                volc_rows=[
                    {
                        "接入点名称": "orphan-fresh",
                        "部署名": "ep-20260822000000-new",
                        "API Key": "ark-test-key",
                        "项目方": "demo",
                        "联系人": "n",
                        "部署时间": fresh,
                    },
                ],
                azure_rows=[],
            )
            items = [
                _balance_item("test", "fresh-stale", "$0", "已启用"),
                _balance_item("test", "old-stale", "$0", "已启用"),
            ]
            data, _ = export_channel_balance_xlsx(items)
            (root / "excel" / "渠道余额_20260825120000.xlsx").write_bytes(data)
            payload = analyze_overview_workbook(dest, root_dir=root, now=now)
            down_actions = {
                row["渠道名称"]: row
                for row in payload["actions"]
                if row.get("类别") == "下游无上游"
            }
            self.assertEqual(down_actions["fresh-stale"]["级别"], "中")
            self.assertIn("新部署未满 10 天", down_actions["fresh-stale"]["问题"])
            self.assertEqual(down_actions["old-stale"]["级别"], "高")
            up_actions = [
                row for row in payload["actions"] if row.get("类别") == "上游无下游"
            ]
            self.assertTrue(up_actions)
            self.assertEqual(up_actions[0]["级别"], "低")
            unused = {
                row["渠道名称"]: row
                for row in payload["actions"]
                if row.get("类别") == "渠道余额"
            }
            self.assertEqual(unused["fresh-stale"]["级别"], "低")
            self.assertEqual(unused["old-stale"]["级别"], "中")
            self.assertGreaterEqual(payload["stats"].get("relaxed_new_deploy") or 0, 1)
            disable_channels = {row.get("test渠道") for row in payload.get("disable_rows") or []}
            self.assertIn("old-stale\tg", disable_channels)
            self.assertNotIn("fresh-stale\tg", disable_channels)

    def test_format_name_group_line(self):
        self.assertEqual(format_name_group_line("gpt-4o", "azure"), "gpt-4o\tazure")
        self.assertEqual(format_name_group_line("  a  ", " b "), "a\tb")
        self.assertEqual(format_name_group_line("only-name", ""), "")

    def test_build_batch_disable_rows_filters(self):
        from lib.overview_analyzer import _NEW_DEPLOY_HOLD

        rows = build_batch_disable_rows(
            [
                {
                    "级别": "高",
                    "类别": "下游无上游",
                    "站点": "test",
                    "渠道名称": "ch-a",
                    "分组": "g",
                    "状态": "已启用",
                    "令牌名称": "shared-tk，owned-tk",
                    "建议": "建议在 test 禁用或删除该渠道",
                },
                {
                    "级别": "高",
                    "类别": "下游无上游",
                    "站点": "test",
                    "渠道名称": "ch-new",
                    "分组": "g",
                    "状态": "已启用",
                    "令牌名称": "",
                    "建议": "建议禁用；" + _NEW_DEPLOY_HOLD,
                },
                {
                    "级别": "中",
                    "类别": "上游无下游",
                    "站点": "三站皆无",
                    "渠道名称": "orphan",
                    "分组": "g",
                    "建议": "可在火山侧停用密钥",
                },
                {
                    "级别": "中",
                    "类别": "令牌对照",
                    "站点": "uat",
                    "令牌名称": "missing-tk",
                    "分组": "grp",
                    "建议": "渠道已删或改名时，同步改令牌或删除令牌",
                },
                {
                    "级别": "低",
                    "类别": "渠道余额",
                    "站点": "prod",
                    "渠道名称": "rare-ch",
                    "分组": "g",
                    "状态": "已启用",
                    "建议": "确认是否仍需保留",
                },
            ],
            {
                "test": [
                    {"名称": "shared-tk", "分组": "g", "包含渠道": "无限制", "状态": "已启用"},
                    {"名称": "owned-tk", "分组": "g", "包含渠道": "ch-a", "状态": "已启用"},
                ],
                "uat": [
                    {"名称": "missing-tk", "分组": "grp", "包含渠道": "gone", "状态": "已启用"},
                ],
                "prod": [],
            },
        )
        self.assertEqual(rows[0]["test渠道"], "ch-a\tg")
        self.assertEqual(rows[0]["test令牌"], "owned-tk\tg")
        self.assertEqual(rows[0]["uat令牌"], "missing-tk\tgrp")
        self.assertFalse(any("ch-new" in str(row.get("test渠道") or "") for row in rows))
        self.assertFalse(any("shared-tk" in str(row.get("test令牌") or "") for row in rows))
        self.assertFalse(any(row.get("prod渠道") for row in rows))
        from lib.channel_batch_delete import parse_channel_delete_rows
        from lib.token_batch_delete import parse_token_name_group_rows

        channel_text = "\n".join(row["test渠道"] for row in rows if row.get("test渠道"))
        token_text = "\n".join(row["test令牌"] for row in rows if row.get("test令牌"))
        self.assertEqual(parse_channel_delete_rows(channel_text), [{"名称": "ch-a", "分组": "g"}])
        self.assertEqual(parse_token_name_group_rows(token_text), [{"名称": "owned-tk", "分组": "g"}])

    def test_batch_disable_skips_token_with_remaining_live_channel(self):
        rows = build_batch_disable_rows(
            [
                {
                    "级别": "中",
                    "类别": "渠道余额",
                    "站点": "test",
                    "渠道名称": "dsv4pro",
                    "分组": "biz",
                    "状态": "已启用",
                    "令牌名称": "dsv4pro&glm5.1-孙浏淘",
                    "建议": "建议在 test 禁用或删除该渠道",
                },
                {
                    "级别": "中",
                    "类别": "渠道余额",
                    "站点": "uat",
                    "渠道名称": "clinai-dev-gpt-5.1-codex",
                    "分组": "azure",
                    "状态": "已启用",
                    "令牌名称": "临研-gpt-5.1-费仕忆，Clinical-Assistant-dev-李梦珂",
                    "建议": "建议在 uat 禁用或删除该渠道",
                },
            ],
            {
                "test": [
                    {
                        "名称": "dsv4pro&glm5.1-孙浏淘",
                        "分组": "biz",
                        "包含渠道": "dsv4pro, glm-5.1",
                        "状态": "已启用",
                    }
                ],
                "uat": [
                    {
                        "名称": "临研-gpt-5.1-费仕忆",
                        "分组": "azure",
                        "包含渠道": "clinai-dev-gpt-5.1, clinai-dev-gpt-5.1-codex",
                        "状态": "已启用",
                    },
                    {
                        "名称": "Clinical-Assistant-dev-李梦珂",
                        "分组": "azure",
                        "包含渠道": "clinai-dev-gpt-5.1, clinai-dev-gpt-5.1-codex",
                        "状态": "已启用",
                    },
                ],
            },
            {
                "test": [
                    {"名称": "dsv4pro", "状态": "已启用"},
                    {"名称": "glm-5.1", "状态": "已启用"},
                ],
                "uat": [
                    {"名称": "clinai-dev-gpt-5.1", "状态": "已启用"},
                    {"名称": "clinai-dev-gpt-5.1-codex", "状态": "已启用"},
                ],
            },
        )
        self.assertEqual(rows[0]["test渠道"], "dsv4pro\tbiz")
        self.assertEqual(rows[0]["uat渠道"], "clinai-dev-gpt-5.1-codex\tazure")
        self.assertFalse(any(row.get("test令牌") for row in rows))
        self.assertFalse(any(row.get("uat令牌") for row in rows))

    def test_partition_disabled_actions(self):
        tokens = {
            "test": [
                {"名称": "tk-off", "分组": "g", "包含渠道": "ch-off", "状态": "已禁用"},
                {"名称": "tk-on", "分组": "g", "包含渠道": "ch-partial", "状态": "已启用"},
                {"名称": "tk-free", "分组": "g", "包含渠道": "无限制", "状态": "已启用"},
            ]
        }
        archived_row = {
            "类别": "下游无上游",
            "站点": "test",
            "渠道名称": "ch-off",
            "状态": "已禁用",
            "令牌名称": "tk-off",
        }
        pending_partial = {
            "类别": "下游无上游",
            "站点": "test",
            "渠道名称": "ch-partial",
            "状态": "已禁用",
            "令牌名称": "tk-on",
        }
        pending_live = {
            "类别": "下游无上游",
            "站点": "test",
            "渠道名称": "ch-live",
            "状态": "已启用",
            "令牌名称": "tk-on",
        }
        token_off = {
            "类别": "令牌对照",
            "站点": "test",
            "令牌名称": "tk-off",
            "状态": "已禁用",
        }
        up_keep = {
            "类别": "上游无下游",
            "站点": "三站皆无",
            "状态": "已禁用",
        }
        self.assertTrue(is_action_already_disabled(archived_row, tokens))
        self.assertFalse(is_action_already_disabled(pending_partial, tokens))
        self.assertFalse(is_action_already_disabled(pending_live, tokens))
        self.assertTrue(is_action_already_disabled(token_off, tokens))
        self.assertFalse(is_action_already_disabled(up_keep, tokens))
        no_token = {
            "类别": "渠道余额",
            "站点": "test",
            "渠道名称": "ch-off",
            "状态": "已禁用",
            "令牌名称": "",
        }
        self.assertTrue(is_action_already_disabled(no_token, tokens))
        free_only = {
            "类别": "已禁用仍登记",
            "站点": "test",
            "渠道名称": "ch-off",
            "状态": "已禁用",
            "令牌名称": "tk-free",
        }
        self.assertTrue(is_action_already_disabled(free_only, tokens))
        pending, archived = partition_disabled_actions(
            [archived_row, pending_partial, pending_live, token_off, up_keep],
            tokens,
        )
        self.assertEqual([row["渠道名称"] for row in pending if row.get("渠道名称")], ["ch-partial", "ch-live"])
        self.assertEqual(len(pending), 3)
        self.assertEqual(len(archived), 2)

    def test_disabled_channel_and_token_leave_actions_sheet(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            dest = root / "excel" / f"{OVERVIEW_STEM}_src.xlsx"
            save_workbook_data(
                dest,
                channel_env_rows={
                    "test": [
                        {
                            "名称": "off-stale",
                            "分组": "g",
                            "类型": "字节火山方舟、豆包通用",
                            "状态": "已禁用",
                            "模型列表": "ep-20990101000000-zzzzz",
                            "模型重定向": "",
                            "base_url": "https://ark.cn-beijing.volces.com",
                            "密钥": "aaaaaaaa-bbbb-cccc-dddd-eeeeeeeeeeee",
                            "部署时间": "",
                        }
                    ],
                    "uat": [],
                    "prod": [],
                },
                token_env_rows={
                    "test": [
                        {
                            "名称": "off-token",
                            "项目环境": "g",
                            "状态": "已禁用",
                            "剩余额度": "",
                            "分组": "g",
                            "可用模型": "off-stale",
                            "包含渠道": "off-stale",
                            "IP限制": "",
                            "创建时间": "",
                            "过期时间": "永不过期",
                        }
                    ],
                    "uat": [],
                    "prod": [],
                },
                channel_headers=CHANNEL_HEADERS_WITH_KEY,
            )
            save_overview_workbook(dest, gateway_path=dest, volc_rows=[], azure_rows=[])
            payload = analyze_overview_workbook(dest)
            action_names = {row["渠道名称"] for row in payload["actions"] if row.get("类别") == "下游无上游"}
            done_names = {
                row["渠道名称"]
                for row in payload.get("disabled_actions") or []
                if row.get("类别") == "下游无上游"
            }
            self.assertNotIn("off-stale", action_names)
            self.assertIn("off-stale", done_names)
            self.assertEqual(payload["stats"].get("disabled_actions"), 2)
            saved, _ = run_overview_analysis(root, dest)
            from openpyxl import load_workbook

            wb = load_workbook(saved)
            try:
                self.assertIn(SHEET_DISABLED_ACTIONS, wb.sheetnames)
                done_names_xlsx = {
                    row[3]
                    for row in wb[SHEET_DISABLED_ACTIONS].iter_rows(min_row=2, values_only=True)
                    if row and row[3]
                }
                self.assertIn("off-stale", done_names_xlsx)
                live_names_xlsx = {
                    row[3]
                    for row in wb[SHEET_ACTIONS].iter_rows(min_row=2, values_only=True)
                    if row and row[3]
                }
                self.assertNotIn("off-stale", live_names_xlsx)
            finally:
                wb.close()


def _balance_item(
    env: str,
    name: str,
    used_label: str,
    status: str,
    activity_label: str = "",
) -> dict:
    return {
        "env": env,
        "site_label": env,
        "id": 1,
        "name": name,
        "group": "g",
        "type_label": "自定义渠道",
        "status_label": status,
        "used_usd_label": used_label,
        "remain_usd_label": "$1.00",
        "created_time_label": "—",
        "balance_updated_time_label": "—",
        "activity_label": activity_label,
    }


if __name__ == "__main__":
    unittest.main()
