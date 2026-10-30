# -*- coding: utf-8 -*-
"""⑥ 一表总览工作簿测试（无网络）。"""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from lib.data_summary import SHEET_CHANNELS, SHEET_TOKENS
from lib.data_summary_exporter import save_workbook_data
from lib.overview_exporter import (
    OVERVIEW_STEM,
    SHEET_INDEX,
    build_overview_filename,
    latest_overview_excel_path,
    list_overview_excel_files,
    overview_workbook_stats,
    save_overview_workbook,
)
from lib.upstream_exporter import SHEET_AZURE, SHEET_VOLC


class OverviewWorkbookTests(unittest.TestCase):
    def test_merge_gateway_and_upstream(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            gateway = root / "excel" / "ai网关令牌及渠道列表_20260101000000.xlsx"
            save_workbook_data(
                gateway,
                channel_env_rows={
                    "test": [{"名称": "ch-a", "分组": "g", "类型": "自定义渠道", "状态": "已启用", "模型列表": "m", "模型重定向": "", "base_url": "https://x", "备注": "渠道管理员备注", "部署时间": ""}],
                    "uat": [],
                    "prod": [],
                },
                token_env_rows={
                    "test": [{"名称": "tk-a", "项目环境": "test", "状态": "已启用", "剩余额度": "", "分组": "g", "可用模型": "m", "包含渠道": "", "IP限制": "", "创建时间": "", "过期时间": ""}],
                    "uat": [],
                    "prod": [],
                },
            )
            dest = root / "excel" / f"{OVERVIEW_STEM}_test.xlsx"
            save_overview_workbook(
                dest,
                gateway_path=gateway,
                volc_rows=[{"接入点名称": "ep-a", "部署名": "ep-1"}],
                azure_rows=[{"订阅": "订阅1", "部署名": "d1"}],
            )
            stats = overview_workbook_stats(dest)
            self.assertEqual(stats["volc_rows"], 1)
            self.assertEqual(stats["azure_rows"], 1)
            self.assertIn(SHEET_CHANNELS["test"], stats["sheets"])
            self.assertIn(SHEET_TOKENS["test"], stats["sheets"])
            self.assertGreaterEqual(stats["sheets"][SHEET_CHANNELS["test"]], 1)
            self.assertGreater(stats["sheets"][SHEET_INDEX], 0)
            from openpyxl import load_workbook
            from lib.upstream_azure import AZURE_HEADERS

            wb = load_workbook(dest, read_only=True)
            try:
                self.assertEqual(wb.sheetnames[0], SHEET_INDEX)
                self.assertIn(SHEET_VOLC, wb.sheetnames)
                self.assertIn(SHEET_AZURE, wb.sheetnames)
                self.assertEqual(wb.sheetnames[-2], SHEET_VOLC)
                self.assertEqual(wb.sheetnames[-1], SHEET_AZURE)
                headers = [cell.value for cell in next(wb[SHEET_AZURE].iter_rows(min_row=1, max_row=1))]
                self.assertEqual(headers, AZURE_HEADERS)
                values = [cell.value for cell in next(wb[SHEET_AZURE].iter_rows(min_row=2, max_row=2))]
                self.assertEqual(values[headers.index("项目")], "N/A")
                self.assertEqual(values[headers.index("负责人")], "N/A")
                channel_headers = [cell.value for cell in next(wb[SHEET_CHANNELS["test"]].iter_rows(min_row=1, max_row=1))]
                channel_values = [cell.value for cell in next(wb[SHEET_CHANNELS["test"]].iter_rows(min_row=2, max_row=2))]
                self.assertEqual(channel_headers[channel_headers.index("base_url") + 1], "备注")
                self.assertEqual(channel_values[channel_headers.index("备注")], "渠道管理员备注")
            finally:
                wb.close()

            (root / "excel" / "上游信息_汇总_20260101000000.xlsx").write_bytes(b"x")
            files = list_overview_excel_files(root)
            self.assertEqual([p.name for p in files], [dest.name])
            self.assertEqual(latest_overview_excel_path(root), dest)

    def test_overview_without_gateway(self):
        with tempfile.TemporaryDirectory() as tmp:
            dest = Path(tmp) / "overview.xlsx"
            save_overview_workbook(
                dest,
                volc_rows=[{"接入点名称": "a"}],
                azure_rows=[],
            )
            stats = overview_workbook_stats(dest)
            self.assertEqual(stats["volc_rows"], 1)
            self.assertEqual(stats["azure_rows"], 0)
            self.assertEqual(stats["gateway_sheets"], [])

    def test_overview_adds_empty_remark_column_for_legacy_gateway(self):
        from openpyxl import Workbook, load_workbook

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            gateway = root / "legacy-gateway.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = SHEET_CHANNELS["test"]
            ws.append(["名称", "base_url", "部署时间"])
            ws.append(["legacy-channel", "https://legacy.example.com", ""])
            wb.save(gateway)
            wb.close()

            dest = root / "overview.xlsx"
            save_overview_workbook(dest, gateway_path=gateway, volc_rows=[], azure_rows=[])
            wb = load_workbook(dest, read_only=True, data_only=True)
            try:
                ws = wb[SHEET_CHANNELS["test"]]
                headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
                values = [cell.value for cell in next(ws.iter_rows(min_row=2, max_row=2))]
                self.assertEqual(headers[headers.index("base_url") + 1], "备注")
                self.assertEqual(values[headers.index("备注")], None)
                base_cell = ws.cell(row=2, column=headers.index("base_url") + 1)
                remark_cell = ws.cell(row=2, column=headers.index("备注") + 1)
                self.assertEqual(base_cell.fill.fgColor.rgb, remark_cell.fill.fgColor.rgb)
                self.assertNotEqual(getattr(base_cell.font.color, "rgb", None), "FF7030A0")
                self.assertEqual(remark_cell.font.color.rgb, "FF7030A0")
            finally:
                wb.close()

    def test_filename_stamp_beats_mtime(self):
        import os

        with tempfile.TemporaryDirectory() as tmp:
            excel = Path(tmp) / "excel"
            excel.mkdir()
            newer_stamp = excel / f"{OVERVIEW_STEM}_20260825120000.xlsx"
            older_stamp = excel / f"{OVERVIEW_STEM}_20260825100000.xlsx"
            newer_stamp.write_bytes(b"new")
            older_stamp.write_bytes(b"old")
            os.utime(older_stamp, None)
            files = list_overview_excel_files(Path(tmp))
            self.assertEqual(files[0].name, newer_stamp.name)
            self.assertEqual(latest_overview_excel_path(Path(tmp)), newer_stamp)

    def test_resolve_prefers_newer_source_file(self):
        from lib.overview_exporter import resolve_azure_rows, resolve_gateway_source, resolve_volc_rows
        from lib.upstream_exporter import save_combined_upstream_workbook, save_upstream_workbook
        from lib.upstream_volc import UPSTREAM_HEADERS
        from lib.upstream_azure import AZURE_HEADERS

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            excel = root / "excel"
            excel.mkdir()
            old_combined = excel / "上游信息_汇总_20260825100000.xlsx"
            save_combined_upstream_workbook(
                old_combined,
                [{"接入点名称": "old-ep", "部署名": "old"}],
                [{"订阅": "订阅1", "部署名": "old-az", "Azure OpenAI": "old"}],
            )
            new_volc = excel / "上游信息_火山云_20260825120000.xlsx"
            save_upstream_workbook(
                new_volc,
                [{"接入点名称": "new-ep", "部署名": "new"}],
                headers=list(UPSTREAM_HEADERS),
                sheet_name=SHEET_VOLC,
            )
            new_azure = excel / "上游信息_Azure_20260825130000.xlsx"
            save_upstream_workbook(
                new_azure,
                [{"订阅": "订阅1", "部署名": "new-az", "Azure OpenAI": "acc", "项目": "产研", "负责人": "乔加奇"}],
                headers=list(AZURE_HEADERS),
                sheet_name=SHEET_AZURE,
            )
            volc_rows = resolve_volc_rows(root)
            azure_rows = resolve_azure_rows(root)
            self.assertEqual(volc_rows[0]["接入点名称"], "new-ep")
            self.assertEqual(azure_rows[0]["部署名"], "new-az")
            self.assertEqual(azure_rows[0]["项目"], "产研")

            gw_old = excel / "ai网关令牌及渠道列表_20260825100000.xlsx"
            save_workbook_data(
                gw_old,
                channel_env_rows={"test": [{"名称": "old-ch", "分组": "g", "类型": "自定义渠道", "状态": "已启用", "模型列表": "m", "模型重定向": "", "base_url": "https://x", "部署时间": ""}], "uat": [], "prod": []},
                token_env_rows={"test": [], "uat": [], "prod": []},
            )
            overview_new = excel / f"{OVERVIEW_STEM}_20260825140000.xlsx"
            save_overview_workbook(
                overview_new,
                gateway_path=gw_old,
                volc_rows=[{"接入点名称": "ov-ep"}],
                azure_rows=[{"订阅": "订阅1", "部署名": "ov-az"}],
            )
            chosen = resolve_gateway_source(root)
            self.assertEqual(chosen.name, overview_new.name)

    def test_overview_filename_and_legacy(self):
        from datetime import datetime

        self.assertEqual(
            build_overview_filename(datetime(2026, 8, 25, 10, 32, 0)),
            "ai网关令牌及渠道列表_一表总览_20260825103200.xlsx",
        )
        self.assertTrue(OVERVIEW_STEM.startswith("ai网关令牌及渠道列表_"))
        with tempfile.TemporaryDirectory() as tmp:
            excel = Path(tmp) / "excel"
            excel.mkdir()
            legacy = excel / "一表总览_20260101000000.xlsx"
            current = excel / f"{OVERVIEW_STEM}_20260825103200.xlsx"
            legacy.write_bytes(b"old")
            current.write_bytes(b"new")
            names = [p.name for p in list_overview_excel_files(Path(tmp))]
            self.assertEqual(names, [current.name, legacy.name])

    def test_scan_channel_key_status_missing_and_filled(self):
        from lib.data_summary import CHANNEL_HEADERS_WITH_KEY, scan_channel_key_status

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            dest = root / "excel" / f"{OVERVIEW_STEM}_keys.xlsx"
            save_overview_workbook(
                dest,
                volc_rows=[{"接入点名称": "ep-keep"}],
                azure_rows=[],
            )
            save_workbook_data(
                dest,
                channel_env_rows={
                    "test": [
                        {"名称": "ch-a", "分组": "g", "类型": "自定义渠道", "状态": "已启用", "模型列表": "m", "模型重定向": "", "base_url": "https://x", "密钥": "sk-a", "部署时间": ""},
                        {"名称": "ch-b", "分组": "g", "类型": "自定义渠道", "状态": "已启用", "模型列表": "m", "模型重定向": "", "base_url": "https://x", "密钥": "", "部署时间": ""},
                    ],
                    "uat": [],
                    "prod": [],
                },
                token_env_rows={"test": [], "uat": [], "prod": []},
                channel_headers=CHANNEL_HEADERS_WITH_KEY,
            )
            status = scan_channel_key_status(dest)
            self.assertEqual(status["envs"]["test"]["keys_filled"], 1)
            self.assertEqual(status["envs"]["test"]["keys_missing"], 1)
            self.assertEqual(status["pending_total"], 1)
            self.assertTrue(status["can_continue"])
            from openpyxl import load_workbook

            wb = load_workbook(dest, read_only=True)
            try:
                self.assertIn(SHEET_VOLC, wb.sheetnames)
            finally:
                wb.close()

    def test_no_key_column_counts_as_missing(self):
        from lib.data_summary import scan_channel_key_status

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            gateway = root / "excel" / "ai网关令牌及渠道列表_20260101000000.xlsx"
            save_workbook_data(
                gateway,
                channel_env_rows={
                    "test": [{"名称": "ch-a", "分组": "g", "类型": "自定义渠道", "状态": "已启用", "模型列表": "m", "模型重定向": "", "base_url": "https://x", "部署时间": ""}],
                    "uat": [],
                    "prod": [],
                },
                token_env_rows={"test": [], "uat": [], "prod": []},
            )
            status = scan_channel_key_status(gateway)
            self.assertEqual(status["envs"]["test"]["channels"], 1)
            self.assertEqual(status["envs"]["test"]["keys_filled"], 0)
            self.assertEqual(status["pending_total"], 1)


if __name__ == "__main__":
    unittest.main()
